# -*- coding: utf-8 -*-
import sys
import os
import base64
import io
import datetime
import time
import dash
import dash_bootstrap_components as dbc
import dash_core_components as dcc
import dash_daq as daq  # #
import dash_html_components as html
import dash_table  # #
import pandas as pd
import plotly.graph_objects as go
import plotly.express as px
from dash import no_update
from dash.dependencies import Input, Output, State,MATCH
from dash.exceptions import PreventUpdate
import numpy as np
from numpy import trapz
from flask import send_file
from openpyxl import Workbook, load_workbook
from dash_extensions.enrich import ServersideOutput
import OpenOPC
from sshtunnel import SSHTunnelForwarder
import mariadb
import mysql.connector
import pywintypes
import win32timezone

pywintypes.datetime = pywintypes.TimeType


def find_data_file(filename):
    if getattr(sys, 'frozen', False):
        # The application is frozen
        datadir = os.path.dirname(sys.executable)
    else:
        # The application is not frozen
        # Change this bit to match where you store your data files:
        datadir = os.path.dirname(__file__)
    return os.path.join(datadir, filename)

BS =[ "https://stackpath.bootstrapcdn.com/bootstrap/4.5.1/css/bootstrap.min.css",
    {
        'href': 'https://use.fontawesome.com/releases/v5.8.1/css/all.css',
        'rel': 'stylesheet',
        'integrity': 'sha384-50oBUHEmvpQ+1lW4y57PTFmhCaXp0ML5d60M1M7uH2+nqUivzIebhndOJK28anvf',
        'crossorigin': 'anonymous'
    }
]

# Initialize the app
app = dash.Dash(__name__,external_stylesheets=BS, suppress_callback_exceptions=True, assets_folder=find_data_file('assets/'), update_title='Loading...',
                meta_tags=[{'name': 'viewport', 'content': 'width=device-width, initial-scale=2.0, maximum-scale=1.2, minimum-scale=0.5'}],
                )
server = app.server
app.config.suppress_callback_exceptions = True

# Referance values
extra_data_list = ['General Ref Point',
    'Mass de Bois', 'Volume gaz', 'Vitesse de rotation', 'Puissance Thermique',
    'Puissance Electrique', 'CO', 'CO2', 'NO', 'NOX', 'Temperature de Fumée'
]
# Layout design
app.layout = html.Div([
    dcc.Location(id='url', refresh=False),
    html.Div(id='page-content'),
])
index_page = html.Div([
    html.Div(html.Div(html.Div(children=[
        html.H3('RCC/KN'),
        dcc.Link('Start', href='/RCC_KN')
    ], className='content'), className='box'), className='card'),
    html.Div(html.Div(html.Div(children=[
        html.H3('Enerbat'),
        html.P('Works on RealTime in Local'),
        dcc.Link('Start', href='/realTime_Enerbat')
    ], className='content'), className='box'), className='card'),

    html.Div(html.Div(html.Div(children=[
        html.H3('DATABASE'),
        html.P('Choose your Database and make analysis'),
        dcc.Link('Start', href='/Database')
    ], className='content'), className='box'), className='card'),
    html.Div(html.Div(html.Div(
        children=[html.H3('Data Analysis'),
                  html.P('Upload and run files such as .xlsx/.xls/.csv'),
                  dcc.Link('Start', href="/File")
                  ], className='content'), className='box'), className='card'),

], className='container')

colors = {
    'background': '#e6f7f6',
    'text': '#E1E2E5',
    'figure_text': 'black'}
page_1_layout = html.Div(
    className='main_container',
    children=[
        html.Div(id='fourcolumnsdivusercontrols', className='four-columns-div-user-controls',
                 children=[
                     html.Div([html.Div([daq.PowerButton(id='my-toggle-switch',
                                               label={'label': 'Open page',
                                                      'style': {'fontSize': '22px', 'fontWeight': 'bold',},},
                                               labelPosition='bottom', on=False, size=100, color='green',style = {'marginTop':'1rem'},
                                               className='dark-theme-control'),]),
                               html.Div([html.Div([html.Div([html.I(className="fas fa-home"), ]),
                                                                    dcc.Link('Home', href='/', id='link1')],
                                                                   className='icon_position1'),
                                                          html.Div([html.Div([html.I(className="fas fa-database"), ]),
                                                                    dcc.Link('Database', href='/Database', id='link2')],
                                                                   className='icon_position1'),
                                                        html.Div([html.Div([html.I(className="fas fa-server"), ]),
                                                                     dcc.Link('Enerbat', href='/realTime_Enerbat', id='link3')],
                                                                   className='icon_position1'),
                                                          html.Div([html.Div([html.I(className="fas fa-chart-line"), ]),
                                                                    dcc.Link('RCC/KN', href='/RCC_KN',
                                                                             id='link4')], className='icon_position1'),
                                         ], style = {'margin' : '2rem 2rem 0 2rem'},className='abpage1'),
                               ], className='abpower'),
                     html.Div(
                         dcc.Upload(
                             id='upload-data',
                             children=html.Div([
                                 'Drag and Drop or ',
                                 html.A('Select Files for work')
                             ]),
                             style={
                                 'marginLeft': '-2rem',
                                 'visibility': 'hidden',
                             },
                             # Allow multiple files to be uploaded
                             multiple=True,
                         ),
                     ),
                     html.Div(id="openOPCDiv", children=[], style={'visibility': 'hidden'}),
                     html.Div(className='userControlDownSide',
                              children=[
                                  html.Div(className='userControlDownLeftSide',
                                           children=[
                                                html.Div(id="opcLoad",
                                                        className='div-for-dropdown',
                                                        children=[], ),
                                                html.Div(dcc.Interval(
                                                   id='interval',
                                                   interval=5000,
                                                   n_intervals=3,
                                               )),
                                                html.Div(id='pointLeftFirst', children=[], style={'display': 'None'}),
                                                html.Div(id='pointLeftSecond', children=[], style={'display': 'None'}),
                                                html.Div(id='pointRightFirst', children=[], style={'display': 'None'}),
                                                html.Div(id='pointRightSecond', children=[], style={'display': 'None'}),
                                                html.Div(id='pointLeftFirstTab4', children=[],style={'display': 'None'}),
                                                html.Div(id='pointLeftSecondTab4', children=[],style={'display': 'None'}),
                                                html.Div(id='pointRightFirstTab4', children=[],style={'display': 'None'}),
                                                html.Div(id='pointRightSecondTab4', children=[],style={'display': 'None'}),
                                                html.Div(id='leftSideChecklistValue', children=[],style={'display': 'None'}),
                                                html.Div(id='leftSidedroptValue', children=[],style={'display': 'None'}),
                                                html.Div(id='leftSideChecklistValueHidden', children=[],style={'display': 'None'}),
                                                html.Div(id='leftSideChecklistValueHiddendb', children=[],style={'display': 'None'}),
                                                html.Div(id='deletedval', children=[], style={'display': 'None'}),
                                                html.Div(id='deletedvalref', children=[], style={'display': 'None'}),
                                                html.Div(id='leftSideChecklistValueHiddenTab4', children=[],style={'display': 'None'}),
                                                html.Div(id='tab2hiddenValuex_axis', style={'display': 'None'},children=[]),
                                                html.Div(id='tab2hiddenValuey_axis', style={'display': 'None'}, children=[]),
                                                html.Div(id='tab4hiddenValuex_axissecond', style={'display': 'None'},children=[]),
                                                html.Div(id='tab4hiddenValuey_axissecond', style={'display': 'None'},children=[]),
                                                html.Div(id='tab4hiddenValuex_axis', style={'display': 'None'},children=[]),
                                                html.Div(id='hiddenTextHeader', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenTextNote', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenTextxaxis', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenTextyaxis', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenTexty1axis', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenTexty2axis', children=[], style={'display': 'None'}),
                                                html.Div(id='selectyaxishidden', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenTextHeader4', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenTextNote4', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenTextxaxis4', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenTextyaxis4', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenShapeVal', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenTexty1axis4', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenTexty2axis4', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenShapeVal', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenShapeVal', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenShapeDate', children=[],style={'display': 'None'}), ], ),
                                                html.Div(id='hiddenDifferance', children=[], style={'display': 'None'}),
                                                dcc.Store(id='datastore'),
                                                html.Div(id='retrieve', children=[], style={'display': 'None'}),
                                                html.Div(id='datatablehidden', children=[], style={'display': 'None'}),
                                                html.Div(id='radiographhidden', children=[], style={'display': 'None'}),
                                                html.Div(id='sliderHeightTab1hidden', children=[], style={'display': 'None'}),
                                                html.Div(id='sliderWidthTab1hidden', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenShapeValtab4', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenShapeDatetab4', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenDifferancetab4', children=[], style={'display': 'None'}),
                                                html.Div(id='retrievetab4', children=[], style={'display': 'None'}),
                                                html.Div(id='datatablehiddentab4', children=[], style={'display': 'None'}),
                                                html.Div(id='radiographhiddentab4', children=[], style={'display': 'None'}),
                                                html.Div(id='sliderHeightTab1hiddentab4', children=[], style={'display': 'None'}),
                                                html.Div(id='sliderWidthTab1hiddentab4', children=[], style={'display': 'None'}),
                                                html.Div(id='minimumValueGraphhiddenfirst', children=[], style={'display': 'None'}),
                                                html.Div(id='minimumValueGraphhiddensecond', children=[], style={'display': 'None'}),
                                                html.Div(id='firstchoosenvalhidden', children=[], style={'display': 'None'}),
                                                html.Div(id='secondchoosenvalhidden', children=[], style={'display': 'None'}),
                                                html.Div(id='leftintegralfirsthidden', children=[], style={'display': 'None'}),
                                                html.Div(id='leftintegralsecondhidden', children=[], style={'display': 'None'}),
                                                html.Div(id='rightintegralfirsthidden', children=[], style={'display': 'None'}),
                                                html.Div(id='rightintegralsecondhidden', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddencharactere', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenvaluechange', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddencleanshape', children=[], style={'display': 'None'}),
                                                html.Div(id='firstchoosenvalhiddentab4', children=[], style={'display': 'None'}),
                                                html.Div(id='secondchoosenvalhiddentab4', children=[], style={'display': 'None'}),
                                                html.Div(id='leftintegralfirsthiddentab4', children=[], style={'display': 'None'}),
                                                html.Div(id='leftintegralsecondhiddentab4', children=[], style={'display': 'None'}),
                                                html.Div(id='rightintegralfirsthiddentab4', children=[], style={'display': 'None'}),
                                                html.Div(id='rightintegralsecondhiddentab4', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenchoosenChecklistLeft', children=[], style={'display': 'None'}),
                                                html.Div(id='writeexcelhidden', children=[], style={'display': 'None'}),
                                                html.Div(id='writeexcelhiddenTab4', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenrecord1', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenrecord2', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenrecord3', children=[], style={'display': 'None'}),
                                                html.Div(id='hiddenrecord4', children=[], style={'display': 'None'}),
                                                html.Div(id='inputRightY_axishidden', children=[], style={'display': 'None'}),
                                                html.Div(id='inputRightX_axishidden', children=[], style={'display': 'None'}),
                                                html.Div(id='checklistvaleurhidden', children=[], style={'display': 'None'}),
                                                html.Div(id='checklistvaleurhidden2', children=[], style={'display': 'None'}),
                                                html.Div(id='shiftaxisdrophidden', children=[], style={'display': 'None'}),
                                                html.Div(id='shift_x_axishidden', children=[], style={'display': 'None'}),
                                                html.Div(id='shift_y_axishidden', children=[], style={'display': 'None'}),
                                                html.Div(id='tab1sendhidden', children=[], style={'display': 'None'}),
                                                html.Div(id='shiftaxisdroptab4hidden', children=[], style={'display': 'None'}),
                                                html.Div(id='shift_x_axistab4hidden', children=[], style={'display': 'None'}),
                                                html.Div(id='shift_y_axistab4hidden', children=[], style={'display': 'None'}),
                                                html.Div(id='output_s', children=[], style={'display': 'None'}),
                                                html.Div(id='radiographtab4hidden', children=[], style={'display': 'None'}),
                                                html.Div(id='reelhidden11', children=[], style={'display': 'None'}),
                                                html.Div(id='reelhidden12', children=[], style={'display': 'None'}),
                                                html.Div(id='reelhidden13', children=[], style={'display': 'None'}),
                                                html.Div(id='reelhidden14', children=[], style={'display': 'None'}),
                                                html.Div(id='reelhidden15', children=[], style={'display': 'None'}),
                                                html.Div(dcc.Graph(id='graphhidden',
                                                     config={'displaylogo': False},
                                                     style={'display': 'None'},
                                                     figure={
                                                         'layout': {'legend': {'tracegroupgap': 0},}
                                                            }
                                                     ), ),
                              ]),]),
        html.Div(id='eightcolumnsdivforcharts', className='eight-columns-div-for-charts',
                 children=[
                     html.Div(
                         className='right-upper',
                         children=[
                             html.Div([
                                 dcc.Tabs(
                                     id="tabs-with-classes",
                                     value='tab-1',
                                     parent_className='custom-tabs',
                                     className='custom-tabs-container',
                                     children=[
                                         dcc.Tab(
                                             id="tab1",
                                             label='Work on unique parameter',
                                             value='tab-1',
                                             className='custom-tab',
                                             selected_className='custom-tab--selected',
                                             children=[],
                                         ),
                                         dcc.Tab(
                                             id="tab4",
                                             label='Work on Different Parameters',
                                             value='tab-4',
                                             className='custom-tab',
                                             # style={'visibility': 'hidden'},
                                             selected_className='custom-tab--selected'
                                         ),
                                         dcc.Tab(
                                             id="tab5",
                                             label='Tab for one option',
                                             value='tab-5',
                                             className='custom-tab',
                                             style={'visibility': 'hidden'},
                                             selected_className='custom-tab--selected'
                                         ),
                                     ]),
                                 html.Div(id='tabs-content-classes'),

                             ]),

                         ]),

                 ]
                 ),
    ]),

page_2_layout = html.Div(
                    [html.Div(
                        [html.Div([
                                 html.Div([html.Div([html.Div([html.I(className="fas fa-home"), ]),
                                                                    dcc.Link('Home', href='/', id='link1')],
                                                                   className='icon_position'),
                                                          html.Div([html.Div([html.I(className="far fa-file-excel"), ]),
                                                                    dcc.Link('Analysis', href='/File', id='link5')],
                                                                   className='icon_position'),
                                                          html.Div([html.Div([html.I(className="fas fa-server"), ]),
                                                                     dcc.Link('Enerbat', href='/realTime_Enerbat', id='link3')],
                                                                   className='icon_position'),
                                                          html.Div([html.Div([html.I(className="fas fa-chart-line"), ]),
                                                                    dcc.Link('RCC/KN', href='/RCC_KN',
                                                                             id='link4')], className='icon_position'),
                                           ],
                                           style = {'margin':'-3rem 0 1rem 0'},className='abpage2'),
                                        dbc.ButtonGroup([dbc.Button("Database Activate", id="activatedb", n_clicks=0,
                                                     size="lg", className='mr-1', color="success",
                                                     style={'width': '25rem', 'marginBottom' : '1rem','borderRadius' : '2rem'}
                                                     ),
                                                     dbc.Button("Database Deactivate", id="deactivatedb", n_clicks=0,
                                                             size="lg", className='mr-1', color="danger",
                                                             style={'width': '25rem', 'marginBottom' : '1rem','borderRadius' : '2rem'}
                                                             )]),
                                 html.Div([dbc.Input(id='db_Ip',
                                           type="text",
                                           debounce=True,
                                           min=-10000, max=10000, step=1,
                                           bs_size="mr",
                                           style={'width': '11rem', 'marginTop': '1.5rem'},
                                           autoFocus=True,
                                           placeholder="Enter your Username"),
                                         dbc.Input(id='db_password',
                                                   type="password",
                                                   debounce=True,
                                                   min=-10000, max=10000, step=1,
                                                   bs_size="mr",
                                                   style={'width': '11rem', 'marginTop': '1.5rem'},
                                                   autoFocus=True,
                                                   placeholder="Enter your password"),], className = 'ab'),

                                ], className='page2design1'),
                   html.Div([
                       dcc.Dropdown(id='db_name',
                                    options=[{'label': i, 'value': i}
                                             for i in ['rcckn', 'enerbat']],
                                    multi=False,
                                    style={'cursor': 'pointer', 'marginTop': '5px'},
                                    className='stockSelectorClass3',
                                    clearable=True,
                                    placeholder='Select Database',

                                    ),
                       dcc.Dropdown(id='dbvalchoosen',
                                    # options=[{'label': i, 'value': i}
                                    #          for i in df.columns],
                                    multi=False,
                                    style={'cursor': 'pointer', 'marginTop': '5px'},
                                    className='stockSelectorClass3',
                                    clearable=True,
                                    placeholder='Select Table ...',

                                    ),

                       dcc.Dropdown(id='dbvalname',
                                    # options=[{'label': i, 'value': i}
                                    #          for i in df.columns],
                                    multi=True,
                                    style={'cursor': 'pointer', 'marginTop': '13px'},
                                    className='stockSelectorClass3',
                                    clearable=True,
                                    placeholder='Select your parameters...',
                                    ),

                       dcc.Dropdown(id='dbvaldate',
                                    # options=[{'label': i, 'value': i}
                                    #          for i in df.columns],
                                    multi=True,
                                    style={'cursor': 'pointer', 'marginTop': '13px'},
                                    className='stockSelectorClass3',
                                    clearable=False,
                                    placeholder='Select your time...',
                                    ), ], className='aadb'),
                   html.Div([
                            daq.BooleanSwitch(
                                  id="calculintegraldb",
                                  on=False,
                                  label="Calculate Integral",
                                  color= '#1f78b4',
                                  labelPosition="bottom",
                                  style = {'margin': '5rem'}
                            )
                        ])

                ], className='abcdb'),
     dcc.Store(id='memory-output'),
     html.Div(id='dbcheck', children=
     [html.Div([html.Div([dcc.Dropdown(id='firstChoosenValuedb',
                                       options=[{'label': i, 'value': i} for i in
                                                []],
                                       multi=False,
                                       style={'cursor': 'pointer', 'width': '180px'},

                                       clearable=True,
                                       placeholder='First Value...',
                                       ),
                          dbc.Input(id='leftIntegralFirstdb',
                                    type="text",
                                    debounce=True,
                                    min=-10000, max=10000, step=1,
                                    bs_size="sm",
                                    style={'width': '8rem', 'marginTop': '1.5rem'},
                                    autoFocus=True,
                                    placeholder="first point"),
                          dbc.Input(id='leftIntegralSeconddb',
                                    type="text",
                                    debounce=True,
                                    min=-10000, max=10000, step=1,
                                    bs_size="sm",
                                    style={'width': '8rem', 'marginTop': '1.5rem'},
                                    autoFocus=True,
                                    placeholder="second point"),
                          dbc.Input(id='leftIntegraldb',
                                    type="text",
                                    min=-10000, max=10000, step=1,
                                    bs_size="sm",
                                    style={'width': '9rem', 'marginTop': '1.5rem'},
                                    autoFocus=True,
                                    placeholder="total integration"),
                          ]), html.Div([html.Button("Save", id="write_exceldb", n_clicks=0,
                                                    style={'fontSize': '1rem', 'width': '4rem',
                                                           'margin': '1rem'},
                                                    ),
                                        html.A(html.Button("Download Data",
                                                           id='download_datadb',
                                                           n_clicks=0,
                                                           style={'fontSize': '1rem',
                                                                  'width': '9rem',
                                                                  'margin': '1rem'}, ),
                                               id='download_exceldb',
                                               # # download="rawdata.csv",
                                               href="/download_exceldb/",
                                               # target="_blank"
                                               ),
                                        ], className='ad')

                ]),
      html.Div([dbc.Checklist(
          id='operateurdb',
          options=[{'label': i, 'value': i} for i in
                   ['Plus', 'Moins', 'Multiplie', 'Division']],
          value=[],
          labelStyle={'display': 'Block'},
      ), ]),
      html.Div([
          dcc.Dropdown(id='secondChoosenValuedb',
                       options=[{'label': i, 'value': i} for i in
                                []],
                       multi=False,
                       style={'cursor': 'pointer', 'width': '180px'},

                       clearable=True,
                       placeholder='Second Value...',
                       ),
          dbc.Input(id='rightIntegralFirstdb',
                    type="text",
                    min=-10000, max=10000, step=1,
                    bs_size="sm",
                    style={'width': '8rem', 'marginTop': '1.5rem'},
                    autoFocus=True,
                    placeholder="first point"),
          dbc.Input(id='rightIntegralSeconddb',
                    type="text",
                    min=-10000, max=10000, step=1,
                    bs_size="sm",
                    style={'width': '8rem', 'marginTop': '1.5rem'},
                    autoFocus=True,
                    placeholder="second point"),
          dbc.Input(id='rightIntegraldb',
                    type="text",
                    min=-10000, max=10000, step=1,
                    bs_size="sm",
                    style={'width': '9rem', 'marginTop': '1.5rem'},
                    autoFocus=True,
                    placeholder="total integration")
      ]),
      html.Div([dbc.Input(id='operationdb',
                          type="text",
                          min=-10000, max=10000, step=1,
                          bs_size="sm",
                          style={'width': '10rem', 'marginTop': '2rem',
                                 'height': '5rem', 'textAlign': 'center'},
                          autoFocus=True,
                          placeholder="result"),
                dbc.Input(id='intersectiondb',
                          type="text",
                          min=-10000, max=10000, step=1,
                          bs_size="sm",
                          style={'width': '10rem', 'marginTop': '2rem',
                                 'height': '2rem', 'textAlign': 'center'},
                          autoFocus=True,
                          placeholder="Intersection")], className='aa')
      ], style={'display': 'None'},
              className='abdbase'),
     html.Div([html.Div([html.Div([html.Div(dcc.RadioItems(id="radiographdb",
                               options=[
                                   {'label': 'Point', 'value': 'markers'},
                                   {'label': 'Line', 'value': 'lines'},
                                   {'label': 'Line + Point', 'value': 'lines+markers'}],
                               value='markers',
                               labelClassName='groupgraph2',
                               labelStyle={'margin': '10px'},
                               inputStyle={'margin': '10px'}
                               ), className='abdbside-db'),
                        html.Div([dcc.Dropdown(id="dropadddb",
                                                       options=[
                                                           {'label': 'Note', 'value': 'note'},
                                                           {'label': 'Header', 'value': 'header'},
                                                           {'label': 'x-axis', 'value': 'x_axis'},
                                                           {'label': 'y-axis', 'value': 'y_axis'},
                                                           {'label': 'y2-axis', 'value': 'y2_axis'},
                                                           {'label': 'y3-axis', 'value': 'y3_axis'},],
                                                       value='header',
                                                       ),
                                          dcc.Textarea(
                                              id='textareadb',
                                              value='',
                                              style={'width': '15rem', 'marginTop': '0.5rem'},
                                          ),
                                          dbc.Button("Add Text", id='addTextdb', n_clicks=0,
                                                                   className='mr-1', color="primary", style={'margin': '1rem 1rem 0 6vw'}),

                                          dcc.Dropdown(id='characteredb',
                                                    options=[{'label': i, 'value': i} for i in [8,12,14,16,18,20,22,24,26,28]],
                                                    value = 16,
                                                    multi=False,
                                                    style={'cursor': 'pointer', 'width': '120px', 'margin': '1rem'},
                                                    clearable=True,
                                                    placeholder='Charactere...',
                                                    ),
                                          ], className='aatabdb'),], className='abpagedb'),

               html.Div([html.Div([
                         html.Div(dcc.Loading(type='cube',children = dcc.Graph(id="getdbgraph",
                                            config={'displayModeBar': True,
                                                    'scrollZoom': True,
                                                    'displaylogo': False,
                                                    'modeBarButtonsToAdd': [
                                                        'drawline',
                                                        'drawrect',
                                                        'drawopenpath',
                                                        'select2d',
                                                        'eraseshape',

                                                    ]},
                                            style={'marginTop': '20px', },
                                            figure={
                                                'layout': {'legend': {'tracegroupgap': 0},

                                                           }
                                            }

                                            ), ),),
                                        html.Div([daq.BooleanSwitch(
                                                             id="selectyaxisdb",
                                                             label="Multiple Y-axis",
                                                             labelPosition="bottom",
                                                             color='red',

                                                         )
                                                             ,
                                                             dbc.Tooltip(
                                                                 "You can select maximum 3 y-axis",
                                                                 target="selectyaxisdb",
                                                                 placement='bottom',
                                                             ),
                                                         ], className='pluraly'),
                         html.Div(daq.Slider(id="sliderHeightdb",
                                             max=2100,
                                             min=400,
                                             value=530,
                                             step=100,
                                             size=400,
                                             vertical=True,
                                             updatemode='drag'), style={'margin': '20px'})], className='abcdb_graph'),

               html.Div([daq.Slider(id="sliderWidthdb",
                                             max=2000,
                                             min=600,
                                             value=1000,
                                             step=100,
                                             size=600,

                                             updatemode='drag')],className='design_page1'),
                         ],className='design_page2'),],className='acsecond'),
               html.Div(dash_table.DataTable(id="getdbtable",
                                             editable=True,
                                             page_size=50,
                                             style_table={'height': '500px', 'overflowY': 'auto', 'width': '98%'},
                                             style_cell={
                                                 'overflow': 'hidden',
                                                 'textOverflow': 'ellipsis',
                                                 'maxWidth': 0,
                                                 'fontSize': '1rem',
                                                 'TextAlign': 'center',
                                             },

                                             fixed_rows={'headers': True},
                                             style_header={
                                                 'backgroundColor': 'rgb(230, 230, 230)',
                                                 'fontWeight': 'bold'
                                             },
                                             filter_action="native",
                                             sort_action="native",
                                             sort_mode="multi",
                                             column_selectable="single",
                                             selected_columns=[],
                                             selected_rows=[],
                                             page_action="native",
                                             page_current=0,
                                             export_format='xlsx',
                                             export_headers='display',
                                             merge_duplicate_headers=True),style = {'width': '80%', 'margin': 'auto'}),
               html.Div(id='pointLeftFirstdb', children=[], style={'display': 'None'}),
               html.Div(id='pointLeftSeconddb', children=[], style={'display': 'None'}),
               html.Div(id='pointRightFirstdb', children=[], style={'display': 'None'}),
               html.Div(id='pointRightSeconddb', children=[], style={'display': 'None'}),
               html.Div(id='hiddenrecord1db', children=[], style={'display': 'None'}),
               html.Div(id='hiddenrecord3db', children=[], style={'display': 'None'}),
               html.Div(id='hiddenrecord4db', children=[], style={'display': 'None'}),
               html.Div(id='writeexcelhiddendb', children=[], style={'display': 'None'}),
               html.Div(id='hiddenTextHeaderdb', children=[], style={'display': 'None'}),
               html.Div(id='hiddenTextNotedb', children=[], style={'display': 'None'}),
               html.Div(id='hiddenTextxaxisdb', children=[], style={'display': 'None'}),
               html.Div(id='hiddenTextyaxisdb', children=[], style={'display': 'None'}),
               html.Div(id='hiddenTexty1axisdb', children=[], style={'display': 'None'}),
               html.Div(id='hiddenTexty2axisdb', children=[], style={'display': 'None'}),

               ],className='four-columns-div-user-controlsreel' ), ], ),

page_3_layout = html.Div([html.Div([
                            html.Div([html.Div([html.Div([
                                                          html.Div([html.Div([html.I(className="fas fa-home"), ]),
                                                                    dcc.Link('Home', href='/', id='link1')],
                                                                   className='icon_position3'),
                                                          html.Div([html.Div([html.I(className="far fa-file-excel"), ]),
                                                                    dcc.Link('Analysis', href='/File', id='link5')],
                                                                   className='icon_position3'),
                                                          html.Div([html.Div([html.I(className="fas fa-database"), ]),
                                                                    dcc.Link('Database', href='/Database', id='link2')],
                                                                   className='icon_position3'),
                                                          html.Div([html.Div([html.I(className="fas fa-chart-line"), ]),
                                                                    dcc.Link('RCC/KN', href='/RCC_KN',
                                                                             id='link4')], className='icon_position3'),
                                                          ],style = {'marginTop':'1rem'},className='ab-page-3'),
                                                daq.PowerButton(id='my-toggle-switch-reel',
                                                                       label={'label': 'Connection',
                                                                              'style': {'fontSize': '22px', 'fontWeight': 'bold',},},
                                                                       labelPosition='bottom', on=False, size=100, color="green",style = {'margin':'3rem 0'},
                                                                       className='dark-theme-control'),]),
                                               dcc.Store(id='get_data_from_modbus'),
                                               html.Div(id='data_to_store_id', children=[], style={'display': 'None'}),
                                               html.Div(id='data_to_store_value', children=[], style={'display': 'None'}),
                                               html.Div(id='data_to_store_qualite', children=[], style={'display': 'None'}),
                                               html.Div(id='data_to_store_date', children=[], style={'display': 'None'}),
                                               dcc.Interval(
                                                   id='interval_component',
                                                   disabled=True,
                                                   interval=1 * 1000,  # in milliseconds
                                                   n_intervals=0),
                            html.Div([
                                     dcc.Dropdown(id='realvalue',
                                                       multi=True,
                                                       style={'cursor': 'pointer'},
                                                       className='stockSelectorClassPage3',
                                                       clearable=True,
                                                       placeholder='Select Value',
                                                       ),
                                          html.Div([html.P('Enter interval value (Second)', style={'margin':'1em 2em 2em 5em'}),
                                                  dbc.Input(id='interval_value', type="text", value='1',
                                                            min=0, max=1000000000, step=1, bs_size="lg", style={'width': '6rem', 'margin':'1em'}, ),], className='acsecond'),

                                      html.Div([
                                                html.Div(
                                                    [
                                                        dbc.Button("Send Values to Database", id='download_reel_db', n_clicks=0, size="lg",
                                                                   className='mr-1', color="primary", style={'margin': '-3rem 1rem 0 6vw'}),
                                                        dbc.Modal(
                                                            [
                                                                dbc.ModalHeader("Save Your Table In Database"),
                                                                dbc.Input(id='input_databasename',
                                                                          type="text",
                                                                          min=-10000, max=10000, step=1, bs_size="sm",
                                                                          style={'width': '31rem', },
                                                                          placeholder = 'Enter Database name',
                                                                          autoFocus=True, ),
                                                                dbc.Input(id='input_tablename',
                                                                          type="text",
                                                                          min=-10000, max=10000, step=1, bs_size="sm",
                                                                          style={'width': '31rem', },
                                                                          placeholder = 'Enter Table name',
                                                                          autoFocus=True, ),
                                                                dbc.ModalFooter([
                                                                    dbc.Button("Okey", id="ok_reel", className='ml-auto'),
                                                                    dbc.Button("Close", id="close_reel", className='ml-auto')]
                                                                ),
                                                            ],
                                                            id="modal_reel",
                                                        ), ]),

                                                ],style = {'margin':'3rem'}, className='abcd'),

                                      ]),html.Div([html.Div([html.Div([daq.Knob(id='HVA1IN',color='blue',label={"label": "HV A1 IN", "style": {'color': 'blue','marginBottom':'-1rem'}},
                                                                                value=0,size= 150,max=100, labelPosition = 'top'),
                                                                      daq.LEDDisplay(id = 'HVA1INLED',value = 0, color= 'blue', size=30,style = {'margin':'-2rem 0 2rem 0'})],
                                                                      className = 'aadbknob'),
                                                             html.Div([daq.Knob(id='HVA2IN',color='blue', label={"label": "HV A2 IN", "style": {'color': 'blue','marginBottom':'-1rem'}},
                                                                                value=0,size=150, max=100, labelPosition='top'),
                                                                       daq.LEDDisplay(id='HVA2INLED', value = 0, color='blue', size=30, style={'margin': '-2rem 0 2rem 0'})],
                                                                      className='aadbknob'),], className = 'abc'),
                                        html.Div([
                                                html.Div(
                                                    [
                                                        dbc.Button("Send Valve Values to Server", id='download_reel_valve', n_clicks=0, size="lg",
                                                                   className='mr-1', color="primary", style={'margin': '-3rem 1rem 0 6vw'}),
                                                        ]),

                                                ],style = {'margin':'3rem'}, className='abcd'),
                                                   html.Div([html.Div([daq.Knob(id='HVA1OUT',color='red', label={"label": "HV A1 OUT", "style": {'color': 'red','marginBottom':'-1rem'}},
                                                                                value=0,size= 150,max=100, labelPosition = 'top'),
                                                                      daq.LEDDisplay(id = 'HVA1OUTLED',value = 0, color= 'blue', size=30,style = {'margin':'-2rem 0 2rem 0'})],
                                                                      className = 'aadbknob'),
                                                             html.Div([daq.Knob(id='HVA2OUT',color='red',  label={"label": "HV A2 OUT", "style": {'color': 'red','marginBottom':'-1rem'}},
                                                                                value=0,size=150, max=100, labelPosition='top'),
                                                                       daq.LEDDisplay(id='HVA2OUTLED',value = 0, color='blue', size=30, style={'margin': '-2rem 0 2rem 0'})],
                                                                      className='aadbknob'), ], className = 'abc'),]),
                                                  ], className='acreel'),

              html.Div([html.Div([ dcc.Graph(id="graphreal",
                                                     config={'displayModeBar': True,
                                                             'scrollZoom': True,
                                                             'displaylogo': False,
                                                             'modeBarButtonsToAdd': [
                                                                 'drawline',
                                                                 'drawrect',
                                                                 'drawopenpath',
                                                                 'select2d',
                                                                 'eraseshape',
                                                             ]},
                                                     style={'marginTop': '20px', },
                                                     figure={
                                                         'layout': {'legend': {'tracegroupgap': 0},

                                                                    }
                                                     }
                                                     ),
                                  daq.Slider(id="sliderHeightreel",
                                                      max=2100,
                                                      min=400,
                                                      value=530,
                                                      step=100,
                                                      size=400,
                                                      vertical=True,
                                                      updatemode='drag'), ],style={'marginTop': '3rem'},className='abcdb_graph2'),
                        html.Div([html.Div(daq.Slider(id="sliderWidthreel",
                                                      max=2000,
                                                      min=600,
                                                      value=1000,
                                                      step=100,
                                                      size=600,
                                                      updatemode='drag'), style={'marginLeft': '2rem'}),
                                  ]),
                        html.Div(dash_table.DataTable(id="getrealtable",
                                                      editable=True,
                                                      page_size=50,
                                                      style_table={'height': '500px', 'overflowY': 'auto',
                                                                   'width': '60vw', 'margin': 'auto'},
                                                      style_cell={
                                                          'overflow': 'hidden',
                                                          'textOverflow': 'ellipsis',
                                                          'maxWidth': 0,
                                                          'fontSize': '1rem',
                                                          'TextAlign': 'center',
                                                      },

                                                      fixed_rows={'headers': True},
                                                      style_header={
                                                          'backgroundColor': 'rgb(230, 230, 230)',
                                                          'fontWeight': 'bold'
                                                      },
                                                      filter_action="native",
                                                      sort_action="native",
                                                      sort_mode="multi",
                                                      column_selectable="single",
                                                      selected_columns=[],
                                                      selected_rows=[],
                                                      page_action="native",
                                                      page_current=0,
                                                      export_format='xlsx',
                                                      export_headers='display',
                                                      merge_duplicate_headers=True))
                        ],className='four-columns-div-user-controlsreel'),
              html.Div(id='reelhidden1', children=[], style={'display': 'None'}),
              html.Div(id='reelhidden2', children=[], style={'display': 'None'}),
              html.Div(id='reelhidden3', children=[], style={'display': 'None'}),
              html.Div(id='reelhidden4', children=[], style={'display': 'None'}),
              html.Div(id='reelhidden5', children=[], style={'display': 'None'}),
              dcc.Store(id='reelhidden6'),
              html.Div(id='reelhidden8', children=[], style={'display': 'None'}),
              html.Div(id='reelhidden9', children=[], style={'display': 'None'}),
              html.Div(id='reelhidden10', children=[], style={'display': 'None'}),],className='main_container',),
              ])

page_4_layout = html.Div([html.Div([html.Div([html.Div([html.Div([
                                                        html.Div([html.Div([html.I(className="fas fa-home"),], ),
                                                                  dcc.Link('Home', href='/', id='link1')], className='icon_position4'),
                                                        html.Div([html.Div([html.I(className="far fa-file-excel"),]),
                                                                  dcc.Link('Analysis', href='/File', id='link4')], className='icon_position4'),
                                                        html.Div([html.Div([html.I(className="fas fa-database"), ]),
                                                                  dcc.Link('Database', href='/Database', id='link2')], className='icon_position4'),
                                                        html.Div([html.Div([html.I(className="fas fa-server"), ]),
                                                                  dcc.Link('Enerbat', href='/realTime_Enerbat', id='link3')], className='icon_position4'),

                                                                                    ],
                                                        className = 'abcdbpr' ),
                                                        html.Div(dbc.Button("Who Are You?", id='enter_pr',
                                                                        n_clicks=0,color="danger",), style = {'float' : 'right', 'marginTop' : '-3rem'})]),

                                                        dbc.Modal(
                                                            [
                                                                dbc.ModalHeader("Enter your username and password"),
                                                                dbc.Input(id='username_pr',
                                                                          type="text",
                                                                          min=-10000, max=10000, step=1, bs_size="sm",
                                                                          style={'width': '31rem', },
                                                                          placeholder = 'Username',
                                                                          autoFocus=True, ),
                                                                dbc.Input(id='password_pr',
                                                                          type="password",
                                                                          min=-10000, max=10000, step=1, bs_size="sm",
                                                                          style={'width': '31rem', },
                                                                          placeholder = 'Password',
                                                                          autoFocus=True, ),
                                                                dbc.ModalFooter([
                                                                    dbc.Button("Okey", id="ok_reel_pr", className='ml-auto'),
                                                                    dbc.Button("Close", id="close_reel_pr", className='ml-auto')]
                                                                ),
                                                            ],
                                                            id="modal_reel_pr",
                                                        ),

                                    html.Div([html.Div([html.Div([daq.PowerButton(id='my-toggle-switch-pr-db',
                                                                                  label={'label': 'Connection Database',
                                                                                         'style': {'fontSize': '22px',
                                                                                                   'fontWeight': 'bold'}},
                                                                                  labelPosition='bottom', on=False, size=100,
                                                                                  color="green",
                                                                                  className='dark-theme-control'),

                                                                                ], style={'marginLeft': '1rem'}),
                                                        html.Div([ dcc.Dropdown(id='prname',
                                                                                options=[{'label': i, 'value': i}
                                                                                         for i in ['rcckn', 'enerbat']],
                                                                                multi=False,
                                                                                style={'cursor': 'pointer', 'marginTop': '5px'},
                                                                                className='stockSelectorClass3',
                                                                                clearable=True,
                                                                                placeholder='Select Database',
                                                                                ),
                                                                   dcc.Dropdown(id='prvalchoosen',
                                                                                # options=[{'label': i, 'value': i}
                                                                                #          for i in df.columns],
                                                                                multi=False,
                                                                                style={'cursor': 'pointer', 'marginTop': '5px'},
                                                                                className='stockSelectorClass3',
                                                                                clearable=True,
                                                                                placeholder='Select Table...',
                                                                                ),
                                                                   dcc.Dropdown(id='prvalname',
                                                                                # options=[{'label': i, 'value': i}
                                                                                #          for i in df.columns],
                                                                                multi=True,
                                                                                style={'cursor': 'pointer','marginTop': '13px'},
                                                                                className='stockSelectorClass3',
                                                                                clearable=True,
                                                                                placeholder='Select your parameters...',
                                                                                ),
                                                                   dcc.Dropdown(id='prvaldate',
                                                                                # options=[{'label': i, 'value': i}
                                                                                #          for i in df.columns],
                                                                                multi=True,
                                                                                style={'cursor': 'pointer', 'marginTop': '13px'},
                                                                                className='stockSelectorClass3',
                                                                                clearable=False,
                                                                                placeholder='Select your time...',
                                                                                ),
                                                                   html.P('Enter interval value (Second)'),
                                                                    dcc.Interval(
                                                                                 id='interval_component_pr_db',
                                                                                 disabled=True,
                                                                                 interval=1 * 1000,  # in milliseconds
                                                                                 n_intervals=0),
                                                                   dbc.Input(id='interval_value_pr_db', type="text", value='5',
                                                                             min=0, max=1000000000, step=1, bs_size="lg", style={'width': '6rem'}, ),
                                                                               ], className='page4reel'),

                                                                ],className='abcdbpage4upleft'),

                                            html.Div([html.Div([
                                                                    html.Div([daq.PowerButton(id='my-toggle-switch-pr',
                                                                                     label={'label': "Conn. Enerbat",
                                                                                            'style': {'fontSize': '22px',
                                                                                                      'fontWeight': 'bold'}},
                                                                                     labelPosition='bottom', on=False, size=100,
                                                                                     color="green",
                                                                                     className='dark-theme-control'),
                                                                            dbc.Tooltip(
                                                                                    "!!! Before the connection, fulfill database and table name, "
                                                                                    "Then send to Database button",
                                                                                    target = "my-toggle-switch-pr",
                                                                                            ),
                                                                     ], className='abpower',style = {'margin':'0 1rem 0 5rem'}),
                                                                           dcc.Store(id='get_data_from_modbus_pr'),
                                                                           html.Div(id='data_to_store_id_pr', children=[], style={'display': 'None'}),
                                                                           html.Div(id='data_to_store_value_pr', children=[],
                                                                                    style={'display': 'None'}),
                                                                           html.Div(id='data_to_store_qualite_pr', children=[],
                                                                                    style={'display': 'None'}),
                                                                           html.Div(id='data_to_store_date_pr', children=[],
                                                                                    style={'display': 'None'}),
                                                                           dcc.Interval(
                                                                               id='interval_component_pr',
                                                                               disabled=True,
                                                                               interval=1 * 1000,  # in milliseconds
                                                                               n_intervals=0), ],className='page4reel'),
                                                            html.Div([
                                                                 dcc.Dropdown(id='realvalue_pr',
                                                                           options = [{'label': i[16:], 'value': i} for i in ['sauter.EY6AS680.Tb1', 'sauter.EY6AS680.Tb2',
                                                                                                        'sauter.EY6AS680.Tb3', 'sauter.EY6AS680.Tb4',
                                                                                                        'sauter.EY6AS680.Tec', 'sauter.EY6AS680.Teev',
                                                                                                        'sauter.EY6AS680.Teg', 'sauter.EY6AS680.Tsc',
                                                                                                        'sauter.EY6AS680.Tsev', 'sauter.EY6AS680.Tsg' ]],
                                                                           multi=True,
                                                                           style={'cursor': 'pointer', 'margin': '5px 5px 10px 0',
                                                                                  },
                                                                           className='stockSelectorClass3',
                                                                           clearable=True,
                                                                           placeholder='Select Value',

                                                                           ),
                                                              html.P('Enter interval value (Second)'),
                                                              dbc.Input(id='interval_value_pr', type="text", value='5',
                                                                        min=0, max=1000000000, step=1, bs_size="lg", style={'width': '6rem'}, ),

                                                              html.Div(
                                                                        html.Div(
                                                                            [
                                                                                dbc.Button("Send to Database", id='download_pr',
                                                                                           n_clicks=0, size="lg",
                                                                                           className='mr-1', color="primary",
                                                                                           style={'margin': '1rem 1rem 1rem 0'}),
                                                                                dbc.Tooltip(
                                                                                    "!!! Enter a Database and Table name",
                                                                                    "If entered, disregard this message ",
                                                                                    target = "download_pr",
                                                                                            ),
                                                                                ]),
                                                                        className='abcd'),
                                                              html.Div([html.Div([html.P('Enter Table Name'),
                                                                                dbc.Input(id='filenametodb', type="text", value='',
                                                                                min=0, max=1000000000, step=1, bs_size="lg", style={'width': '10rem'}, ),]),
                                                                        html.Div([html.P('Enter Database Name', style={'marginLeft': '1rem', }),
                                                                                dbc.Input(id='nametodb', type="text", value='',
                                                                                min=0, max=1000000000, step=1, bs_size="lg",
                                                                                style={'width': '10rem', 'marginLeft': '1rem', }, ),])],className='abcd',),
                                                          ], className='page4reel', ),

                                  html.Div(id='reelhidden1pr', children=[], style={'display': 'None'}),
                                  html.Div(id='reelhidden2pr', children=[], style={'display': 'None'}),
                                  html.Div(id='reelhidden3pr', children=[], style={'display': 'None'}),
                                  html.Div(id='reelhidden4pr', children=[], style={'display': 'None'}),
                                  ],className='abcdbpage4upleft')
                                ],className='prstyle')
                                         ], className='page4reel'),
                               ], className='abcdbpage4'),
     dcc.Store(id='memory-outputpr'),
     html.Div([html.Div([html.Div([html.Div(dcc.Dropdown(id='firstgraph_pr_real',
                                            options=[{'label': i, 'value': i} for i in []],
                                            multi=True,
                                            style={'cursor': 'pointer', 'width': '30rem', 'margin' : '1rem 0 0 5rem'},
                                            clearable=True,
                                            placeholder='Values of Real Time',
                                                            ),),
                                html.Div(dcc.Dropdown(id='firstgraph_pr_db',
                                                options=[{'label': i, 'value': i} for i in []],
                                                multi=True,
                                                style={'cursor': 'pointer', 'width': '30rem', 'margin' : '1rem 0 0 5rem'},
                                                clearable=True,
                                                placeholder='Values of Database',
                                                            ),),
                       html.Div([html.Div([html.Div(dcc.Graph(id="getprgraph",
                                            config={'displayModeBar': True,
                                                    'scrollZoom': True,
                                                    'displaylogo': False,
                                                    'modeBarButtonsToAdd': [
                                                        'drawline',
                                                        'drawrect',
                                                        'drawopenpath',
                                                        'select2d',
                                                        'eraseshape',
                                                    ]},
                                            style={'marginTop': '20px', },
                                            figure={
                                                'layout': {'legend': {'tracegroupgap': 0},}
                                            }
                                            ), ),
                         html.Div(daq.Slider(id="sliderHeightpr",
                                             max=1200,
                                             min=200,
                                             value=400,
                                             step=100,
                                             size=300,
                                             vertical=True,
                                             updatemode='drag'), style={'margin': '2rem 2rem 0 1rem'})], style= {'margin' : '1rem 1rem 1rem 5rem'}, className='page4graph1'),
                        html.Div(daq.Slider(id="sliderWidthpr",
                                            max=1600,
                                            min=500,
                                            value=800,
                                            step=100,
                                            size=500,
                                                     updatemode='drag'), style={'marginLeft': '7rem'}),
                        ],style = {'margin':'1rem','padding':'10px'}, className='boxdesign1'),
                         ]),
                       html.Div(dcc.Dropdown(id='secondgraph_pr_real',
                                    options=[{'label': i, 'value': i} for i in []],
                                    multi=True,
                                    style={'cursor': 'pointer', 'width': '30rem','margin' : '3rem 0 0 5rem'},
                                    clearable=True,
                                    placeholder='Values of Real Time',
                               ),),
                               html.Div(dcc.Dropdown(id='secondgraph_pr_db',
                                    options=[{'label': i, 'value': i} for i in []],
                                    multi=True,
                                    style={'cursor': 'pointer', 'width': '30rem','margin' : '1rem 0 0 5rem'},
                                    clearable=True,
                                    placeholder='Values of Database',
                               ),),
                               html.Div([html.Div([html.Div(dcc.Graph(id="getprgraph2",
                                                    config={'displayModeBar': True,
                                                            'displaylogo': False,
                                                            'scrollZoom': True,
                                                            'modeBarButtonsToAdd': [
                                                                'drawline',
                                                                'drawrect',
                                                                'drawopenpath',
                                                                'select2d',
                                                                'eraseshape',
                                                            ]},
                                                    style={'marginTop': '20px' },
                                                    figure={
                                                        'layout': {'legend': {'tracegroupgap': 0},}
                                                    }
                                                    ), ),
                                                 html.Div(daq.Slider(id="sliderHeightpr2",
                                                                     max=1200,
                                                                     min=200,
                                                                     value=400,
                                                                     step=100,
                                                                     size=300,
                                                                     vertical=True,
                                                                     updatemode='drag'), style={'margin': '2rem 2rem 0 1rem'})],style = {'margin' : '1rem 1rem 1rem 5rem'}, className='page4graph1'),
                                                html.Div(daq.Slider(id="sliderWidthpr2",
                                                                    max=1600,
                                                                    min=500,
                                                                    value=800,
                                                                    step=100,
                                                                    size=500,
                                                                             updatemode='drag'), style={'marginLeft': '6rem'}),
                                                ],style = {'margin':'1rem','padding':'10px'}, className='boxdesign2'),
                                                 ]),

                html.Div([html.Div([html.Div([
                                html.Div([dcc.Dropdown(id='thirdgraph_pr_real',
                                                options=[{'label': i, 'value': i} for i in []],
                                                multi=True,
                                                style={'cursor': 'pointer', 'width': '30rem', 'margin' : '1rem 0 0 5rem'},
                                                clearable=True,
                                                placeholder='Values of Real Time',
                                                    ),
                                        dcc.Dropdown(id='thirdgraph_pr_db',
                                                options=[{'label': i, 'value': i} for i in []],
                                                multi=True,
                                                style={'cursor': 'pointer', 'width': '30rem', 'margin' : '1rem 0 0 5rem' },
                                                clearable=True,
                                                placeholder='Values of Database',
                                                )    ,], className = 'thirdgraphpr_db'),
                                       html.Div([html.Div([html.Div(dcc.Graph(id="getprgraph3",
                                                            config={'displayModeBar': True,
                                                                    'scrollZoom': True,
                                                                     'displaylogo': False,
                                                                    'modeBarButtonsToAdd': [
                                                                        'drawline',
                                                                        'drawrect',
                                                                        'drawopenpath',
                                                                        'select2d',
                                                                        'eraseshape',
                                                                    ]},
                                                            style={'marginTop': '20px', },
                                                            figure={
                                                                'layout': {'legend': {'tracegroupgap': 0},}
                                                            }
                                                            ), ),
                                         html.Div(daq.Slider(id="sliderHeightpr3",
                                                             max=1200,
                                                             min=200,
                                                             value=400,
                                                             step=100,
                                                             size=300,
                                                             vertical=True,
                                                             updatemode='drag'), style={'margin': '2rem 2rem 0 1rem'})],style = {'margin':'1rem 3rem 1rem 1rem'}, className='page4graph1'),
                                        html.Div(daq.Slider(id="sliderWidthpr3",
                                                            max=1600,
                                                            min=500,
                                                            value=800,
                                                            step=100,
                                                            size=500,
                                                                     updatemode='drag'), style={'marginLeft': '5rem'}),
                                        ],style = {'margin':'1rem','padding':'10px'}, className='boxdesign3'),
                                         ]),

                                       html.Div([html.Div(dcc.Dropdown(id='fourgraph_pr_real',
                                                                    options=[{'label': i, 'value': i} for i in []],
                                                                    multi=True,
                                                                    style={'cursor': 'pointer', 'width': '30rem', 'margin' : '3rem 0 0 5rem'},
                                                                    clearable=True,
                                                                    placeholder='Values of Real Time',
                                                                            ),),
                                               html.Div(dcc.Dropdown(id='fourgraph_pr_db',
                                                                options=[{'label': i, 'value': i} for i in []],
                                                                multi=True,
                                                                style={'cursor': 'pointer', 'width': '30rem', 'margin' : '1rem 0 0 5rem'},
                                                                clearable=True,
                                                                placeholder='Values of Database',
                                                                ),),], className = 'fourgraphpr_db'),
                                               html.Div([html.Div([html.Div(dcc.Graph(id="getprgraph4",
                                                                    config={'displayModeBar': True,
                                                                            'scrollZoom': True,
                                                                            'displaylogo': False,
                                                                            'modeBarButtonsToAdd': [
                                                                                'drawline',
                                                                                'drawrect',
                                                                                'drawopenpath',
                                                                                'select2d',
                                                                                'eraseshape',
                                                                            ]},
                                                                    style={'margin': '20px ' },
                                                                    figure={
                                                                        'layout': {'legend': {'tracegroupgap': 0},}
                                                                    }
                                                                    ), ),
                                                                 html.Div(daq.Slider(id="sliderHeightpr4",
                                                                                     max=1200,
                                                                                     min=200,
                                                                                     value=400,
                                                                                     step=100,
                                                                                     size=300,
                                                                                     vertical=True,
                                                                                     updatemode='drag'), style={'margin': '2rem 2rem 0 1rem'})],style = {'margin':'0 3rem 0 0'}, className='page4graph1'),
                                                                html.Div(daq.Slider(id="sliderWidthpr4",
                                                                                             max=1600,
                                                                                             min=500,
                                                                                             value=800,
                                                                                             step=100,
                                                                                             size=500,
                                                                                             updatemode='drag'), style={'marginLeft': '5rem'}),
                                                                ],style = {'margin':'1rem','padding':'10px'}, className='boxdesign4'),
                                                                 ], ), ]),], className='abcdbprGraph'),

               html.Div(dash_table.DataTable(id="getprtable",
                                             editable=True,
                                             page_size=50,
                                             style_table={'height': '500px', 'overflowY': 'auto', 'width': '98%'},
                                             style_cell={
                                                 'overflow': 'hidden',
                                                 'textOverflow': 'ellipsis',
                                                 'maxWidth': 0,
                                                 'fontSize': '1rem',
                                                 'TextAlign': 'center',
                                             },

                                             fixed_rows={'headers': True},
                                             style_header={
                                                 'backgroundColor': 'rgb(230, 230, 230)',
                                                 'fontWeight': 'bold' },
                                             filter_action="native",
                                             sort_action="native",
                                             sort_mode="multi",
                                             column_selectable="single",
                                             selected_columns=[],
                                             selected_rows=[],
                                             page_action="native",
                                             page_current=0,
                                             export_format='xlsx',
                                             export_headers='display',
                                             merge_duplicate_headers=True)),
               ], style = {'overflow-x' : 'auto'} ),



@app.callback([ Output("HVA1INLED", "value"),Output("HVA2INLED", "value"),
                Output("HVA1OUTLED", "value"),Output("HVA2OUTLED", "value")],
              [ Input("HVA1IN", "value"), Input("HVA2IN", "value"),
                Input("HVA1OUT", "value"), Input("HVA2OUT", "value")],)
def knobvalues(v1,v2,v3,v4):
    if v1 != None or v2 != None or v3 != None or v4 != None:
        v1,v2,v3,v4 = int((v1/100)*28000),int((v2/100)*28000),int((v3/100)*28000),int((v4/100)*28000)
        return v1,v2,v3,v4
    else: raise PreventUpdate

@app.callback(  Output("reelhidden6", "children"),
              [ Input("download_reel_valve", "n_clicks")],
              [ State("HVA1INLED", "value"), State("HVA2INLED", "value"),
                State("HVA1OUTLED", "value"), State("HVA2OUTLED", "value")])
def toserver(nc, v1,v2,v3,v4):
    if nc > 0 :
        opc = OpenOPC.client()
        opc.servers()
        opc.connect('Kepware.KEPServerEX.V6')
        opc.write([('Siemens.PLC1.Vanne3voies1', v1), ('Siemens.PLC1.Vanne3voies2', v2),('Siemens.PLC1.Vanne3voies3', v3), ('Siemens.PLC1.Vanne3voies4', v4)])

@app.callback(  [Output("reelhidden3", "children"),Output("reelhidden5", "children")],
                [Input("ok_reel", "n_clicks"), ],
                [State("input_tablename", "value"),State("input_databasename", "value")],)
def toggle_modal_1(nc, tbname, databasename):
    if tbname == None  or databasename==None:
        raise PreventUpdate
    if nc != None:
        return tbname,databasename

@app.callback(
    Output("modal_reel", "is_open"),
    [Input("download_reel_db", "n_clicks"), Input("close_reel", "n_clicks"), Input("ok_reel", "n_clicks")],
    [State("modal_reel", "is_open")],
)
def toggle_modal_2(n1, n2, n3, is_open):
    if n1 or n2 or n3:
        return not is_open
    return is_open

@app.callback(
    Output("modal_reel_analysis", "is_open"),
    [Input("save_to_db", "n_clicks"), Input("close_reel_analysis", "n_clicks"), Input("ok_reel_analysis", "n_clicks")],
    [State("modal_reel_analysis", "is_open")],
)
def toggle_modal_analysis(n1, n2, n3, is_open):
    if n1 or n2 or n3:
        return not is_open
    return is_open

@app.callback(  [Output("reelhidden11", "children"),Output("reelhidden12", "children"),
                 Output("reelhidden14", "children"),Output("reelhidden15", "children")],
                [Input("ok_reel_analysis", "n_clicks"), ],
                [State("input_tblname", "value"),State("input_dbname", "value"),
                 State("input_user", "value"),State("input_pswrd", "value")],)
def toggle_modal_analaysis(nc, tbname, databasename,user, pswrd):
    if tbname == None  or databasename==None:
        raise PreventUpdate
    if nc != None:
        return tbname,databasename, user, pswrd

@app.callback(
    [Output("reelhidden3pr", "children"),Output("reelhidden4pr", "children")],
    [Input('download_pr', 'n_clicks')],
    [State("filenametodb", "value"),
     State("nametodb", "value")],
)
def toggle_modal_3(n1, name, nametodb):
    if n1 > 0 :
        if name != '' and nametodb == '' :
            return name , 'enerbat'
        elif name == '' and nametodb != '' :
            return 'LERMAB_test' , nametodb
        elif name != '' and nametodb != '' :
            return name , nametodb
        else :
            return 'LERMAB_test', 'enerbat'
    else : return no_update, no_update

@app.callback(Output('interval_component', 'disabled'),
              [Input("my-toggle-switch-reel", "on")],
              )
def intervalcontrol(on):
    if on == 1:
        return False
    else:
        return True

@app.callback([Output('activatedb', 'color'),Output('deactivatedb', 'color')],
              [Input('activatedb', 'n_clicks'),Input('deactivatedb', 'n_clicks')],
              [State('activatedb', 'color'),State('deactivatedb', 'color')],
              )
def buttoncontrol(on1, on2, col1, col2):
    q1 = dash.callback_context.triggered[0]["prop_id"].split(".")[0]
    if q1 == 'activatedb':
        return 'secondary', 'danger'
    if q1 == 'deactivatedb':
        return 'success','secondary'
    else : return no_update, no_update

@app.callback(Output('interval_component_pr', 'disabled'),
              [Input("my-toggle-switch-pr", "on")],
              )
def intervalcontrolpr(on):
    if on == 1:
        return False
    else:
        return True

@app.callback(Output('interval_component_pr_db', 'disabled'),
              [Input("my-toggle-switch-pr-db", "on")],
              )
def intervalcontrolpr_db(on):
    if on == 1:
        return False
    else:
        return True

@app.callback(Output('interval_component', 'interval'),
              [Input("interval_value", "value")],
              )
def intervalcontrol2(val):
    val = int(val)*1000
    return val

@app.callback(Output('interval_component_pr', 'interval'),
              [Input("interval_value_pr", "value")],
              )
def intervalcontrol2_pr(val):
    val = int(val)*1000
    return val

@app.callback(Output('interval_component_pr_db', 'interval'),
              [Input("interval_value_pr_db", "value")],
              )
def intervalcontrol2_pr_db(val):
    val = int(val)*1000
    return val

@app.callback(Output("interval_value_pr_db", "value"),
              Input('prname', 'value')
              )
def intervalcontrol2_pr_db(val):
    if val == 'rcckn' :
        return '10'
    return '5'


@app.callback([Output('data_to_store_id', 'children'),
               Output('data_to_store_value', 'children'),
               Output('data_to_store_qualite', 'children'),
               Output('data_to_store_date', 'children'), ],
              [Input("my-toggle-switch-reel", "on"), Input('interval_component', 'n_intervals')],
              [State('data_to_store_id', 'children'),
               State('data_to_store_value', 'children'),
               State('data_to_store_qualite', 'children'),
               State('data_to_store_date', 'children'), ])
def values(on, n_intervals, id, val, qual, date):
    if on == 1:
        opc = OpenOPC.client()
        opc.servers()
        opc.connect('Kepware.KEPServerEX.V6')
        for ID, value, Quality, Timestamp in opc.iread(
                ['sauter.EY6AS680.Tb1', 'sauter.EY6AS680.Tb2', 'sauter.EY6AS680.Tb3', 'sauter.EY6AS680.Tb4',
                 'sauter.EY6AS680.Tec', 'sauter.EY6AS680.Teev', 'sauter.EY6AS680.Teg', 'sauter.EY6AS680.Tsc',
                 'sauter.EY6AS680.Tsev', 'sauter.EY6AS680.Tsg' ]):
            id.append(ID[16:])
            val.append(value)
            qual.append(Quality)
            date.append(Timestamp)

    return id, val, qual, date

@app.callback([Output('data_to_store_id_pr', 'children'),
               Output('data_to_store_value_pr', 'children'),
               Output('data_to_store_qualite_pr', 'children'),
               Output('data_to_store_date_pr', 'children'),
               ],
              [Input("my-toggle-switch-pr", "on"),Input('realvalue_pr', 'value'), Input('interval_component_pr', 'n_intervals')],
              [State('data_to_store_id_pr', 'children'),
               State('data_to_store_value_pr', 'children'),
               State('data_to_store_qualite_pr', 'children'),
               State('data_to_store_date_pr', 'children'), ])
def values_pr(on,realval, n_intervals, id, val, qual, date):
    if realval == None :
        raise PreventUpdate
    if on == 1:
        opc = OpenOPC.client()
        opc.servers()
        opc.connect('Kepware.KEPServerEX.V6')
        for ID, value, Quality, Timestamp in opc.iread(realval):
            id.append(ID[16:])
            val.append(value)
            qual.append(Quality)
            date.append(Timestamp)
    return id, val, qual, date

@app.callback([Output('get_data_from_modbus', 'data'), Output('realvalue', 'options')],
              [Input('data_to_store_id', 'children'),
               Input('data_to_store_value', 'children'),
               Input('data_to_store_qualite', 'children'),
               Input('data_to_store_date', 'children'), ], )
def storedata(id, val, qual, date):
    zipped_val = list(zip(id, val, qual, date))
    df = pd.DataFrame(list(zip(id, val, qual, date)),
                      columns=['ID', 'Value', 'Quality', 'TIMESTAMP'])
    val = df['ID'].unique()
    return zipped_val, [{'label': i, 'value': i} for i in val]

@app.callback( Output('get_data_from_modbus_pr', 'data'),
              [Input('data_to_store_id_pr', 'children'),
               Input('data_to_store_value_pr', 'children'),
               Input('data_to_store_qualite_pr', 'children'),
               Input('data_to_store_date_pr', 'children'), ], )
def storedata_pr(id, val, qual, date):
    zipped_val = list(zip(id, val, qual, date))
    return zipped_val

@app.callback(Output('reelhidden1pr', 'children'),
              [Input("write_excel_pr", "n_clicks")],
              [State('get_data_from_modbus_pr', 'data')],
              )
def intervalcontrol2_pr(nc, data):
    if nc > 0:
        df = pd.DataFrame(data, columns=['ID', 'Value', 'Quality', 'date'])
        df.to_excel('real.xlsx')


@app.server.route("/download_excel_pr/")
def download_excel_pr():
    # Create DF
    dff = pd.read_excel("real.xlsx")
    # Convert DF
    buf = io.BytesIO()
    excel_writer = pd.ExcelWriter(buf, engine="xlsxwriter")
    dff.to_excel(excel_writer, sheet_name="sheet1")
    excel_writer.save()
    excel_data = buf.getvalue()
    buf.seek(0)
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        attachment_filename="real.xlsx",
        as_attachment=True,
        cache_timeout=0
    )

@app.callback(Output('reelhidden2', 'children'),
              [Input("reelhidden3", "children"), Input("reelhidden5", "children"),],
              [State('get_data_from_modbus', 'data'),State('username_pr', 'value'), State('password_pr', 'value')])
def pandastosql(name,dbname, data, user, passw):
    if name == None or dbname == None:
        raise PreventUpdate
    df = pd.DataFrame(data)
    if name != None:
        df = pd.DataFrame(data, columns=['variable_name', 'variable_num_value', 'quality', 'TIMESTAMP'])
        a = [i for i in range(len(df.index))]  # for ID
        b = [i for i in df['variable_name']]  # name of variable
        c = [i for i in df['variable_num_value']]
        d = [i for i in df['TIMESTAMP']]
        df['TIMESTAMP'] = df['TIMESTAMP'].apply(lambda x: pd.Timestamp(x).strftime('%Y-%m-%d %H:%M:%S'))
        sql_insert = list(zip(a, df['variable_name'], df['variable_num_value'], df['TIMESTAMP']))
        if dbname == None:
            dbname = 'enerbat'
        ipadress = "193.54.2.211"
        server = SSHTunnelForwarder(
            (ipadress, 22),
            ssh_username='soudani',
            ssh_password="univ484067152",
            remote_bind_address=(ipadress, 3306))
        server.start()
        try:
            db_connection = mariadb.connect(
                user='dashapp' if user == None else user,
                password='dashapp' if passw == None else passw,
                host=ipadress,
                port=3306,
                database=dbname)
            db_cursor = db_connection.cursor()
            # Here creating database table '
            db_cursor.execute(
                f"CREATE OR REPLACE TABLE {name} (id BIGINT PRIMARY KEY, variable_name VARCHAR(255), variable_num_value DOUBLE, TIMESTAMP TIMESTAMP)")
            sql_query = f" INSERT INTO {name} (id, variable_name,variable_num_value,TIMESTAMP) VALUES (%s, %s, %s, %s)"
            # Get database table'
            db_cursor.executemany(sql_query, sql_insert)
            db_connection.commit()
            print(db_cursor.rowcount, "Record inserted successfully into ENERBAT Database")
        except mariadb.Error as e:
            print(f"Error connecting to MariaDB Platform: {e}")
            sys.exit(1)

@app.callback(Output('reelhidden13', 'children'),
              [Input("reelhidden11", "children"), Input("reelhidden12", "children"),
               Input("reelhidden14", "children"),Input("reelhidden15", "children")],)
def pandastosql_analysis(name,dbname, user, passw):
    if name == None or dbname == None:
        raise PreventUpdate
    print(name)
    print(dbname)
    print(user)
    print(passw)
    df = pd.read_excel('new_fichier.xlsx',index_col=[0])
    df.columns = df.iloc[0]
    df = df[1:]

    if name != None:
        #df = pd.DataFrame(data, columns=['variable_name', 'variable_num_value', 'quality', 'TIMESTAMP'])
        a = [i for i in range(0,len(df.index))]  # for ID
        index = a
        sql_insert = list(zip(df['firstChoosenValue'],df['leftIntegralFirst'],df['leftIntegralSecond'], df['leftIntegral'],
                              df['secondChoosenValue'],df['rightIntegralFirst'],df['rightIntegralSecond'], df['rightIntegral'],
                              df['operation'],df['intersection'],df['referance X'], df['referance Y'],df['time']))
        print(sql_insert)
        df.rename({'index':'0'},axis='columns',inplace=True)
        print(df)
        if dbname == None:
            dbname = 'enerbat'
        ipadress = "193.54.2.211"
        server = SSHTunnelForwarder(
            (ipadress, 22),
            ssh_username='soudani',
            ssh_password="univ484067152",
            remote_bind_address=(ipadress, 3306))
        server.start()
        try:
            db_connection = mariadb.connect(
                            user='dashapp' if user == None else user,
                            password='dashapp' if passw == None else passw,
                            host=ipadress,
                            port=3306,
                            database=dbname)
            db_cursor = db_connection.cursor()
            # Here creating database table '
            db_cursor.execute(
                f"CREATE OR REPLACE TABLE {name} (id BIGINT PRIMARY KEY ,First_Choosen_Value VARCHAR(255), First_Left_Point VARCHAR(255), Second_Left_Point VARCHAR(255),\
                Left_Integral FLOAT(10), Second_Choosen_Value VARCHAR(255), First_Right_Point VARCHAR(255), Second_Right_Point VARCHAR(255),\
                Right_Integral FLOAT(10), Math_Ops FLOAT(10), Intersection FLOAT(10),Referance_X BIGINT,Referance_Y BIGINT, Time TIMESTAMP)")
            print('burda mi')
            sql_query = f" INSERT INTO {name} (id, First_Choosen_Value, First_Left_Point, Second_Left_Point, Left_Integral, Second_Choosen_Value, First_Right_Point,\
                        Second_Right_Point, Right_Integral, Math_Ops, Intersection,Referance_X,Referance_Y, Time) VALUES (%s, %s, %s, %s, %s, %s, %s, %s,%s, %s,%s, %s, %s, %s);"
            # Get database table'
            sql_insert = tuple(sql_insert)
            db_cursor.execute(sql_query, sql_insert)
            db_connection.commit()
            print(db_cursor.rowcount, "Record inserted successfully into ENERBAT Database")
        except mariadb.Error as e:
            print(f"Error connecting to MariaDB Platform: {e}")
            sys.exit(1)


@app.callback(Output('reelhidden10', 'children'),
              [Input("reelhidden8", "children"), Input("reelhidden9", "children"),],
              [State("reelhidden6", "children"),])
def pandastosql_valve(name,dbname, data, ):
    if data == None:
        raise PreventUpdate
    df = pd.DataFrame(data)
    df.columns=['HV_A1_IN', 'HV_A2_IN', 'HV_A1_OUT', 'HV_A2_OUT']
    HV_A1_IN = [i for i in df['HV_A1_IN']]  # name of variable
    HV_A2_IN = [i for i in df['HV_A2_IN']]
    HV_A1_OUT = [i for i in df['HV_A1_OUT']]
    HV_A2_OUT = [i for i in df['HV_A2_OUT']]
    sql_insert = list(zip(HV_A1_IN, HV_A2_IN, HV_A1_OUT,  HV_A2_OUT))
    ipadress = "193.54.2.211"
    server = SSHTunnelForwarder(
        (ipadress, 22),
        ssh_username='soudani',
        ssh_password="univ484067152",
        remote_bind_address=(ipadress, 3306))

    server.start()
    try:
        db_connection =  mariadb.connect(
                            user='dashapp',
                            password='dashapp',
                            host=ipadress,
                            port=3306,
                            database=dbname)
        db_cursor = db_connection.cursor()
            # Here creating database table '
        db_cursor.execute(
            f"CREATE OR REPLACE TABLE {name} (HV_A1_IN BIGINT PRIMARY KEY, HV_A2_IN BIGINT, HV_A1_OUT BIGINT, HV_A2_OUT BIGINT)")
        sql_query = f" INSERT INTO {name} (HV_A1_IN,HV_A2_IN,HV_A1_OUT,HV_A2_OUT) VALUES (%s, %s, %s, %s)"
            # Get database table'
        db_cursor.executemany(sql_query, sql_insert)
        db_connection.commit()
        print(db_cursor.rowcount, f"Record inserted successfully into {name} Database")
    except mariadb.Error as e:
        print(f"Error connecting to MariaDB Platform: {e}")
        sys.exit(1)

@app.callback(Output('reelhidden2pr', 'children'),
              [Input("my-toggle-switch-pr", "on"),Input('interval_component_pr', 'n_intervals')],
              [State("reelhidden3pr", "children"),State("reelhidden4pr", "children"), State('get_data_from_modbus_pr', 'data'),
               State('username_pr', 'value'), State('password_pr', 'value')])
def pandastosql_pr(on,interval, name,nametodb, data, user, passw):
    if name == None :
        raise PreventUpdate
    if on == 1:
        if name != None:
            df = pd.DataFrame(data, columns=['variable_name', 'variable_num_value', 'quality', 'TIMESTAMP'])
            a = [i for i in range(len(df.index))]  # for ID
            b = [i for i in df['variable_name']]  # name of variable
            c = [i for i in df['variable_num_value']]
            d = [i for i in df['TIMESTAMP']]
            df['TIMESTAMP'] = df['TIMESTAMP'].apply(lambda x: pd.Timestamp(x).strftime('%Y-%m-%d %H:%M:%S'))
            sql_insert = list(zip(a, df['variable_name'], df['variable_num_value'], df['TIMESTAMP']))
            ipadress = "193.54.2.211"
            server = SSHTunnelForwarder(
                (ipadress, 22),
                ssh_username='soudani',
                ssh_password="univ484067152",
                remote_bind_address=(ipadress, 3306))

            server.start()
            try:
                db_connection = mariadb.connect(
                            user='dashapp' if user == None else user,
                            password='dashapp' if passw == None else passw,
                            host=ipadress,
                            port=3306,
                            database=nametodb)
                db_cursor = db_connection.cursor()

                # db_cursor.execute(f"REPAIR TABLE {name}")
                db_cursor.execute(f"CREATE OR REPLACE TABLE  {name} (ID BIGINT PRIMARY KEY, VARIABLE_NAME VARCHAR(255), VARIABLE_NUM_VALUE DOUBLE, TIMESTAMP TIMESTAMP)")
                sql_query = f" INSERT INTO {name} (id, variable_name, variable_num_value,TIMESTAMP) VALUES (%s, %s, %s, %s)"
                # Get database table'
                db_cursor.executemany(sql_query, sql_insert)
                db_connection.commit()
                print(db_cursor.rowcount, f"Record inserted successfully into {nametodb} Database")
            except mariadb.Error as e:
                print(f"Error connecting to MariaDB Platform: {e}")
                sys.exit(1)


# surf between pages
# Update the index
@app.callback(dash.dependencies.Output('page-content', 'children'),
              [dash.dependencies.Input('url', 'pathname')])
def display_page(pathname):
    if pathname == '/':
        return index_page
    elif pathname == '/Database':
        return page_2_layout
    elif pathname == '/realTime_Enerbat':
        return page_3_layout
    elif pathname == '/RCC_KN':

        return page_4_layout
    elif pathname == '/File':
        return page_1_layout

@app.callback(
            Output("modal_reel_pr", "is_open"),
            [Input("close_reel_pr", "n_clicks"), Input("ok_reel_pr", "n_clicks"),Input('enter_pr', 'n_clicks')],
            [State("modal_reel_pr", "is_open")],
        )
def toggle_modal_pr(n2, n3, nc, is_open):
    if nc > 0 or n2 or n3:
        return not is_open
    return is_open


def parse_contents(contents, filename, date):
    content_type, content_string = contents.split(',')
    decoded = base64.b64decode(content_string)

    try:
        if 'csv' in filename:
            # Assume that the user uploaded a CSV file
            df = pd.read_csv(io.StringIO(decoded.decode('utf-8')))
            df['index'] = df.index
            df = df.reindex(columns=sorted(df.columns, reverse=True))
            df.to_excel("appending.xlsx")
        elif 'xls' in filename:
            df = pd.read_excel(io.BytesIO(decoded))
            df = df.filter(regex='^\D') ### very important for filter the excel columns which is starts with number
            df['index'] = df.index
            df = df.reindex(columns=sorted(df.columns, reverse=True))
            df.to_excel("appending.xlsx")

    except Exception as e:
        print(e)
        return html.Div([
            'There was an error processing this file.'
        ])
    return html.Div([
        html.H5(filename, style = {'color' : 'black'}),
        html.H6(datetime.datetime.fromtimestamp(date), style = {'color' : 'black'}),
        dash_table.DataTable(
            id='datatable-interactivity',
            data=df.to_dict('records'),
            columns=[{'name': i, 'id': i, "deletable": True, "selectable": True} for i in df.columns if
                     i[1:].isdigit() != 1 and i.startswith('Unn') != 1],

            editable=True,
            page_size=50,
            style_table={'height': '500px', 'overflowY': 'auto', 'width': '98%'},
            style_cell={
                'overflow': 'hidden',
                'textOverflow': 'ellipsis',
                'maxWidth': 0,
                'fontSize': '1rem',
                'textAlign': 'center',
                'color' : 'black'
            },
            fixed_rows={'headers': True},
            tooltip_data=[
                {
                    column: {'value': str(value), 'type': 'markdown'}
                    for column, value in row.items()
                } for row in df.to_dict('records')
            ],
            style_cell_conditional=[
                {
                    'if': {'column_id': c},
                    'textAlign': 'center',
                    'width': '8%',

                    'if': {'column_id': 'Date'},
                     'width': '18%'} for c in df.columns ],
            style_header={
                'backgroundColor': 'rgb(230, 230, 230)',
                'fontWeight': 'bold',
                'color': 'black'
            },
            filter_action="native",
            sort_action="native",
            sort_mode="multi",
            column_selectable="single",
            selected_columns=[],
            selected_rows=[],
            page_action="native",
            page_current=0,
            export_format='xlsx',
            export_headers='display',
            merge_duplicate_headers=True
        ),

        html.Hr(),  # horizontal line
    ])

@app.callback([Output('datatablehidden', 'children'), Output('retrieve', 'children')],
              [Input('upload-data', 'contents'), Input("my-toggle-switch", "on"), ],
              [State('upload-data', 'filename'),
               State('upload-data', 'last_modified'),
               State('retrieve', 'children'),
               State('datatablehidden', 'children')])
def update_output(list_of_contents, on, list_of_names, list_of_dates, retrieve, content):
    if on == 0:
        raise PreventUpdate
    if list_of_contents is not None:
        pairs = zip(list_of_contents, list_of_names, list_of_dates)
        content = [
            parse_contents(c, n, d) for c, n, d in pairs]
        retrieve = list_of_names
        return content, retrieve
    else:
        return (no_update, no_update)

@app.callback(Output('output-data-upload', 'children'),
              [Input('datatablehidden', 'children')]
              )
def retrieve(retrieve):
    if retrieve == None or retrieve == []:
        raise PreventUpdate
    return retrieve

@app.callback(ServersideOutput('datastore', 'data'),
              [Input('datatablehidden', 'children')],)

def retrieve(retrieve):
    if retrieve == None or retrieve == []:
        raise PreventUpdate
    else:
        xx = retrieve[0]['props']['children'][2]['props']['data']
        return xx

@app.callback(Output('tab4DashTable', 'children'),
              [Input('datatablehidden', 'children')],)
def retrieve4(retrieve):
    return retrieve

@app.callback(
    Output('datatable-interactivity', 'style_data_conditional'),
    [Input('datatable-interactivity', 'selected_columns')]
)
def update_styles(selected_columns):
    return [{
        'if': {'column_id': i},
        'background_color': 'red'
    } for i in selected_columns]

    # Output("opcLoad","children") : for load left and right side,
    # for this created a hiddev div as opcLoad,
    # Output('tab2','children') : also hidden tab, for the graph

@app.callback([Output("opcLoad", "children"), Output('upload-data', 'style')],
              [Input("my-toggle-switch", "on")]
              )
def opcLoadingData(on):
    ocploadlist = []

    if on == 1:
        visibilty = {
            'height': '35px',
            'lineHeight': '25px',
            'borderWidth': '1px',
            'borderStyle': 'dashed',
            'borderRadius': '5px',
            'textAlign': 'center',
            'margin': '20px 20px 20px -100px',

            'visibility': 'visible'}
        data_list = []

        ocploadlist = html.Div(className='userControlDownSideCreated',
                               children=[html.Div(className='userControlDownLeftSide',

                                                  children=[html.Div(className='aa',
                                                                     children=[html.Div(
                                                                         dcc.Dropdown(id='dropdownLeft',
                                                                                      options=[{'label': i, 'value': i}
                                                                                               for i in data_list if
                                                                                               i != 'date'],
                                                                                      multi=False,
                                                                                      style={'cursor': 'pointer',},
                                                                                      className='stockSelectorClass',
                                                                                      clearable=False,
                                                                                      placeholder='Select your parameters...',

                                                                                      ),
                                                                     ),
                                                                         html.Div([html.Button('Show', id='showLeft',
                                                                                               n_clicks=0,
                                                                                               style={'height': '40px',
                                                                                                      'width': '80px',
                                                                                                      'fontSize': '1.2rem'}),

                                                                                   html.Button('Delete', id='clearLeft',
                                                                                               n_clicks=0,
                                                                                               style={'height': '40px',
                                                                                                      'width': '80px',
                                                                                                      'fontSize': '1.2rem'})],
                                                                                  className='buttons'),
                                                                         html.Div(id='leftSideDropdownHidden',
                                                                                  children=[],
                                                                                  style={'display': 'None'}),
                                                                         # html.Div(id='leftSideDropdown', children=[]),
                                                                         html.Div([dbc.Checklist(
                                                                             id='choosenChecklistLeft',
                                                                             options=[{'label': i, 'value': i} for i in
                                                                                      []],
                                                                             value=[],
                                                                             labelStyle={'display': 'Block'},
                                                                         ), ], style={'marginTop': '8px',
                                                                                      'marginLeft': '8px',
                                                                                      'visibility': 'hidden'}),
                                                                         html.Div(
                                                                             [

                                                                                 dbc.Modal(
                                                                                     [
                                                                                         dbc.ModalHeader("INFORMATION"),
                                                                                         dbc.ModalBody(
                                                                                             "Vous pouvez choisir maximum 20 valeur"),
                                                                                         dbc.ModalFooter(
                                                                                             dbc.Button("Close",
                                                                                                        id="close",
                                                                                                        className='ml-auto')
                                                                                         ),
                                                                                     ],
                                                                                     id="modal",
                                                                                 ),
                                                                             ])
                                                                     ])], ),
                                         html.Div(className='userControlDownRightSide',
                                                  children=[
                                                      html.Div(
                                                          className='div-for-dropdown',
                                                          children=[
                                                              html.Div(
                                                                  dcc.Dropdown(id='dropdownRight',
                                                                               options=[{'label': i, 'value': i} for i
                                                                                        in extra_data_list],
                                                                               multi=False,
                                                                               value='',
                                                                               style={'cursor': 'pointer',},
                                                                               className='stockSelectorClass',
                                                                               clearable=False,
                                                                               placeholder='Select your parameters...',

                                                                               ),
                                                              ),
                                                              html.P('Enter Y and X values', style = {'margin' : '1rem 6rem', }),
                                                              html.Div([dbc.Input(
                                                                  id='inputRightY_axis',
                                                                  type="text",
                                                                  min=-10000, max=10000, step=1, bs_size="sm",
                                                                  style={'width': '7rem'},
                                                                  placeholder='Y axis value',
                                                                  autoFocus=True, ),
                                                                        dbc.Input(id='inputRightX_axis',
                                                                                  type="text",
                                                                                  min=-10000, max=10000, step=1,
                                                                                  bs_size="sm", style={'width': '7rem'},
                                                                                  placeholder='X axis value',
                                                                                  autoFocus=True, ),
                                                                  dbc.Tooltip(
                                                                      "For using 'date', the format is  "
                                                                      "'yyyy-mm-dd hh:mm:ss'"
                                                                      " 2021-07-21 12:15:18",
                                                                      target="inputRightX_axis",
                                                                      placement='right',
                                                                  ),
                                                                        ], id="styled-numeric-input",className = "add_design2" ),
                                                              html.Div([dbc.Button('Show', id='showRight',color = 'primary',
                                                                                    n_clicks=0, className = 'mr-1', style = {'width':'5rem', 'fontSize':'1rem'}),
                                                                        dbc.Button('Clear', id='valueClearRight', color = 'warning',
                                                                                    n_clicks=0, className = 'mr-1',style = {'width':'5rem','fontSize':'1rem'}),
                                                                        dbc.Button('Delete', id='clearRight',color = 'danger',
                                                                                    n_clicks=0, className = 'mr-1',style = {'width':'5rem','fontSize':'1rem'}),

                                                                        ],
                                                                       className='buttons'),
                                                              html.Div([dbc.Checklist(
                                                                  id='choosenChecklistRef',
                                                                  options=[{'label': i, 'value': i} for i in
                                                                           []],
                                                                  value=[],
                                                                  labelStyle={'display': 'Block','marginTop': '2rem',},
                                                              ), ], style={'marginTop': '2rem',
                                                                           'marginLeft': '1rem',
                                                                           }),
                                                              html.Div(id='rightSideDropdownHidden', children=[],
                                                                       style={'visibility': 'hidden'}),
                                                              html.Div(id="rightSideDropdown", children=[])
                                                          ],
                                                      ),
                                                  ]),
                                         ])
        return (ocploadlist, visibilty)

    else:
        return (ocploadlist, {'visibility': 'hidden'})


@app.callback(Output("dropdownLeft", "options"),
              [Input("datastore", "data")])
def dropdownlistcontrol(retrieve):
    if retrieve == None:
        raise PreventUpdate

    df = pd.DataFrame(retrieve)
    if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
        dff = [{'label': i, 'value': i} for i in df['ID'].unique() if i.startswith('Un') != 1 and i != 'index' and i != 'Date']
        return dff

    else :
        dff = [{'label': i, 'value': i} for i in df.columns if i.startswith('Un') != 1 and i != 'index' and i != 'date']
        return dff

@app.callback(
    [Output('choosenChecklistLeft', 'options'),
     Output('choosenChecklistLeft', 'style'),
     Output('choosenChecklistLeft', 'value'),
     Output("leftSideDropdownHidden", "children"),
     Output("leftSidedroptValue", "children"),
     Output("deletedval", "children")],
    [Input("showLeft", "n_clicks"),
     Input("clearLeft", "n_clicks"),
     ],
    [State("dropdownLeft", "value"),
     State("leftSideDropdownHidden", "children"),
     State('choosenChecklistLeft', 'value'),
     State('deletedval', 'children')],
)
# left side dropdown-checklist relation
#########

def displayLeftDropdown(n_clicks1, nc2, dropval, valeur, value, deletedval):
    if dropval == None or deletedval == None:
        raise PreventUpdate
    q1 = dash.callback_context.triggered[0]["prop_id"].split(".")[0]
    a = []
    a.append(dropval)
    for i in a:
        if q1 == 'showLeft' and i not in valeur:
            valeur.append(dropval)
        if q1 == 'clearLeft' and i not in deletedval:
            pass
    if q1 == 'showLeft':
        return [{'label': i, 'value': i} for i in valeur], {'visibility': 'visible'}, [], valeur, valeur, deletedval

    if q1 == 'clearLeft':

        for k in range(len(value)):
            valeur.remove(value[k])
            deletedval.append(value[k])

        return [{'label': i, 'value': i} for i in valeur], {'visibility': 'visible'}, [], valeur, valeur, deletedval
    else:
        no_update, no_update, no_update, no_update, no_update, no_update

@app.callback(
    Output("modal", "is_open"),
    [Input("showLeft", "n_clicks"), Input("close", "n_clicks")],
    [State("modal", "is_open"),
     State("leftSideDropdownHidden", "children")],
)
def toggle_modal_5(n1, n2, is_open, children):
    if len(children) > 20:
        return not is_open
    return is_open

@app.callback(
    [Output('choosenChecklistRef', 'options'),
     Output('choosenChecklistRef', 'style'),
     Output('choosenChecklistRef', 'value'),
     Output("checklistvaleurhidden", "children"),
     Output("deletedvalref", "children"),

     ],
    [Input("showRight", "n_clicks"),
     Input("clearRight", "n_clicks"),
     ],
    [
    State('inputRightY_axis', "value"),
    State('inputRightX_axis', "value"),
    State("dropdownRight", "value"),
    State("checklistvaleurhidden", "children"),
    State('choosenChecklistRef', 'value'),
    State('deletedvalref', 'children'),],)
# left side dropdown-checklist relation
#########

def displayLeftDropdownRight(n_clicks1, nc2, valy, valx,dropval, valeur, value, deletedval):
    if dropval == None or deletedval == None or valy == None or valx == None:
        raise PreventUpdate
    def mesure1(textRight):
        if textRight == "Mass de Bois":
            return "g"
        elif textRight == 'Volume gaz':
            return 'm3'

        elif textRight == 'Vitesse de rotation':
            return 'tour/mn'

        elif textRight in {'Puissance Thermique', 'Puissance Electrique'}:
            return "W"

        elif textRight in {'CO', 'CO2', 'NO', 'NOX', 'CX'}:
            return "% MOL"

        elif textRight == 'Temperature de Fumée':
            return '°K'
    q1 = dash.callback_context.triggered[0]["prop_id"].split(".")[0]
    if dropval != None:
        x = f'''Entered values for {dropval} ( {mesure1(dropval)} ) are : Y-axis : {valy} and X-axis : {valx}'''
    a = []
    a.append(x)
    for i in a:
        if q1 == 'showRight' and i not in valeur:
            valeur.append(x)



        if q1 == 'clearRight' and i not in deletedval:
            pass
    if q1 == 'showRight':

        return [{'label': i, 'value': i} for i in valeur], {'visibility': 'visible'}, [], valeur, deletedval

    if q1 == 'clearRight':

        for k in range(len(value)):
            valeur.remove(value[k])
            deletedval.append(value[k])

        return [{'label': i, 'value': i} for i in  valeur], {'visibility': 'visible'}, [], valeur, deletedval
    else:
        no_update, no_update, no_update, no_update, no_update


@app.callback([Output("inputRightY_axishidden", "children"),
               Output("inputRightX_axishidden", "children"),
               Output('checklistvaleurhidden2', "children"),],
              [Input('choosenChecklistRef','value'),],)
def Inputaxis(okval):
    if okval != []:
        y_val = []
        x_val = []
        checklist = []
        for i in okval:
            m = i.split(' ')
            for k in m:
                if k != 'Y-axis':
                    checklist.append(k)
                    break
            i = []
            if len(m[-1]) < 7:
                x_val.append(m[-1])
                i = m[:-1]
            else :
                a = ''.join([i for i in m[-2:]])
                b = a[:10] + ' ' + a[10:]
                x_val.append(b)
                i = m[:-2]
            for j in range(len(i)):
                if (i[j]).isdecimal():
                    y_val.append(i[j])
                if '.' in i[j]:
                    y_val.append(float(i[j]))
        return y_val,x_val,checklist
    else : return [],[],[]

@app.callback([Output("inputRightY_axis", "value"),
               Output("inputRightX_axis", "value")],
              [Input('valueClearRight', 'n_clicks')],
              [State("inputRightY_axis", "value"),
              State("inputRightX_axis", "value")]
              )
def clear(nclick, st1, st2):

    if st1 == None or st2 == None:
        raise PreventUpdate
    if nclick > 0:
        st1 = 0
        st2 = 0
        return (st1, st2)
    else:
        (no_update, no_update)

@app.callback(Output('tabs-content-classes', 'children'),
              [Input('tabs-with-classes', 'value')],
              )
def render_content(tab):
    if tab == 'tab-1':
        return html.Div([
            html.Div(id='tab1Data')
        ])


    if tab == 'tab-4':
        return html.Div([
            html.Div(id='tab4Data')
        ])
    else:
        pass

@app.callback(Output('tab1Data', 'children'),
              [Input("my-toggle-switch", "on"),
               Input("leftSidedroptValue", "children")],
              [State('tabs-with-classes', 'value')]
              )
def LoadingDataTab1(on, dropdownhidden, tab):
    if on == 1 and tab == 'tab-1':
        loadTab1 = html.Div([html.Div([html.Div([html.Div([
            dcc.Dropdown(id='firstChoosenValue',
                         options=[{'label': i, 'value': i} for i in
                                  dropdownhidden],
                         multi=False,
                         style={'cursor': 'pointer', 'width': '180px'},

                         clearable=True,
                         placeholder='First Value...',
                         ),
            dbc.Input(id='leftIntegralFirst',
                      type="text",
                      debounce=True,
                      min=-10000, max=10000, step=1,
                      bs_size="sm",
                      style={'width': '8rem', 'marginTop': '1.5rem'},
                      autoFocus=True,
                      placeholder="first point"),
            dbc.Input(id='leftIntegralSecond',
                      type="text",
                      debounce=True,
                      min=-10000, max=10000, step=1,
                      bs_size="sm",
                      style={'width': '8rem', 'marginTop': '1.5rem'},
                      autoFocus=True,
                      placeholder="second point"),
            dbc.Input(id='leftIntegral',
                      type="text",
                      min=-10000, max=10000, step=1,
                      bs_size="sm",
                      style={'width': '9rem', 'marginTop': '1.5rem'},
                      autoFocus=True,
                      placeholder="total integration"),
        ]),
            html.Div([html.Button("Save", id="write_excel", n_clicks=0,
                                  style={'fontSize': '1rem', 'width': '4rem',
                                         'margin': '1rem'},
                                  ),
                      html.A(html.Button("Download Data", id='download_data',
                                         n_clicks=0,
                                         style={'fontSize': '1rem',
                                                'width': '9rem',
                                                'margin': '1rem'}, ),
                             id='download_excel',
                             # # download="rawdata.csv",
                             href="/download_excel/",
                             # target="_blank"
                             ),
                      html.Button("Save DB", id="save_to_db", n_clicks=0,
                                  style={'fontSize': '1rem', 'width': '5rem',
                                         'margin': '1rem'},
                                  ),
                      ], className='ad')

        ]),
            html.Div([dbc.Checklist(
                id='operateur',
                options=[{'label': i, 'value': i} for i in
                         ['Plus', 'Moins', 'Multiplie', 'Division']],
                value=[],
                labelStyle={'display': 'Block'},
            ), ]),
            html.Div([dcc.Dropdown(id='secondChoosenValue',
                                   options=[{'label': i, 'value': i} for i in
                                            dropdownhidden],
                                   multi=False,
                                   style={'cursor': 'pointer','width': '180px'},

                                   clearable=True,
                                   placeholder='Second Value...',
                                   ),
                      dbc.Input(id='rightIntegralFirst',
                                type="text",
                                min=-10000, max=10000, step=1,
                                bs_size="sm",
                                style={'width': '8rem', 'marginTop': '1.5rem'},
                                autoFocus=True,
                                placeholder="first point"),
                      dbc.Input(id='rightIntegralSecond',
                                type="text",
                                min=-10000, max=10000, step=1,
                                bs_size="sm",
                                style={'width': '8rem', 'marginTop': '1.5rem'},
                                autoFocus=True,
                                placeholder="second point"),
                      dbc.Input(id='rightIntegral',
                                type="text",
                                min=-10000, max=10000, step=1,
                                bs_size="sm",
                                style={'width': '9rem', 'marginTop': '1.5rem'},
                                autoFocus=True,
                                placeholder="total integration")]),
            html.Div([dbc.Input(id='operation',
                                type="text",
                                min=-10000, max=10000, step=1,
                                bs_size="sm",
                                style={'width': '10rem', 'marginTop': '2rem',
                                       'height': '5rem', 'textAlign': 'center'},
                                autoFocus=True,
                                placeholder="result"),
                      dbc.Input(id='intersection',
                                type="text",
                                min=-10000, max=10000, step=1,
                                bs_size="sm",
                                style={'width': '10rem', 'marginTop': '2rem',
                                       'height': '2rem', 'textAlign': 'center'},
                                autoFocus=True,
                                placeholder="Intersection")], className='aa')],
            className='ab-tab-1'),
            dbc.Modal(
                [
                    dbc.ModalHeader("Save Your Table In Database"),

                    dbc.Input(id='input_user',
                              type="text",
                              min=-10000, max=10000, step=1, bs_size="sm",
                              style={'width': '31rem', },
                              placeholder='Enter Username',
                              autoFocus=True, ),
                    dbc.Input(id='input_pswrd',
                              type="password",
                              min=-10000, max=10000, step=1, bs_size="sm",
                              style={'width': '31rem', },
                              placeholder='Enter Password',
                              autoFocus=True, ),
                    dbc.Input(id='input_dbname',
                              type="text",
                              min=-10000, max=10000, step=1, bs_size="sm",
                              style={'width': '31rem', },
                              placeholder='Enter Database name',
                              autoFocus=True, ),
                    dbc.Input(id='input_tblname',
                              type="text",
                              min=-10000, max=10000, step=1, bs_size="sm",
                              style={'width': '31rem', },
                              placeholder='Enter Table name',
                              autoFocus=True, ),
                    dbc.ModalFooter([
                        dbc.Button("Okey", id="ok_reel_analysis", className='ml-auto'),
                        dbc.Button("Close", id="close_reel_analysis", className='ml-auto')]
                    ),
                ],
                id="modal_reel_analysis",
            ),
            html.Div([dcc.RadioItems(id="radiograph",
                                     options=[
                                         {'label': 'Point', 'value': 'markers'},
                                         {'label': 'Line', 'value': 'lines'},
                                         {'label': 'Line + Point', 'value': 'lines+markers'},

                                     ],
                                     value='markers',
                                     labelClassName='groupgraph2',
                                     labelStyle={'margin': '10px', 'display': 'inline-block'},
                                     inputStyle={'margin': '10px', }
                                     ),
                      html.Div([html.P('Shift Shaded Area (First-Second)'),
                                dbc.Input(id='minimumValueGraphFirst',
                                          type="text",
                                          min=-10000, max=10000, step=1,
                                          bs_size="sm",
                                          value=0,
                                          style={'width': '8rem', 'marginLeft': '20px'},
                                          placeholder="Minimum Value of Graph for First..."),
                                dbc.Tooltip(
                                    "For the first parameter's (left) shaded area  ",
                                    target="minimumValueGraphFirst",
                                    placement='bottom',
                                ),
                                dbc.Input(id='minimumValueGraphSecond',
                                          type="text",
                                          min=-10000, max=10000, step=1,
                                          bs_size="sm",
                                          value=0,
                                          style={'width': '8rem', 'marginLeft': '20px'},
                                          placeholder="Minimum Value of Graph for Second..."),
                                dbc.Tooltip(
                                    "For the second parameter's(right) shaded area ",
                                    target="minimumValueGraphSecond",
                                    placement='right',
                                ),
                                html.Div([dcc.Dropdown(id="dropadd1",
                                                       options=[
                                                           {'label': 'Note', 'value': 'note'},
                                                           {'label': 'Header', 'value': 'header'},
                                                           {'label': 'x-axis', 'value': 'x_axis'},
                                                           {'label': 'y-axis', 'value': 'y_axis'},
                                                           {'label': 'y2-axis', 'value': 'y2_axis'},
                                                           {'label': 'y3-axis', 'value': 'y3_axis'}

                                                       ],
                                                       value='header',
                                                       ),
                                          dcc.Textarea(
                                              id='textarea1',
                                              value='',
                                              style={'width': '15rem', 'marginTop': '0.5rem'},
                                              autoFocus='Saisir',
                                          ),
                                          ], className='aatab1'),
                                html.Button('Add Text', id='addText1', n_clicks=0,
                                            style={'fontSize': '1rem', 'width': '4rem',
                                                   'margin': '1rem'},),
                                ], className='shift'),

                      ], className='abcd'),

            html.Div([dcc.Dropdown(id='shiftaxisdrop',
                                   options=[{'label': i, 'value': i} for i in
                                            dropdownhidden],
                                   multi=False,
                                   style={'cursor': 'pointer', 'width': '180px', 'margin': '1rem'},

                                   clearable=True,
                                   placeholder='Choose Value...',
                                   ),

                      html.Div(id='shiftaxis',
                               children=[
                                   dbc.Input(id='shift_x_axis',
                                             type="number",
                                             min=-100000, max=100000, step=1,
                                             bs_size="sm",
                                             value=0,
                                             style={'width': '8rem', },
                                             placeholder="Shift X axis..."),
                                   dbc.Input(id='shift_y_axis',
                                             type="number",
                                             min=-100000, max=100000, step=1,
                                             bs_size="sm",
                                             value=0,
                                             style={'width': '8rem', },
                                             placeholder="Shift Y axis..."),
                                   dbc.Button("Ok", id="tab1send", outline=True, n_clicks=0,
                                              color="primary",
                                              className='mr-2'),
                               ], className='abcd', style={'display': 'None'}),
                      dbc.Button("See Surface", id="valuechange", n_clicks=0,
                                 color="warning", style={'height': '2.5em', 'margin': '1.8rem'}),
                      dbc.Button("Clean Surface", id="cleanshape", n_clicks=0,
                                 color="danger", style={'height': '2.5em', 'margin': '1.8rem'}),
                      dcc.Dropdown(id='charactere',
                                   options=[{'label': i, 'value': i} for i in [8,12,14,16,18,20,22,24,26,28]],
                                   value = 16,
                                   multi=False,
                                   style={'cursor': 'pointer', 'width': '120px', 'margin': '1rem'},

                                   clearable=True,
                                   placeholder='Charactere...',
                                   ),

                      ], className='abcd'),


            html.Div([html.Div([dcc.Graph(id='graph',
                                config={'displayModeBar': True,
                                        'scrollZoom': True,
                                        'displaylogo': False,
                                        'modeBarButtonsToAdd': [
                                            'drawline',
                                            'drawrect',
                                            'drawopenpath',
                                            'select2d',
                                            'eraseshape',

                                        ]},
                                style={'marginTop': '20px'},
                                figure={
                                    'layout': {'legend': {'tracegroupgap': 0},

                                               }
                                }

                                ),

                    html.Div([daq.BooleanSwitch(
                                        id="selectyaxis",
                                        label="Multiple Y-axis",
                                        labelPosition="bottom",
                                        color='red',

                                    ),

                        dbc.Tooltip(
                            "You can select maximum 3 y-axis",
                            target="selectyaxis",
                            placement='bottom',
                        ),
                    ], className='pluraly'),
                                html.Div(daq.Slider(id="sliderHeightTab1",
                                                    max=2100,
                                                    min=400,
                                                    value=530,
                                                    step=100,
                                                    size=420,
                                                    vertical=True,
                                                    updatemode='drag'), style={'margin': '20px'}),
                                ],
                     className='abcdb_page1'),

            html.Div([daq.Slider(id="sliderWidthTab1",
                                 max=2000,
                                 min=600,
                                 value=1000,
                                 step=100,
                                 size=500,
                                 updatemode='drag'),], className = 'design_page1')], className = 'design_page2'),
            html.Div(id='output-data-upload', children=[]),
        ]),
        return loadTab1

@app.callback(Output("leftSideChecklistValueHidden", "children"),
              [Input('choosenChecklistLeft', 'value'), ],
              [State("leftSideChecklistValueHidden", "children")]
              )
def res(val, hiddenval):
    if val == None:
        raise PreventUpdate
    hiddenval = val
    return hiddenval

@app.callback(Output("radiographhidden", "children"),
              [Input("radiograph", "value")],
              )
def radio(radiograph):
    return radiograph


@app.callback(Output("radiographhiddentab4", "children"),
              [Input("radiograph4", "value")],
              )
def radiotab4(radiograph):
    if radiograph == None:
        raise PreventUpdate
    return radiograph


@app.callback(Output("sliderHeightTab1hidden", "children"),
              [Input("sliderHeightTab1", "value")],
              )
def tabheight(height):
    return height


@app.callback(Output("sliderWidthTab1hidden", "children"),
              [Input("sliderWidthTab1", "value")],
              )
def tabwidth(width):
    return width


@app.callback(Output("sliderHeightTab1hiddentab4", "children"),
              [Input("sliderHeightTab4", "value")],
              )
def tabheighttab4(height):
    return height

@app.callback(Output("hiddenvaluechange", "children"),
              [Input("valuechange", "n_clicks")],
              )
def valuechange(val):
    if val == None:
        raise PreventUpdate
    return val

@app.callback(Output("hiddencharactere", "children"),
              [Input("charactere", "value")],
              )
def charactereval(val):
    if val == None:
        raise PreventUpdate
    return val

@app.callback(Output("hiddencleanshape", "children"),
              [Input("cleanshape", "n_clicks")],
              )
def cleanshape(val):
    if val == None:
        raise PreventUpdate
    return val

@app.callback(Output("sliderWidthTab1hiddenTab4", "children"),
              [Input("sliderWidthTab4", "value")],
              )
def tabwidthtab4(width):
    return width


@app.callback(Output("minimumValueGraphhiddenfirst", "children"),
              [Input("minimumValueGraphFirst", "value")],
              )
def minfirst(minvalfirst):
    return minvalfirst


@app.callback(Output("minimumValueGraphhiddensecond", "children"),
              [Input("minimumValueGraphSecond", "value")],
              )
def minsecond(minvalsecond):
    return minvalsecond


@app.callback(Output("firstchoosenvalhidden", "children"),
              [Input("firstChoosenValue", "value")],
              [State("firstchoosenvalhidden", "children")]
              )
def firstchleft(firstchoosen, hiddenfirstchoosen):
    hiddenfirstchoosen.append(firstchoosen)
    return hiddenfirstchoosen


@app.callback(Output("firstchoosenvalhiddentab4", "children"),
              [Input("firstChoosenValueTab4", "value")],
              [State("firstchoosenvalhiddentab4", "children")]
              )
def firstchlefttab4(firstchoosen4, hiddenfirstchoosen4):
    hiddenfirstchoosen4.append(firstchoosen4)
    return hiddenfirstchoosen4

@app.callback(Output("secondchoosenvalhidden", "children"),
              [Input("secondChoosenValue", "value")],
              )
def secondchleft(secondchoosen):
    return secondchoosen


@app.callback(Output("secondchoosenvalhiddentab4", "children"),
              [Input("secondChoosenValueTab4", "value")],
              )
def secondchleftTab4(secondchoosen):
    return secondchoosen

@app.callback(Output("secondchoosenvalhiddenpr", "children"),
              [Input("secondChoosenValuedb", "value")],
              )
def secondchleftpr(secondchoosen):
    return secondchoosen


@app.callback(Output("leftintegralfirsthidden", "children"),
              [Input("leftIntegralFirst", "value")],
              )
def firstchright(leftintfirst):
    return leftintfirst

@app.callback(Output("leftintegralfirsthiddenTab4", "children"),
              [Input("leftIntegralFirstTab4", "value")],
              )
def firstchrighttab4(leftintfirst):
    return leftintfirst

@app.callback(Output("leftintegralsecondhidden", "children"),
              [Input("leftIntegralSecond", "value")],
              )
def secondchright(leftintsecond):
    return leftintsecond


@app.callback(Output("leftintegralsecondhiddentab4", "children"),
              [Input("leftIntegralSecondTab4", "value")],
              )
def secondchright(leftintsecond):
    return leftintsecond

@app.callback(Output("rightintegralfirsthidden", "children"),
              [Input("rightIntegralFirst", "value")],
              )
def rightfrst(rightintfirst):
    return rightintfirst


@app.callback(Output("rightintegralfirsthiddentab4", "children"),
              [Input("rightIntegralFirstTab4", "value")],
              )
def rightfrsttab4(rightintfirst):
    return rightintfirst

@app.callback(Output("rightintegralsecondhidden", "children"),
              [Input("rightIntegralSecond", "value")],
              )
def rightscnd(rightintsecond):
    return rightintsecond


@app.callback(Output("rightintegralsecondhiddentab4", "children"),
              [Input("rightIntegralSecondTab4", "value")],
              )
def rightscndtab4(rightintsecond):
    return rightintsecond

@app.callback(Output("hiddenchoosenChecklistLeft", "children"),
              [Input("choosenChecklistLeft", "value")],
              )
def checklistval(val):
    return val

# for show graph and changement

@app.callback(Output('shiftaxisdrophidden', 'children'),
              [Input('shiftaxisdrop', 'value')], )
def relay(val):
    return val


@app.callback(Output('shift_x_axishidden', 'children'),
              [Input('shift_x_axis', 'value')], )
def relay2(val):
    return val


@app.callback(Output('shift_y_axishidden', 'children'),
              [Input('shift_y_axis', 'value')], )
def relay3(val):
    return val


@app.callback(Output('tab1sendhidden', 'children'),
              [Input('tab1send', 'n_clicks')], )
def relay7(val):
    return val

@app.callback(Output('selectyaxishidden', 'children'),
              [Input("selectyaxis", "on")], )
def relay8(val):
    return val

@app.callback(Output('shiftaxis', 'style'),
              [Input('shiftaxisdrop', 'value')])
def shiftingaxes(val):
    if val == None:
        return {'display': 'None'}
    return {'visibility': 'visible', 'marginTop': '2rem'}



@app.callback( Output('graphhidden', 'figure'),
              [Input('selectyaxishidden', 'children'),
               Input('hiddenTextxaxis', 'children'),
               Input('hiddenTextyaxis', 'children'),
               Input('hiddenTextHeader', 'children'),
               Input('hiddenTextNote', 'children'),
               Input('hiddenTexty1axis', 'children'),
               Input('hiddenTexty2axis', 'children'),
               Input("hiddenchoosenChecklistLeft", "children"),
               Input("radiographhidden", "children"),
               Input("sliderHeightTab1hidden", "children"),
               Input("sliderWidthTab1hidden", "children"),
               Input('minimumValueGraphhiddenfirst', 'children'),
               Input('minimumValueGraphhiddensecond', 'children'),
               Input('firstchoosenvalhidden', 'children'),
               Input('secondchoosenvalhidden', 'children'),
               Input('checklistvaleurhidden2', "children"),
               Input('inputRightY_axishidden', 'children'),
               Input('inputRightX_axishidden', 'children'),
               Input('tab1sendhidden', 'children'),
               Input('hiddenvaluechange', 'children'),
               Input('hiddencleanshape', 'children'),
               Input('hiddencharactere', 'children'),
               ],
              [State('shiftaxisdrophidden', 'children'),
               State('shift_x_axishidden', 'children'),
               State('shift_y_axishidden', 'children'),
               State('datastore', 'data'),
               State('leftintegralfirsthidden', 'children'),
               State('leftintegralsecondhidden', 'children'),
               State('rightintegralfirsthidden', 'children'),
               State('rightintegralsecondhidden', 'children'),
               State('pointLeftFirst', 'children'),
               State('pointRightFirst', 'children'),


               ]
              )
def res2(on,g1, g2, head, note,y1,y2, val, radiograph, sliderheight, sliderwidth,
         minValfirst, minValsecond, firstchoosen, secondchoosen, rightsidedrop, right_y_axis, right_x_axis,
         nclick, nc, cleanclick,car, axis, shift_x, shift_y, retrieve, leftfirstval, leftsecondval,
         rightfirstval, rightsecondval, firstshape, secondshape):
    if retrieve == None or retrieve == [] or nc == None or note == None:
        raise PreventUpdate
    if retrieve != []:
        df = pd.DataFrame(retrieve)
        df['index'] = df.index
        df = df.reindex(columns=sorted(df.columns, reverse=True))
        baseval = ''
        if 'date' not in df.columns:
            for col in df.columns:
                if 'Temps' in col:
                    baseval += col
                    dt = df[baseval]
                else :
                    dt = np.arange(0, df.shape[0])
            if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns :
                dff = df[df['ID'] == firstchoosen[-1]]
                dff = dff.copy()
                index = np.arange(0, len(dff['ID']))
                dff.reset_index(drop=True, inplace=True)
                dff.set_index(index, inplace=True)
                dt = dff[['Date']]
                dt.columns = ['Date']
                dt = dt['Date'].apply(lambda x : x[:10] + '_' + x[11:])
                dff2 = df[df['ID'] == secondchoosen]
                dff2 = dff2.copy()
                index = np.arange(0, len(dff2['ID']))
                dff2.reset_index(drop=True, inplace=True)
                dff2.set_index(index, inplace=True)
                dt2 = dff2[['Date']]
                dt2.columns = ['Date']
                dt2 = dt2['Date'].apply(lambda x: x[:10] + '_' + x[11:])

        if 'date' in df.columns:
            if type(df['date'][0]) == 'str':
                df_shape = df.copy()
                df_shape['newindex'] = df_shape.index
                df_shape.index = df_shape['date']
                dt = ["{}-{:02.0f}-{:02.0f}_{:02.0f}:{:02.0f}:{:02.0f}".format(d.year, d.month, d.day, d.hour, d.minute,
                                                                               d.second) for d in df_shape.index]

            else :
                dt = df['date']
        fig = go.Figure()
        color = {0: 'blue', 1: 'red', 2: 'green', 3: 'purple', 4: 'orange', 5: '#ff69b4',
                    6: 'black', 7: 'cyan', 8: 'magenta', 9: '#faebd7',
                    10: '#2e8b57', 11: '#eeefff', 12: '#da70d6', 13: '#ff7f50', 14: '#cd853f',
                    15: '#bc8f8f', 16: '#5f9ea0', 17: '#daa520'}
        if rightsidedrop != None :
            if right_x_axis != [] and right_y_axis != []:
                for k in range(len(rightsidedrop)):
                    if len(rightsidedrop) == len(right_x_axis) and len(rightsidedrop) == len(right_y_axis):
                        if len(right_x_axis[k]) < 6 :
                            x = [right_x_axis[k]]
                        else :
                            x = [right_x_axis[k]]
                        y = [right_y_axis[k]]
                        fig.add_trace(go.Scatter(mode="markers", x=x, y=y, marker_symbol='diamond-x',
                                                 marker_line_width=2, marker_size=8,hovertemplate=f"Ref_point{k}: %{y}/%{x}",
                                                 name=f"Ref_point{k}: {y}/{x}",
                                      )),
                        fig.update_layout(margin=dict(b=0, r=0),)
        for i_val in range(len(val)):
            if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                y_axis = df[df['ID'] == val[i_val]]['Value']
            else :
                y_axis = df[val[i_val]]
            if 'date' not in df.columns:
                if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                    x_axis = df[df['ID'] == val[i_val]]['Date']
                elif baseval != '' :
                    x_axis = df[baseval]
                else: x_axis = np.arange(0, df.shape[0])
            if 'date' in df.columns:
                x_axis = df['date']

            if nclick > 0:
                if axis == val[i_val]:
                    j = []
                    for i in df[axis]:
                        if shift_y == None:
                            raise PreventUpdate
                        else:
                            i += float(shift_y)
                            j.append(i)
                    df[axis] = pd.DataFrame(j)
                    y_axis = df[axis]
                    df.to_excel("appending.xlsx")

                    if 'date' not in df.columns:
                        p = []
                        for i in df[baseval]:
                            if shift_x == None:
                                raise PreventUpdate
                            else:
                                i += float(shift_x)
                                p.append(i)
                        df['New x-axis'] = pd.DataFrame(p)
                        x_axis = df['New x-axis']
                        df.to_excel("appending.xlsx")
                    else:
                        x_axis = df['date']
            if on == 1:
                if len(val) == 1:
                    fig.add_trace(
                        go.Scattergl(x=x_axis, y=y_axis, mode=radiograph, marker=dict(line=dict(width=0.2, color='white')),
                                     name=val[i_val]))
                elif len(val)>1 :
                    fig.add_trace(
                        go.Scattergl(x=x_axis, y=y_axis, mode=radiograph, marker=dict(line=dict(width=0.2, color='white')),
                                     name=val[i_val], yaxis = f'y{i_val+1}'))
            else :
                fig.add_trace(
                    go.Scattergl(x=x_axis, y=y_axis, mode=radiograph, marker=dict(line=dict(width=0.2, color='white')),
                                 name=val[i_val]))


            if len(firstshape) == 2 and leftfirstval != firstshape[0] and leftfirstval != []:
                if leftfirstval.startswith('T') == 1:
                    del firstshape[0]
                    firstshape.append(float(leftfirstval[2:]))
                    firstshape = sorted(firstshape)
                elif leftfirstval.isnumeric() == 1:
                    del firstshape[0]
                    firstshape.append(float(leftfirstval))
                    firstshape = sorted(firstshape)
                elif leftfirstval != None:
                    del firstshape[0]
            if len(firstshape) == 2 and leftsecondval != firstshape[
                1] and leftsecondval != None and leftsecondval != []:
                if leftsecondval.startswith('T') == 1:
                    del firstshape[1]
                    firstshape.append(float(leftsecondval[2:]))
                    firstshape = sorted(firstshape)
                elif leftsecondval.isnumeric() == 1:
                    del firstshape[1]
                    firstshape.append(float(leftsecondval))
                    firstshape = sorted(firstshape)
                elif leftsecondval != None:
                    del firstshape[1]

            if len(secondshape) == 2 and rightfirstval != secondshape[
                0] and rightfirstval != None and rightfirstval != []:
                if rightfirstval.startswith('T') == 1:
                    del secondshape[0]
                    secondshape.append(float(rightfirstval[2:]))
                    secondshape = sorted(secondshape)
                elif rightfirstval.isnumeric() == 1:
                    del secondshape[0]
                    secondshape.append(float(rightfirstval))
                    secondshape = sorted(secondshape)
                elif rightfirstval != None:
                    del secondshape[0]
            if len(secondshape) == 2 and rightsecondval != secondshape[
                1] and rightsecondval != None and rightsecondval != []:
                if rightsecondval.startswith('T') == 1:
                    del secondshape[1]
                    secondshape.append(float(rightsecondval[2:]))
                    secondshape = sorted(secondshape)
                elif rightsecondval.isnumeric() == 1:
                    del secondshape[1]
                    secondshape.append(float(rightsecondval))
                    secondshape = sorted(secondshape)
                elif rightsecondval != None:
                    del secondshape[1]
            if len(secondshape) == 2 and secondchoosen == None:
                del secondshape[1]
            if len(firstshape) == 2 and firstchoosen == None:
                del firstshape[1]

            def controlShape_Tab():
                if retrieve != []:
                    df = pd.DataFrame(retrieve)
                    df['index'] = df.index
                    df = df.reindex(columns=sorted(df.columns, reverse=True))

                pathline = ''
                pathline2 = ''
                df = pd.DataFrame(retrieve)
                if firstchoosen[-1] != None and secondchoosen != None:
                    if len(firstshape) == 2 and leftfirstval != None and leftsecondval != None:
                        if int(firstshape[1]) > int(firstshape[0]):
                            pathline = ''
                            rangeshape = range(int(firstshape[0]), int(firstshape[1]) + 2)
                            if ':' or '-' in dt[0]:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                                            pathline += 'M ' + dt[k] + ', ' + str(minValfirst) + ' L' + \
                                                        dt[k] + ', ' + str(list(dff[dff.index == k]['Value'])[0]) + ' '
                                        else:
                                            pathline += 'M ' + str(dt[k]) + ', ' + str(minValfirst) + ' L' + str(
                                                dt[k]) + ', ' + str(df[firstchoosen[-1]][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                                            pathline += ' L' + dt[k] + ', ' + str(list(dff[dff.index == k]['Value'])[0])

                                        else:
                                            pathline += ' L' + str(dt[k]) + ', ' + str(df[firstchoosen[-1]][k])
                                    elif k == rangeshape[-1]:
                                        if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                                            pathline += ' L' + dt[k - 1] + ', ' + str(minValfirst)
                                            pathline += ' Z'

                                        else:
                                            pathline += ' L' + str(dt[k - 1]) + ', ' + str(minValfirst)
                                            pathline += ' Z'

                            else:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        pathline += 'M ' + str(dt[k]) + ', ' + str(minValfirst) + ' L' + \
                                                    str(dt[k]) + ', ' + str(df[firstchoosen[-1]][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(int(dt[k])) + ', ' + str(df[firstchoosen[-1]][k])
                                pathline += ' L' + str(int(dt[k - 1])) + ', ' + str(minValfirst)
                                pathline += ' Z'

                    if len(secondshape) == 2 and rightsecondval != None and rightfirstval != None:
                        if int(secondshape[1]) > int(secondshape[0]):
                            rangeshape = range(int(secondshape[0]), int(secondshape[1] + 2))
                            if ':' or '-' in dt[0]:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                                            pathline2 += 'M ' + dt2[k] + ', ' + str(minValfirst) + ' L' + \
                                                         dt2[k] + ', ' + str(
                                                list(dff2[dff2.index == k]['Value'])[0]) + ' '

                                        else:
                                            pathline2 += 'M ' + str(dt[k]) + ', ' + str(minValsecond) + ' L' + str(
                                                dt[k]) + ', ' + str(df[secondchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                                            pathline2 += ' L' + dt2[k] + ', ' + str(
                                                list(dff2[dff2.index == k]['Value'])[0])

                                        else:
                                            pathline2 += ' L' + str(dt[k]) + ', ' + str(df[secondchoosen][k])
                                    elif k == rangeshape[-1]:
                                        if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                                            pathline2 += ' L' + dt2[k - 1] + ', ' + str(minValfirst)
                                            pathline2 += ' Z'

                                        else:
                                            pathline2 += ' L' + str(dt[k - 1]) + ', ' + str(minValsecond)
                                            pathline2 += ' Z'
                            else:
                                for k in rangeshape:

                                    if k == rangeshape[0]:
                                        pathline2 += 'M ' + str(dt[k]) + ', ' + str(minValsecond) + ' L' + str(
                                            dt[k]) + ', ' + str(df[secondchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline2 += ' L' + str(int(dt[k])) + ', ' + str(df[secondchoosen][k])
                                pathline2 += ' L' + str(int(dt[k - 1])) + ', ' + str(minValsecond)
                                pathline2 += ' Z'

                    return [dict(
                        type="path",
                        path=pathline,
                        layer='below',
                        fillcolor="#5083C7",
                        opacity=0.8,
                        line_color="#8896BF",
                    ), dict(
                        type="path",
                        path=pathline2,
                        layer='below',
                        fillcolor="#B0384A",
                        opacity=0.8,
                        line_color="#B36873",
                    )]

                if firstchoosen[-1] != None and secondchoosen == None:
                    if len(firstshape) == 2:
                        if int(firstshape[1]) > int(firstshape[0]) or int(firstshape[0]) > int(firstshape[1]):
                            pathline = ''
                            rangeshape = range(int(firstshape[0]), int(firstshape[1]) + 2)

                            if ':' or '-' or '_' in dt[0]:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                                            pathline += 'M ' + dt[k] + ', ' + str(minValfirst) + ' L' + \
                                                        dt[k] + ', ' + str(list(dff[dff.index == k]['Value'])[0]) + ' '

                                        else:
                                            pathline += 'M ' + str(dt[k]) + ', ' + str(minValfirst) + ' L' + str(
                                                dt[k]) + ', ' + str(df[firstchoosen[-1]][k]) + ' '


                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                                            pathline += ' L' + dt[k] + ', ' + str(list(dff[dff.index == k]['Value'])[0])


                                        else:
                                            pathline += ' L' + str(dt[k]) + ', ' + str(df[firstchoosen[-1]][k])

                                    elif k == rangeshape[-1]:
                                        if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                                            pathline += ' L' + dt[k - 1] + ', ' + str(minValfirst)
                                            pathline += ' Z'

                                        else:
                                            pathline += ' L' + str(dt[k - 1]) + ', ' + str(minValfirst)
                                            pathline += ' Z'

                            else:

                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        pathline += 'M ' + str(dt[k]) + ', ' + str(minValfirst) + ' L' + \
                                                    str(dt[k]) + ', ' + str(df[firstchoosen[-1]][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(int(dt[k])) + ', ' + str(df[firstchoosen[-1]][k])
                                pathline += ' L' + str(int(dt[k - 1])) + ', ' + str(minValfirst)
                                pathline += ' Z'

                        return [dict(
                            type="path",
                            path=pathline,
                            layer='below',
                            fillcolor="#5083C7",
                            opacity=0.8,
                            line_color="#8896BF",
                        )]

                if secondchoosen != None and firstchoosen[-1] == None:
                    if len(secondshape) == 2 and rightsecondval != None and rightfirstval != None:
                        if int(secondshape[1]) > int(secondshape[0]) or int(secondshape[0]) > int(secondshape[1]):
                            rangeshape = range(int(secondshape[0]), int(secondshape[1]) + 2)
                            if ':' or '-' in dt[0]:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                                            pathline2 += 'M ' + dt2[k] + ', ' + str(minValfirst) + ' L' + \
                                                         dt2[k] + ', ' + str(
                                                list(dff2[dff2.index == k]['Value'])[0]) + ' '

                                        else:
                                            pathline2 += 'M ' + str(dt[k]) + ', ' + str(minValsecond) + ' L' + str(
                                                dt[k]) + ', ' + str(df[secondchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                                            pathline2 += ' L' + dt2[k] + ', ' + str(
                                                list(dff2[dff2.index == k]['Value'])[0])

                                        else:
                                            pathline2 += ' L' + str(dt[k]) + ', ' + str(df[secondchoosen][k])
                                    elif k == rangeshape[-1]:
                                        if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                                            pathline2 += ' L' + dt2[k - 1] + ', ' + str(minValfirst)
                                            pathline2 += ' Z'

                                        else:
                                            pathline2 += ' L' + str(dt[k - 1]) + ', ' + str(minValsecond)
                                            pathline2 += ' Z'
                            else:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        pathline2 += 'M ' + str(dt[k]) + ', ' + str(minValsecond) + ' L' + str(
                                            dt[k]) + ', ' + str(df[secondchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline2 += ' L' + str(int(dt[k])) + ', ' + str(df[secondchoosen][k])
                                pathline2 += ' L' + str(int(dt[k - 1])) + ', ' + str(minValsecond)
                                pathline2 += ' Z'

                        return [dict(
                            type="path",
                            path=pathline2,
                            layer='below',
                            fillcolor="#5083C7",
                            opacity=0.8,
                            line_color="#8896BF",
                        )]

            a = []

            if nc > 0:
                a = controlShape_Tab()
            fig.update_xaxes(
                tickfont_size= car,
                tickangle=90,
                title_text='' if g1 == [] else g1[-1],
                title_standoff=25,title_font={"size": car},),

            fig.update_yaxes(
                tickfont_size=car,
                title_text='' if g2 == [] else g2[-1],
                title_standoff=25,
                titlefont=dict(
                    color="#636EFA",
                    size = car
                ),
                tickfont=dict(
                    color="#636EFA"
                )),
            fig.update_layout(
                title_text=head[-1] if len(head) > 0 else "{}".format(val[0]),
                title_font={"size": car},
                width=sliderwidth,
                height=sliderheight,
                shapes=a if nc > cleanclick else [],
                legend=dict(

                    traceorder="normal",
                    font=dict(
                        family="sans-serif",
                        size= car,
                        color=colors['figure_text']
                    ),
                    yanchor="bottom",
                    y=-0.4,
                    xanchor="right",
                    x=1,
                    bgcolor=colors['background'],
                    borderwidth=1,
                    orientation = 'h',
                ),
                yaxis2=dict(
                    title_text='y2_axis' if y1 == [] else y1[-1],
                    titlefont=dict(
                        color="#EF553B",
                        size = car
                    ),
                    tickfont=dict(
                        color="#EF553B"
                    ),
                    title_standoff = 10,
                    anchor="free",
                    overlaying="y",
                    side="right",
                    position=1

                ),

                yaxis3=dict(
                    title_text='y3_axis' if y2 ==[] else y2[-1],
                    titlefont=dict(
                        color="#00CC96",
                        size=car
                    ),
                    tickfont=dict(
                        color="#00CC96",
                    ),
                    title_standoff=10,
                    anchor="free",
                    overlaying="y",
                    side="right",
                    position=0.87
                ),


                margin=dict(
                    l=50,
                    r=50,
                    b=100,
                    t=50,
                    pad=4

                ),
                paper_bgcolor="LightSteelBlue",
                plot_bgcolor=colors['background'],),
            note_new = []
            if note != None and len(note) > 0 :
                note = note[-1].split(' ')
            x = ''
            for i in note:
                note_new.append(i)
                if len(note_new) % 4 == 0:
                    note_new[-1] += '<br>'
            x += ' '.join(note_new)
            note = [x]
            fig.add_annotation(text=note[-1] if len(note) > 0 else '',align = 'left',bordercolor = 'black',
                               borderwidth=1,xref="paper", yref="paper",bgcolor = "white", font=dict(
                                family="Arial, sans-serif" , size= car,),
                                x=1, y=1, showarrow=False)
            fig.update_layout()

        return fig

    else:
        return (no_update)


@app.callback(Output('graph', 'figure'),
              [Input("graphhidden", "figure")], )
def aa(a):
    return a

@app.callback(Output('tab4Data', 'children'),
              [Input("my-toggle-switch", "on")],
              [State('tabs-with-classes', 'value')]
              )
def LoadingDataTab4(on, tab):
    if on == 1 and tab == 'tab-4':

        data_list = ['Choose your value firstly']

        loadlist = html.Div([html.Div([
            html.Div(id='tab4first', children=[html.Div([html.Div([html.Div(
                dcc.RadioItems(id="radiographtab4",
                               options=[
                                   {'label': 'X-axis and Y-axis unlimited', 'value': 'optionlibre'},
                                   {'label': 'X-axis for each Y-axis', 'value': 'choosevalue'},
                               ],
                               # value='choosevalue',
                               labelClassName='groupgraph',
                               labelStyle={'margin': '10px'},
                               inputStyle={'margin': '10px'}
                               ), className='abtab4'),

                html.Div(dcc.RadioItems(id="radiograph4",
                               options=[
                                   {'label': 'Point', 'value': 'markers'},
                                   {'label': 'Line', 'value': 'lines'},
                                   {'label': 'Line + Point', 'value': 'lines+markers'}],
                               value='markers',
                               labelClassName='groupgraph2',
                               labelStyle={'margin': '10px'},
                               inputStyle={'margin': '10px'}
                               ), className='abtab4'),
            ], className='abtab4'),
                dcc.Loading(id = 'load1', type = 'default', children = [html.Div([
                          dcc.Dropdown(id='tabDropdownTopTab4',
                                       options=[{'label': i, 'value': i} for i in data_list],
                                       multi=True,
                                       style={'cursor': 'pointer', 'display': 'None'},
                                       className='stockSelectorClass2',
                                       clearable=True,
                                       placeholder='Select your y-axis value...',
                                       ),
                          dcc.Dropdown(id='tabDropdownDownTab4',
                                       options=[{'label': i, 'value': i} for i in data_list],
                                       multi=True,
                                       style={'cursor': 'pointer', 'display': 'None'},
                                       className='stockSelectorClass2',
                                       clearable=True,
                                       placeholder='Select your x-axis value...',
                                       ),
                          dcc.Dropdown(id='tabDropdownTop',
                                       options=[{'label': i, 'value': i} for i in data_list],
                                       multi=True,
                                       style={'cursor': 'pointer', 'display': 'None'},
                                       className='stockSelectorClass2',
                                       clearable=True,
                                       placeholder='Select your y-axis value...',
                                       ),
                          dcc.Dropdown(id='tabDropdownDown',
                                       options=[{'label': i, 'value': i} for i in data_list],
                                       multi=True,
                                       style={'cursor': 'pointer', 'display': 'None'},
                                       className='stockSelectorClass2',
                                       clearable=True,
                                       placeholder='Select your x-axis value...',
                                       ),
                          ], className='abtab4_drop'),]),
                 ], className='ac'),

                html.Div([dcc.Dropdown(id="dropadd4",
                                       options=[
                                           {'label': 'Note', 'value': 'note'},
                                           {'label': 'Header', 'value': 'header'},
                                           {'label': 'x-axis', 'value': 'x_axis'},
                                           {'label': 'y-axis', 'value': 'y_axis'},
                                           {'label': 'y2-axis', 'value': 'y2_axis'},
                                           {'label': 'y3-axis', 'value': 'y3_axis'}

                                       ],
                                       value='header',
                                       ),
                          dcc.Textarea(
                              id='textarea4',
                              value='',
                              style={'width': '15rem', 'marginTop': '0.5rem'},
                              autoFocus='Saisir',
                          ),
                          ], className='aatab4'),

                html.Button('Add Text', id='addText4', n_clicks=0, style={'marginTop': '1.5rem', 'marginLeft': '2rem'}),
                html.Div([
                    daq.BooleanSwitch(
                        id="calculintegraltab4",
                        label="Calculate Integral",
                        labelPosition="bottom",
                        color= 'red',


                    )
                ], className= 'calculIntegral'),

            ], className='tabDesigntab4', ),
            html.Div(id='tab4check', children=
            [html.Div([html.Div([dcc.Dropdown(id='firstChoosenValueTab4',
                                              options=[{'label': i, 'value': i} for i in
                                                       data_list],
                                              multi=False,
                                              style={'cursor': 'pointer', 'width': '180px'},

                                              clearable=True,
                                              placeholder='First Value...',
                                              ),
                                 dbc.Input(id='leftIntegralFirstTab4',
                                           type="text",
                                           debounce=True,
                                           min=-10000, max=10000, step=1,
                                           bs_size='sm',
                                           style={'width': '8rem', 'marginTop': '1.5rem'},
                                           autoFocus=True,
                                           placeholder="first point"),
                                 dbc.Input(id='leftIntegralSecondTab4',
                                           type="text",
                                           debounce=True,
                                           min=-10000, max=10000, step=1,
                                           bs_size="sm",
                                           style={'width': '8rem', 'marginTop': '1.5rem'},
                                           autoFocus=True,
                                           placeholder="second point"),
                                 dbc.Input(id='leftIntegralTab4',
                                           type="text",
                                           min=-10000, max=10000, step=1,
                                           bs_size="sm",
                                           style={'width': '9rem', 'marginTop': '1.5rem'},
                                           autoFocus=True,
                                           placeholder="total integration"),
                                 ]), html.Div([html.Button("Save", id="write_excelTab4", n_clicks=0,
                                                           style={'fontSize': '1rem', 'width': '4rem',
                                                                  'margin': '1rem'},
                                                           ),
                                               html.A(html.Button("Download Data", id='download_dataTab4',
                                                                  n_clicks=0,
                                                                  style={'fontSize': '1rem', 'width': '9rem',
                                                                         'margin': '1rem'}, ),
                                                      id='download_excelTab4',
                                                      # # download="rawdata.csv",
                                                      href="/download_excel/",
                                                      # target="_blank"
                                                      )
                                               ], className='abTab4')

                       ]),
             html.Div([dbc.Checklist(
                 id='operateurTab4',
                 options=[{'label': i, 'value': i} for i in
                          ['Plus', 'Moins', 'Multiplie', 'Division']],
                 value=[],
                 labelStyle={'display': 'Block'},
             ), ]),
             html.Div([
                 dcc.Dropdown(id='secondChoosenValueTab4',
                              options=[{'label': i, 'value': i} for i in
                                       data_list],
                              multi=False,
                              style={'cursor': 'pointer', 'width': '180px'},

                              clearable=True,
                              placeholder='Second Value...',
                              ),
                 dbc.Input(id='rightIntegralFirstTab4',
                           type="text",
                           min=-10000, max=10000, step=1,
                           bs_size="sm",
                           style={'width': '8rem', 'marginTop': '1.5rem'},
                           autoFocus=True,
                           placeholder="first point"),
                 dbc.Input(id='rightIntegralSecondTab4',
                           type="text",
                           min=-10000, max=10000, step=1,
                           bs_size="sm",
                           style={'width': '8rem', 'marginTop': '1.5rem'},
                           autoFocus=True,
                           placeholder="second point"),
                 dbc.Input(id='rightIntegralTab4',
                           type="text",
                           min=-10000, max=10000, step=1,
                           bs_size="sm",
                           style={'width': '9rem', 'marginTop': '1.5rem'},
                           autoFocus=True,
                           placeholder="total integration")
             ]),
             html.Div([dbc.Input(id='operationTab4',
                                 type="text",
                                 min=-10000, max=10000, step=1,
                                 bs_size="sm",
                                 style={'width': '10rem', 'marginTop': '2rem',
                                        'height': '5rem', 'textAlign': 'center'},
                                 autoFocus=True,
                                 placeholder="result"),
                       dbc.Input(id='intersectionTab4',
                                 type="text",
                                 min=-10000, max=10000, step=1,
                                 bs_size="sm",
                                 style={'width': '10rem', 'marginTop': '2rem',
                                        'height': '2rem', 'textAlign': 'center'},
                                 autoFocus=True,
                                 placeholder="Intersection")], className='aa')
             ], style={'display': 'None'},
                     className='abdbase'),

            html.Div(id='tab4second', children=[html.Div([dcc.Dropdown(id='shiftaxisdroptab4',
                                                             options=[{'label': i, 'value': i} for i in
                                                                      []],
                                                             multi=False,
                                                             style={'cursor': 'pointer', 'width':'10vw','marginLeft' : '1vw'},

                                                             clearable=True,
                                                             placeholder='Choose Value...',
                                                             ),
                                                dbc.Tooltip(
                                                    "You can change y-axis values in the same x-axis,    "
                                                    "If you finished your shifting operation, clean variable name and closed dropdown-list",
                                                    target="shiftaxisdroptab4",
                                                    placement='top',
                                                ),]),

                                                html.Div(id='shiftaxistab4',
                                                         children=[
                                                             dbc.Input(id='shift_x_axistab4',
                                                                       type="number",
                                                                       min=-100000, max=100000, step=1,
                                                                       bs_size='sm',
                                                                       value=0,
                                                                       style={'width': '8rem' },
                                                                       placeholder="Shift X axis..."),
                                                             dbc.Input(id='shift_y_axistab4',
                                                                       type="number",
                                                                       min=-100000, max=100000, step=1,
                                                                       bs_size='sm',
                                                                       value=0,
                                                                       style={'width': '8rem' },
                                                                       placeholder="Shift Y axis..."),
                                                             dbc.Button("Ok", id="tab4send", outline=True, n_clicks=0,
                                                                        color='primary',
                                                                        className='mr-1'),
                                                         ], className='abcd',
                                                         style={'display': 'None'}),
                                                dbc.Button("See Surface", id="valuechangetab4", n_clicks=0,
                                                           color="warning",
                                                           style={'height': '2.5em', 'marginLeft': '1.8rem'}
                                                           ),
                                                dbc.Button("Clean Surface", id="cleanshapetab4", n_clicks=0,
                                                           color="danger",
                                                           style={'height': '2.5em', 'marginLeft': '1.8rem'}
                                                           ),
                                                dcc.Dropdown(id='characteretab4',
                                                             options=[{'label': i, 'value': i} for i in
                                                                      [8, 12, 14, 16, 18, 20, 22, 24, 26, 28]],
                                                             value=16,
                                                             multi=False,
                                                             style={'cursor': 'pointer', 'width': '120px',
                                                                    'height': '2.5em', 'marginLeft': '0.8rem'},

                                                             clearable=True,
                                                             placeholder='Charactere...',
                                                             ),

                                                ], className='add_design'),

            html.Div([html.Div(id='tab4third', children=[dcc.Store(id='tab4datastore'),
                                               dcc.Loading(id = 'graph4load', type = 'cube', children = [dcc.Graph(id='graph4', config={'displayModeBar': True,
                                                                              'scrollZoom': True,
                                                                              'displaylogo': False,
                                                                              'modeBarButtonsToAdd': [
                                                                                  'drawopenpath',
                                                                                  'drawcircle',
                                                                                  'eraseshape',
                                                                                  'select2d',
                                                                              ]},
                                                         figure={
                                                             'layout': {'legend': {'tracegroupgap': 0},

                                                                        }
                                                         }
                                                         )]),
                                                         html.Div([daq.BooleanSwitch(
                                                             id="selectyaxistab4",
                                                             label="Multiple Y-axis",
                                                             labelPosition="bottom",
                                                             color='red',

                                                         )
                                                             ,
                                                             dbc.Tooltip(
                                                                 "You can select maximum 3 y-axis",
                                                                 target="selectyaxistab4",
                                                                 placement='bottom',
                                                             ),
                                                         ], className='pluraly'),
                                               html.Div(daq.Slider(id="sliderHeightTab4",
                                                                   max=2100,
                                                                   min=400,
                                                                   value=530,
                                                                   step=100,
                                                                   size=400,
                                                                   vertical=True,
                                                                   updatemode='drag'), style={'margin': '10px'})],
                     className='abcTab4'),

            html.Div([daq.Slider(id="sliderWidthTab4",
                                 max=2000,
                                 min=600,
                                 value=1200,
                                 step=100,
                                 size=600,
                                 updatemode='drag')],className = 'design_page1')],className = 'design_page2'),
            html.Div(id="tab4DashTable", children=[],style = {"width" : '90vw','textAlign': 'left','color': colors['text'],'marginLeft' : '5rem'} )

        ]), ],className = 'four-columns-div-user-controlsreel', style={'backgroundColor': 'white' })

        return loadlist
    else:
        no_update

@app.callback([Output('fourcolumnsdivusercontrols', 'style'),
               Output('eightcolumnsdivforcharts', 'style'), ],
               [Input('tabs-with-classes', 'value')], )
def tab4enlarger(tab):
    if tab == 'tab-4':
        return {'display': 'None'}, {'margin': '1rem'}
    else:
        return {'visibility': 'visible'}, {'visibility': 'visible'}


@app.callback(Output('tab4check', 'style'),
              [Input("calculintegraltab4", "on")],
              )
def showintegral(show):
    if show == True:
        return {'visibility': 'visible'}
    return {'display': 'None'}

@app.callback(Output('dbcheck', 'style'),
              [Input("calculintegraldb", "on")],
              )
def showintegral(show):
    if show == True:
        return {'visibility': 'visible'}
    return {'display': 'None'}

@app.callback([Output("tabDropdownTop", "options"), Output("tabDropdownDown", "options")],
              [Input("datastore", "data")])
def dropdownlistcontrol(retrieve):
    if retrieve == []:
        raise PreventUpdate
    if retrieve != []:
        time.sleep(1)
        df = pd.DataFrame(retrieve)
        if 'ID' and  'Value' and 'Quality' and 'Date' in df.columns:
            return [{'label': i, 'value': i} for i in df['ID'].unique()], [{'label': i, 'value': i} for i in df['ID'].unique()]
        else :
            dff = [{'label': i, 'value': i} for i in df.columns if i.startswith('Un') != 1 and i != 'index' and i != 'date']
            return (dff, dff)
    else:
        return (no_update, no_update)


@app.callback([Output("tabDropdownTopTab4", "options"), Output("tabDropdownDownTab4", "options")],
              [Input("datastore", "data")])
def dropdownlistcontrolTab4Second(retrieve):
    if retrieve == []:
        raise PreventUpdate
    if retrieve != []:
        time.sleep(1)
        df = pd.DataFrame(retrieve)
        if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
            return [{'label': i, 'value': i} for i in df['ID'].unique()], [{'label': i, 'value': i} for i in df['ID'].unique()]

        else:
            dff = [{'label': i, 'value': i} for i in df.columns if
                   i.startswith('Un') != 1 and i != 'index' and i != 'date']
            return (dff, dff)
    else:
        return (no_update, no_update)


@app.callback([Output('tabDropdownTopTab4', 'style'),
               Output('tabDropdownDownTab4', 'style'),
               Output('tabDropdownTop', 'style'),
               Output('tabDropdownDown', 'style')],
              [Input('radiographtab4', 'value')], )
def chooseradio(radio):
    if radio == None:
        raise PreventUpdate
    if radio == 'choosevalue':
        return {'visibility': 'visible'}, {'visibility': 'visible'}, {'display': 'None'}, {'display': 'None'}
    if radio == 'optionlibre':
        return {'display': 'None'}, {'display': 'None'}, {'visibility': 'visible'}, {'visibility': 'visible'},


@app.callback([Output('tab2hiddenValuex_axis', 'children'),
               Output('tab2hiddenValuey_axis', 'children')],
              [Input('tabDropdownTop', 'value'),
               Input('tabDropdownDown', 'value'),
               Input('radiographtab4', 'value')],
              )
def contractdropdown(x, y, radioval):
    if x == [] or x == None or y == None or y == []:
        raise PreventUpdate
    if radioval == 'optionlibre':
        return x, y
    else:
        return [], []


@app.callback([Output('tab4hiddenValuex_axissecond', 'children'),
               Output('tab4hiddenValuey_axissecond', 'children'),
               ],
              [Input('tabDropdownTopTab4', 'value'),
               Input('tabDropdownDownTab4', 'value'),
               Input('radiographtab4', 'value')]
              )
def contractdropdown2(valxsecond, valysecond, radio):
    if valxsecond == None or valysecond == None or radio == None:
        raise PreventUpdate

    if radio == 'choosevalue':
        return valxsecond, valysecond

    else:
        return [], []


@app.callback(
    Output('output_s', 'children'),
    [Input('tabDropdownTopTab4', 'value'),
     Input('tabDropdownTop', 'value'),
     Input('radiographtab4', 'value')], )
def container4(val2, val3, radio):
    if val2 == None and val3 == None or radio == None:
        raise PreventUpdate

    a = ''

    if radio == 'choosevalue':
        if val2 != None:
            a = val2
            return a
        else:
            return ''

    if radio == 'optionlibre':
        if val3 != None:
            a = val3
            return a
        else:
            return ''

@app.callback(
    Output('shiftaxisdroptab4', 'options'),
    [Input('tabDropdownTopTab4', 'value'),
     Input('tabDropdownTop', 'value'),
     Input('radiographtab4', 'value')], )
def container5(val2, val3, radio):
    if val2 == None and val3 == None or radio == None:
        raise PreventUpdate

    a = []

    if radio == 'choosevalue':
        if val2 != None:
            a = val2

    if radio == 'optionlibre':
        if val3 != None:
            a = val3
    return [{'label': i, 'value': i} for i in a]

@app.callback(
    [Output('firstChoosenValueTab4', 'options'),
     Output('secondChoosenValueTab4', 'options')],
    [Input('output_s', 'children'),
     Input('radiographtab4', 'value')],
    [State("datastore", "data")])
def container4_2(val, radio,data):
    if val == None or val == []:
        raise PreventUpdate
    a = []
    df = pd.DataFrame(data)
    if radio == 'choosevalue':
        a = [{'label': i, 'value': i} for i in val], [{'label': i, 'value': i} for i in val]
        return a
    elif radio == 'optionlibre':
        a = [{'label': i, 'value': i} for i in val], [{'label': i, 'value': i} for i in val]
        return a
    else : return no_update


@app.callback([Output('hiddenTextxaxis', 'children'), Output('hiddenTextyaxis', 'children'),
               Output('hiddenTextHeader', 'children'), Output('hiddenTextNote', 'children'),
               Output('hiddenTexty1axis', 'children'),Output('hiddenTexty2axis', 'children')],
              [Input('addText1', 'n_clicks')],
              [State('textarea1', 'value'), State('dropadd1', 'value'),
               State('hiddenTextxaxis', 'children'), State('hiddenTextyaxis', 'children'),
               State('hiddenTextHeader', 'children'), State('hiddenTextNote', 'children'),
               State('hiddenTexty1axis', 'children'),State('hiddenTexty2axis', 'children')]
              )
def detailedGraph(addtextclick, textarea, add, g1, g2, head, note, y1,y2):
    if add == None or g1 == None or g2 == None or head == None or note == None or y1 == None or y2== None:
        raise PreventUpdate

    if addtextclick > 0:
        if add == 'x_axis':
            g1.append(textarea)

        if add == 'y_axis':
            g2.append(textarea)

        if add == 'y2_axis':
            y1.append(textarea)

        if add == 'y3_axis':
            y2.append(textarea)

        if add == 'header':
            head.append(textarea)

        if add == 'note':
            note.append(textarea)
        textarea = ''
        return g1, g2, head, note, y1, y2
    else:
        return (no_update, no_update, no_update, no_update, no_update, no_update)


@app.callback([Output('hiddenTextxaxis4', 'children'), Output('hiddenTextyaxis4', 'children'),
               Output('hiddenTextHeader4', 'children'), Output('hiddenTextNote4', 'children'),
               Output('hiddenTexty1axis4', 'children'),Output('hiddenTexty2axis4', 'children')],
              [Input('addText4', 'n_clicks')],
              [State('textarea4', 'value'), State('dropadd4', 'value'),
               State('hiddenTextxaxis4', 'children'), State('hiddenTextyaxis4', 'children'),
               State('hiddenTextHeader4', 'children'), State('hiddenTextNote4', 'children'),
               State('hiddenTexty1axis4', 'children'),State('hiddenTexty2axis4', 'children')]
              )
def detailedGraph4(addtextclick, textarea, add, g1, g2, head, note,y1,y2):
    if add == None or g1 == None or g2 == None or head == None or note == None or y1 == None or y2== None:
        raise PreventUpdate

    if addtextclick > 0:
        if add == 'x_axis':
            g1.append(textarea)

        if add == 'y_axis':
            g2.append(textarea)

        if add == 'y2_axis':
            y1.append(textarea)

        if add == 'y3_axis':
            y2.append(textarea)

        if add == 'header':
            head.append(textarea)

        if add == 'note':
            note.append(textarea)
        textarea = ''
        return g1, g2, head, note,y1,y2
    else:
        return (no_update, no_update, no_update, no_update, no_update, no_update)

@app.callback([Output('hiddenTextxaxisdb', 'children'), Output('hiddenTextyaxisdb', 'children'),
               Output('hiddenTextHeaderdb', 'children'), Output('hiddenTextNotedb', 'children'),
               Output('hiddenTexty1axisdb', 'children'),Output('hiddenTexty2axisdb', 'children')],
              [Input('addTextdb', 'n_clicks')],
              [State('textareadb', 'value'), State('dropadddb', 'value'),
               State('hiddenTextxaxisdb', 'children'), State('hiddenTextyaxisdb', 'children'),
               State('hiddenTextHeaderdb', 'children'), State('hiddenTextNotedb', 'children'),
               State('hiddenTexty1axisdb', 'children'),State('hiddenTexty2axisdb', 'children')]
              )
def detailedGraphdb(addtextclick, textarea, add, g1, g2, head, note, y1,y2):
    if add == None or g1 == None or g2 == None or head == None or note == None or y1 == None or y2== None:
        raise PreventUpdate

    if addtextclick > 0:
        if add == 'x_axis':
            g1.append(textarea)

        if add == 'y_axis':
            g2.append(textarea)
        if add == 'y2_axis':
            y1.append(textarea)

        if add == 'y3_axis':
            y2.append(textarea)

        if add == 'header':
            head.append(textarea)

        if add == 'note':
            note.append(textarea)
        textarea = ''

        return g1, g2, head, note,y1,y2
    else:
        return (no_update, no_update, no_update, no_update, no_update, no_update)

@app.callback(Output('shiftaxistab4', 'style'),
              [Input('shiftaxisdroptab4', 'value')])
def shiftingaxestab4(val):
    if val == None:
        return {'display': 'None'}
    return {'visibility': 'visible', 'marginTop': '2rem'}


@app.callback(Output('shiftaxisdroptab4hidden', 'children'),
              [Input('shiftaxisdroptab4', 'value')], )
def relay4(val):
    return val


@app.callback(Output('shift_x_axistab4hidden', 'children'),
              [Input('shift_x_axistab4', 'value')], )
def relay5(val):
    return val


@app.callback(Output('shift_y_axistab4hidden', 'children'),
              [Input('shift_y_axistab4', 'value')], )
def relay6(val):
    return val


@app.callback(Output('radiographtab4hidden', 'children'),
              [Input('radiographtab4', 'value')], )
def relay7(valradio):
    if valradio == None:
        raise PreventUpdate
    return valradio


@app.callback(Output('graph4', 'figure'),
              [Input("selectyaxistab4", "on"),
               Input('radiograph4', 'value'),
               Input('radiographtab4hidden', 'children'),
               Input('tab4hiddenValuex_axissecond', 'children'),
               Input('tab4hiddenValuey_axissecond', 'children'),
               Input('sliderHeightTab4', 'value'),
               Input('sliderWidthTab4', 'value'),
               Input('hiddenTextxaxis4', 'children'),
               Input('hiddenTextyaxis4', 'children'),
               Input('hiddenTextHeader4', 'children'),
               Input('hiddenTextNote4', 'children'),
               Input('hiddenTexty1axis4', 'children'),
               Input('hiddenTexty2axis4', 'children'),
               Input('tab4send', 'n_clicks'),
               Input('firstChoosenValueTab4', 'value'),
               Input('secondChoosenValueTab4', 'value'),
               Input('valuechangetab4', 'n_clicks'),
               Input('tab2hiddenValuex_axis', 'children'),
               Input('tab2hiddenValuey_axis', 'children'),
               Input('cleanshapetab4', 'n_clicks'),
               Input('characteretab4', 'value'),
               ],
              [State('shiftaxisdroptab4hidden', 'children'),
               State('shift_x_axistab4hidden', 'children'),
               State('shift_y_axistab4hidden', 'children'),
               State('retrieve', 'children'),
               State('pointLeftFirstTab4', 'children'),
               State('pointRightFirstTab4', 'children'),
               State('leftIntegralFirstTab4', 'value'),
               State('leftIntegralSecondTab4', 'value'),
               State('rightIntegralFirstTab4', 'value'),
               State('rightIntegralSecondTab4', 'value'),
               ]
              )
def detailedGraph4(on,radio, radioval,  valxsecond, valysecond,
                   slideheight, slidewidth, g1, g2, head, note,y1,y2, nclick, firstchoosen, secondchoosen, nc,
                   valx2, valy2, cleanclick,car,  axisdrop, shift_x, shift_y, retrieve, firstshape, secondshape,
                   leftfirstval, leftsecondval, rightfirstval, rightsecondval, ):
    if g1 == None or g2 == None or head == None or note == None or radioval == []:
        raise PreventUpdate
    q1 = dash.callback_context.triggered[0]["prop_id"].split(".")[0]
    if radioval != None:
        if len(retrieve) > 0:
            df = pd.read_excel("appending.xlsx")
            df = df.loc[:, ~df.columns.str.match('Unnamed')]
            if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                a = df['ID'].unique()
                dff2 = pd.DataFrame([])
                for i in a:
                    dff = df[df['ID'] == i]

                    index = np.arange(0, len(dff))
                    dff.reset_index(drop=True, inplace=True)
                    dff.set_index(index, inplace=True)
                    # dff.drop(['Unnamed: 0', 'Unnamed: 0.1', 'Unnamed: 0.1.1'], axis=1, inplace=True)
                    dff = dff.pivot(values='Value', columns='ID')
                    dff2 = pd.concat([dff2, dff], axis=1)
                df = dff2.copy()
                fig = go.Figure()

            else:
                df.dropna(axis=0, inplace=True)
                fig = go.Figure()

            def controlShape():
                pathline = ''
                pathline2 = ''
                minValfirst = 0
                minValsecond = 0
                if firstchoosen != None and secondchoosen != None:
                    if len(firstshape) == 2 and leftfirstval != None and leftsecondval != None:
                        if int(firstshape[1]) > int(firstshape[0]):
                            pathline = ''
                            rangeshape = range(int(firstshape[0]), int(firstshape[1])+2)
                            if ':' or '-' in df[lst[i][0]][0]:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        pathline += 'M ' + str(df[lst[i][0]][k]) + ', ' + str(minValfirst) + ' L' + \
                                                    str(df[lst[i][0]][k]) + ', ' + str(df[firstchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(df[lst[i][0]][k]) + ', ' + str(df[firstchoosen][k])
                                pathline += ' L' + str(df[lst[i][0]][k]) + ', ' + str(minValfirst)
                                pathline += ' Z'
                            else:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        pathline += 'M ' + str(int(df[lst[i][0]][k])) + ', ' + str(minValfirst) + ' L' + \
                                                    str(int(df[lst[i][0]][k])) + ', ' + str(df[firstchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(df[firstchoosen][k])
                                pathline += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(minValfirst)
                                pathline += ' Z'

                    if len(secondshape) == 2 and rightsecondval != None and rightfirstval != None:
                        if int(secondshape[1]) > int(secondshape[0]):
                            rangeshape = range(int(secondshape[0]), int(secondshape[1])+2)
                            if ':' or '-' in df[lst[i][0]][0]:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        pathline2 += 'M ' + str(df[lst[i][0]][k]) + ', ' + str(minValsecond) + ' L' + \
                                                     str(df[lst[i][0]][k]) + ', ' + str(df[secondchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline2 += ' L' + str(df[lst[i][0]][k]) + ', ' + str(df[secondchoosen][k])
                                pathline2 += ' L' + str(df[lst[i][0]][k]) + ', ' + str(minValsecond)
                                pathline2 += ' Z'
                            else:
                                for k in rangeshape:

                                    if k == rangeshape[0]:
                                        pathline2 += 'M ' + str(int(df[lst[i][0]][k])) + ', ' + str(
                                            minValsecond) + ' L' + \
                                                     str(int(df[lst[i][0]][k])) + ', ' + str(df[secondchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline2 += ' L' + str(int(a[k])) + ', ' + str(df[secondchoosen][k])
                                pathline2 += ' L' + str(int(a[k])) + ', ' + str(minValsecond)
                                pathline2 += ' Z'

                    return [dict(
                        type="path",
                        path=pathline,
                        layer='below',
                        fillcolor="#5083C7",
                        opacity=0.8,
                        line_color="#8896BF",
                    ), dict(
                        type="path",
                        path=pathline2,
                        layer='below',
                        fillcolor="#B0384A",
                        opacity=0.8,
                        line_color="#B36873",
                    )]

                if firstchoosen != None and secondchoosen == None:
                    if len(firstshape) == 2:
                        if int(firstshape[1]) > int(firstshape[0]):
                            pathline = ''
                            rangeshape = range(int(firstshape[0]), int(firstshape[1])+2)
                            if ':' or '-' in df[lst[i][0]][0]:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        pathline += 'M ' + str(df[lst[i][0]][k]) + ', ' + str(minValfirst) + ' L' + \
                                                    str(df[lst[i][0]][k]) + ', ' + str(df[firstchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(df[lst[i][0]][k]) + ', ' + str(df[firstchoosen][k])
                                pathline += ' L' + str(df[lst[i][0]][k]) + ', ' + str(minValfirst)
                                pathline += ' Z'
                            else:
                                for k in rangeshape:

                                    if k == rangeshape[0]:
                                        pathline += 'M ' + str(int(df[lst[i][0]][k])) + ', ' + str(
                                            minValfirst) + ' L' + str(
                                            int(df[lst[i][0]][k])) + ', ' + str(
                                            df[firstchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(int(a[k])) + ', ' + str(df[firstchoosen][k])
                                pathline += ' L' + str(int(a[k])) + ', ' + str(minValfirst)
                                pathline += ' Z'

                            return [dict(
                                type="path",
                                path=pathline,
                                layer='below',
                                fillcolor="#5083C7",
                                opacity=0.8,
                                line_color="#8896BF",
                            )]

                        if int(firstshape[0]) > int(firstshape[1]):
                            rangeshape = range(int(firstshape[1]), int(firstshape[0])+2)
                            if ':' or '-' in df[lst[i][0]][0]:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        pathline += 'M ' + str(df[lst[i][0]][k]) + ', ' + str(
                                            minValsecond) + ' L' + str(
                                            df[lst[i][0]][k]) + ', ' + str(
                                            df[firstchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(df[lst[i][0]][k]) + ', ' + str(
                                            df[firstchoosen][k])
                                pathline += ' L' + str(df[lst[i][0]][k]) + ', ' + str(minValsecond)
                                pathline += ' Z'
                            else:
                                for k in rangeshape:

                                    if k == rangeshape[0]:
                                        pathline += 'M ' + str(int(df[lst[i][0]][k])) + ', ' + str(
                                            minValsecond) + ' L' + \
                                                    str(int(df[lst[i][0]][k])) + ', ' + str(df[firstchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(df[firstchoosen][k])
                                pathline += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(minValsecond)
                                pathline += ' Z'

                            return [dict(
                                type="path",
                                path=pathline,
                                layer='below',
                                fillcolor="#5083C7",
                                opacity=0.8,
                                line_color="#8896BF",
                            )]

                if secondchoosen != None and firstchoosen == None:
                    if len(secondshape) == 2 and rightsecondval != None and rightfirstval != None:
                        if int(secondshape[1]) > int(secondshape[0]):
                            rangeshape = range(int(secondshape[0]), int(secondshape[1])+2)
                            if ':' or '-' in df[lst[i][0]][0]:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        pathline2 += 'M ' + str(df[lst[i][0]][k]) + ', ' + str(minValsecond) + ' L' + \
                                                     str(df[lst[i][0]][k]) + ', ' + str(df[secondchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline2 += ' L' + str(df[lst[i][0]][k]) + ', ' + str(df[secondchoosen][k])
                                pathline2 += ' L' + str(df[lst[i][0]][k]) + ', ' + str(minValsecond)
                                pathline2 += ' Z'
                            else:
                                for k in rangeshape:

                                    if k == rangeshape[0]:
                                        pathline2 += 'M ' + str(int(df[lst[i][0]][k])) + ', ' + str(
                                            minValsecond) + ' L' + \
                                                     str(int(df[lst[i][0]][k])) + ', ' + str(df[secondchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline2 += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(
                                            df[secondchoosen][k])
                                pathline2 += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(minValsecond)
                                pathline2 += ' Z'

                            return [dict(
                                type="path",
                                path=pathline2,
                                layer='below',
                                fillcolor="#5083C7",
                                opacity=0.8,
                                line_color="#8896BF",
                            )]

                        if int(secondshape[0]) > int(secondshape[1]):
                            rangeshape = range(int(secondshape[1]), int(secondshape[0]))
                            for k in rangeshape:
                                if k == rangeshape[0]:
                                    pathline2 += 'M ' + str(df[lst[i][0]][k]) + ', ' + str(
                                        minValsecond) + ' L' + str(
                                        df[lst[i][0]][k]) + ', ' + str(
                                        df[secondchoosen][k]) + ' '

                                elif k != rangeshape[0] and k != rangeshape[-1]:
                                    pathline2 += ' L' + str(df[lst[i][0]][k]) + ', ' + str(df[secondchoosen][k])
                            pathline2 += ' L' + str(df[lst[i][0]][k]) + ', ' + str(minValsecond)
                            pathline2 += ' Z'
                        else:
                            rangeshape = range(int(secondshape[1]), int(secondshape[0])+2)
                            for k in rangeshape:

                                if k == rangeshape[0]:
                                    pathline2 += 'M ' + str(int(df[lst[i][0]][k])) + ', ' + str(minValsecond) + ' L' + \
                                                 str(int(df[lst[i][0]][k])) + ', ' + str(df[secondchoosen][k]) + ' '

                                elif k != rangeshape[0] and k != rangeshape[-1]:
                                    pathline2 += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(df[secondchoosen][k])
                            pathline2 += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(minValsecond)
                            pathline2 += ' Z'

                        return [dict(
                            type="path",
                            path=pathline2,
                            layer='below',
                            fillcolor="#5083C7",
                            opacity=0.8,
                            line_color="#8896BF",
                        )]
                else:
                    return no_update

            if len(firstshape) == 2 and leftfirstval != firstshape[0] and leftfirstval != []:
                if leftfirstval.startswith('T') == 1:
                    del firstshape[0]
                    firstshape.append(float(leftfirstval[2:]))
                    firstshape = sorted(firstshape)
                elif leftfirstval.isnumeric() == 1:
                    del firstshape[0]
                    firstshape.append(float(leftfirstval))
                    firstshape = sorted(firstshape)
                elif leftfirstval != None:
                    del firstshape[0]
            if len(firstshape) == 2 and leftsecondval != firstshape[
                1] and leftsecondval != None and leftsecondval != []:
                if leftsecondval.startswith('T') == 1:
                    del firstshape[1]
                    firstshape.append(float(leftsecondval[2:]))
                    firstshape = sorted(firstshape)
                elif leftsecondval.isnumeric() == 1:
                    del firstshape[1]
                    firstshape.append(float(leftsecondval))
                    firstshape = sorted(firstshape)
                elif leftsecondval != None:
                    del firstshape[1]

            if len(secondshape) == 2 and rightfirstval != secondshape[
                0] and rightfirstval != None and rightfirstval != []:
                if rightfirstval.startswith('T') == 1:
                    del secondshape[0]
                    secondshape.append(float(rightfirstval[2:]))
                    secondshape = sorted(secondshape)
                elif rightfirstval.isnumeric() == 1:
                    del secondshape[0]
                    secondshape.append(float(rightfirstval))
                    secondshape = sorted(secondshape)
                elif rightfirstval != None:
                    del secondshape[0]
            if len(secondshape) == 2 and rightsecondval != secondshape[
                1] and rightsecondval != None and rightsecondval != []:
                if rightsecondval.startswith('T') == 1:
                    del secondshape[1]
                    secondshape.append(float(rightsecondval[2:]))
                    secondshape = sorted(secondshape)
                elif rightsecondval.isnumeric() == 1:
                    del secondshape[1]
                    secondshape.append(float(rightsecondval))
                    secondshape = sorted(secondshape)
                elif rightsecondval != None:
                    del secondshape[1]
            if len(secondshape) == 2 and secondchoosen == None:
                del secondshape[1]
            if len(firstshape) == 2 and firstchoosen == None:
                del firstshape[1]
            if radioval == 'optionlibre' and valx2 != None and valy2 != None:
                lst = []
                for j in zip(valy2, valx2):
                    lst.append(j)
                s = -1
                m = ''
                for i in range(len(lst)):
                    if lst[i][0][-2].isdigit() == 1:
                        m = lst[i][0][-2]
                        m = 'T' + m
                    elif lst[i][0][-1].isdigit() == 1:
                        m = lst[i][0][-1]
                        m = 'T' + m
                    s += 1
                    a = df[lst[i][0]]
                    b = df[lst[i][1]]
                    if q1  == "tab4send":
                        if axisdrop in valx2:
                            p = []
                            c = []
                            for y in df[axisdrop]:
                                if shift_y == None:
                                    raise PreventUpdate
                                else:
                                    y += float(shift_y)
                                    c.append(y)
                            c.append(axisdrop)
                            df[axisdrop] = pd.DataFrame(c)
                            b = df[axisdrop]
                            df.to_excel("appending.xlsx")
                    for j in range(len(valy2)):
                        for k in range(len(valx2)):
                            a = df[valy2[j]]
                            b = df[valx2[k]]
                            if on == 1:
                                if len(valx2) == 1:
                                    fig.add_trace(go.Scattergl(x=a, y=b, mode=radio, marker=dict(line=dict(width=0.2, color='white')),
                                    name="{}".format(valx2[k])))
                                elif len(valx2) > 1:
                                    fig.add_trace(go.Scattergl(x=a, y=b, mode=radio,
                                                               marker=dict(line=dict(width=0.2, color='white')),
                                                               name="{}".format(valx2[k]),yaxis = f'y{k+1}'))
                            else :
                                fig.add_trace(go.Scattergl(x=a, y=b, mode=radio, marker=dict(line=dict(width=0.2, color='white')),name="{}".format(valx2[k])))

                            a = []
                            if nc > 0:
                                a = controlShape()
                            fig.update_xaxes(
                                tickfont_size=car,
                                tickangle=90,
                                title_text='' if g1 == [] else g1[-1],
                                title_standoff=25, title_font={"size": car}, ),

                            fig.update_yaxes(
                                tickfont_size=car,
                                title_text='' if g2 == [] else g2[-1],
                                title_standoff=25,
                                titlefont=dict(
                                    color="#636EFA",
                                    size=car
                                ),
                                tickfont=dict(
                                    color="#636EFA"
                                )),

                            fig.update_layout(
                                title_text=head[-1] if len(head) > 0 else "{}".format(valy2[0]),
                                title_font={"size": car},
                                width=slidewidth,
                                height=slideheight,
                                shapes=a if nc > cleanclick else [],
                                legend=dict(

                                    traceorder="normal",
                                    font=dict(
                                        family="sans-serif",
                                        size=car,
                                        color=colors['figure_text']
                                    ),
                                    yanchor="bottom",
                                    y=-0.4,
                                    xanchor="right",
                                    x=1,
                                    bgcolor=colors['background'],
                                    borderwidth=1,
                                    orientation='h',
                                ),
                                yaxis2=dict(
                                    title_text='y2_axis' if y1 == [] else y1[-1],
                                    titlefont=dict(
                                        color="#EF553B",
                                        size=car
                                    ),
                                    tickfont=dict(
                                        color="#EF553B"
                                    ),
                                    title_standoff=10,
                                    anchor="free",
                                    overlaying="y",
                                    side="right",
                                    position=1

                                ),

                                yaxis3=dict(
                                    title_text='y3_axis' if y2 == [] else y2[-1],
                                    titlefont=dict(
                                        color="#00CC96",
                                        size=car
                                    ),
                                    tickfont=dict(
                                        color="#00CC96",
                                    ),
                                    title_standoff=10,
                                    anchor="free",
                                    overlaying="y",
                                    side="right",
                                    position=0.87
                                ),

                                margin=dict(
                                    l=50,
                                    r=50,
                                    b=100,
                                    t=50,
                                    pad=4

                                ),
                                paper_bgcolor="LightSteelBlue",
                                plot_bgcolor=colors['background'], ),
                            note_new = []
                            if note != None and len(note) > 0:
                                note = note[-1].split(' ')
                            x = ''
                            for i in note:
                                note_new.append(i)
                                if len(note_new) % 4 == 0:
                                    note_new[-1] += '<br>'
                            x += ' '.join(note_new)
                            note = [x]
                            fig.add_annotation(text=note[-1] if len(note) > 0 else '', align='left',
                                               bordercolor='black',
                                               borderwidth=1, xref="paper", yref="paper", bgcolor="white", font=dict(
                                    family="Arial, sans-serif", size=car, ),
                                               x=1, y=1, showarrow=False)
                            fig.update_layout()

                    return fig

            if radioval == 'choosevalue' and len(valxsecond) > 0 and len(valysecond) > 0:
                lst = []
                for j in zip(valysecond, valxsecond):
                    lst.append(j)
                s = -1
                m = ''
                for i in range(len(lst)):
                    if lst[i][0][-2].isdigit() == 1:
                        m = lst[i][0][-2]
                        m = 'T' + m
                    elif lst[i][0][-1].isdigit() == 1:
                        m = lst[i][0][-1]
                        m = 'T' + m
                    s += 1
                    a = df[lst[i][0]]
                    b = df[lst[i][1]]

                    if q1  == "tab4send":
                        if axisdrop == lst[i][1]:
                            p = []
                            c = []
                            for t in df[lst[i][0]]:
                                if shift_x == None:
                                    raise PreventUpdate
                                else:
                                    t += float(shift_x)
                                    p.append(t)
                            df[lst[i][0]] = pd.DataFrame(p)
                            a = df[lst[i][0]]
                            df.to_excel("appending.xlsx")
                            for y in df[axisdrop]:
                                if shift_y == None:
                                    raise PreventUpdate
                                else:
                                    y += float(shift_y)
                                    c.append(y)
                            c.append(axisdrop)
                            df[axisdrop] = pd.DataFrame(c)
                            b = df[axisdrop]
                            df.to_excel("appending.xlsx")

                    if on == 1:

                        if len(valxsecond) == 1:
                            fig.add_trace(go.Scattergl(x=a, y=b, mode=radio, marker=dict(line=dict(width=0.2, color='white')),
                                               name="{}".format(valxsecond[s])))
                        if len(valxsecond) > 1:
                            fig.add_trace(
                                go.Scattergl(x=a, y=b, mode=radio, marker=dict(line=dict(width=0.2, color='white')),
                                             name="{}".format(valxsecond[s]), yaxis = f'y{s+1}'))

                    else :
                        fig.add_trace(
                            go.Scattergl(x=a, y=b, mode=radio, marker=dict(line=dict(width=0.2, color='white')),
                                         name="{}".format(valxsecond[s])))

                    a = []
                    if nc > 0:
                        a = controlShape()
                    fig.update_xaxes(
                        tickfont_size=car,
                        tickangle=90,
                        title_text='' if g1 == [] else g1[-1],
                        title_standoff=25, title_font={"size": car}, ),

                    fig.update_yaxes(
                        tickfont_size=car,
                        title_text='' if g2 == [] else g2[-1],
                        title_standoff=25,
                        titlefont=dict(
                            color="#636EFA",
                            size=car
                        ),
                        tickfont=dict(
                            color="#636EFA"
                        )),
                    fig.update_shapes(yref='y'),
                    fig.update_layout(
                        title_text=head[-1] if len(head) > 0 else "{}/{}".format(valxsecond[0], valysecond[0]),
                        title_font={"size": car},
                        width=slidewidth,
                        height=slideheight,
                        shapes=a if nc > cleanclick else [],
                        legend=dict(

                            traceorder="normal",
                            font=dict(
                                family="sans-serif",
                                size=car,
                                color=colors['figure_text']
                            ),
                            yanchor="bottom",
                            y=-0.4,
                            xanchor="right",
                            x=1,
                            bgcolor=colors['background'],
                            borderwidth=1,
                            orientation='h',
                        ),
                        yaxis2=dict(
                            title_text='y2_axis' if y1 == [] else y1[-1],
                            titlefont=dict(
                                color="#EF553B",
                                size=car
                            ),
                            tickfont=dict(
                                color="#EF553B"
                            ),
                            title_standoff=10,
                            anchor="free",
                            overlaying="y",
                            side="right",
                            position=1

                        ),

                        yaxis3=dict(
                            title_text='y3_axis' if y2 == [] else y2[-1],
                            titlefont=dict(
                                color="#00CC96",
                                size=car
                            ),
                            tickfont=dict(
                                color="#00CC96",
                            ),
                            title_standoff=10,
                            anchor="free",
                            overlaying="y",
                            side="right",
                            position=0.87
                        ),

                        margin=dict(
                            l=50,
                            r=50,
                            b=100,
                            t=50,
                            pad=4

                        ),
                        paper_bgcolor="LightSteelBlue",
                        plot_bgcolor=colors['background'], ),
                    note_new = []
                    if note != None and len(note) > 0:
                        note = note[-1].split(' ')
                    x = ''
                    for i in note:
                        note_new.append(i)
                        if len(note_new) % 4 == 0:
                            note_new[-1] += '<br>'
                    x += ' '.join(note_new)
                    note = [x]
                    fig.add_annotation(text=note[-1] if len(note) > 0 else '', align='left', bordercolor='black',
                                       borderwidth=1, xref="paper", yref="paper", bgcolor="white", font=dict(
                            family="Arial, sans-serif", size=car, ),
                                       x=1, y=1, showarrow=False)
                    fig.update_layout()

                return fig
            else:
                return no_update
        else:
            return no_update


@app.callback(
    [Output('pointLeftFirst', 'children'),
     Output('pointLeftSecond', 'children')],
    [Input('graph', 'clickData'),
     Input('firstChoosenValue', 'value'), ],
    [State('leftSideChecklistValueHidden', 'children'),
     State('pointLeftFirst', 'children'),
     State('pointLeftSecond', 'children'),
     State('shift_x_axis', 'value'),
     State('checklistvaleurhidden2', "children"),
     State('retrieve', 'children'),
     ]
)
def valint(clickData, firstchoosen, value, leftchild, rightchild, shift_x,rightside, retrieve):
    if value is [] or value is None or clickData == None or clickData == [] or firstchoosen == None or rightside == None or retrieve == None or retrieve == []:
        raise PreventUpdate

    spaceList1 = []
    zero = 0
    spaceList2 = []
    if len(retrieve) > 0:
        df = pd.read_excel('appending.xlsx')
        df = df.loc[:, ~df.columns.str.match('Unnamed')]
        df['index'] = df.index

        for i in range(len(rightside)):
            spaceList1.append(zero)
            zero += 1
            spaceList2.append(rightside[i])

        for i in range(len(value)):
            spaceList1.append(zero)
            zero += 1
            spaceList2.append(value[i])
        zippedval = [i for i in list(zip(spaceList1, spaceList2))]
        curvenumber = clickData['points'][0]['curveNumber']
        for k in zippedval:
            if k[1] == firstchoosen:
                if k[0] == curvenumber:
                    x_val = clickData['points'][0]['x']
                    dff = pd.DataFrame([])
                    if 'date' in df.columns:
                        dff = df[df['date'] == x_val]
                    elif 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                        dff = df.loc[df['ID'] == firstchoosen]
                        dff = dff.copy()
                        index = np.arange(0, len(dff['ID']))
                        dff.reset_index(drop=True, inplace=True)
                        dff.set_index(index, inplace=True)
                        if x_val[-3] == '.':
                            x_val = x_val + '0000+00:00'
                        elif x_val[-1] == '.':
                            x_val = x_val + '00000+00:00'
                        else : x_val += '000+00:00'
                        dff = dff[dff['Date'] == x_val]
                    else:
                        a = ''
                        for v in df.columns:
                            if 'Temps' in v:
                                a += v
                                dff = df[df[v] == x_val]
                                if shift_x != 0:
                                    x_val -= shift_x
                                    dff = df[df[v] == x_val]
                    a = []
                    if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                        a.append([dff[dff['ID'] == firstchoosen].index][0])
                    else : a.append(dff[firstchoosen].index)
                    for i in range(len(a)):
                        for j in a:
                            leftchild.append(j[i])
                    if len(leftchild) > 2:
                        leftchild.pop(0)
                    return (leftchild, leftchild)
                else:
                    return (no_update, no_update)
            # else : return (no_update, no_update)
    else:
        return (no_update, no_update)

@app.callback(
    [Output('pointLeftFirstdb', 'children'),
     Output('pointLeftSeconddb', 'children')],
    [Input('getdbgraph', 'clickData'),
     Input('firstChoosenValuedb', 'value'), ],
    [State('dbvalname', 'value'), State('pointLeftFirstdb', 'children'),
     State('pointLeftSeconddb', 'children'),
     State('memory-output', 'data'),
     State('dbvalchoosen', 'value'), State('db_name', 'value')
     ]
)
def valintdb(clickData, firstchoosen, value, leftchild, rightchild, retrieve, dbch, dbname):
    if value == [] or value == None or clickData == None or clickData == [] or \
            retrieve == None or retrieve == []:
        raise PreventUpdate
    spaceList1 = []
    zero = 0
    spaceList2 = []
    if retrieve != None:
        df = pd.DataFrame(retrieve)
        if dbname == 'rcckn':
            if dbch == 'send_variablevalues':
                df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'TIMESTAMP',
                              'PROCESSED', 'TIMED_OUT', 'UNREFERENCED']
            if dbch == 'received_variablevalues':
                df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'LOCAL_TIMESTAMP',
                              'REMOTE_ID', 'REMOTE_TIMESTAMP', 'REMOTE_MESSAGE_ID', 'PROCESSED', 'TIMED_OUT',
                              'CONVERTED_NUM_VALUE']
        if dbname == 'enerbat':
            df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
        for i in range(len(value)):
            spaceList1.append(zero)
            zero += 1
            spaceList2.append(value[i])
        zippedval = [i for i in list(zip(spaceList1, spaceList2))]
        curvenumber = clickData['points'][0]['curveNumber']
        for k in zippedval:
            if k[1] == firstchoosen:
                if k[0] == curvenumber:
                    x_val = clickData['points'][0]['x']
                    x_val = x_val[:10] + 'T' + x_val[11:]
                    if len(x_val) <=16 :
                        x_val += ':00'

                    dff = df[df['VARIABLE_NAME'] == firstchoosen]
                    if dbch == 'send_variablevalues':
                        dff = dff[dff.TIMESTAMP.str.startswith(x_val[:10])]
                        index = np.arange(0, len(dff['VARIABLE_NAME']))
                        dff.reset_index(drop=True, inplace=True)
                        dff.set_index(index, inplace=True)
                        dff = dff[(dff['TIMESTAMP'] == x_val)]
                    if dbch == 'received_variablevalues':
                        dff = dff[dff.REMOTE_TIMESTAMP.str.startswith(x_val[:10])]
                        index = np.arange(0, len(dff['VARIABLE_NAME']))
                        dff.reset_index(drop=True, inplace=True)
                        dff.set_index(index, inplace=True)
                        dff = dff[(dff['REMOTE_TIMESTAMP'] == x_val)]
                    elif dbch != 'received_variablevalues' and dbch != 'send_variablevalues':
                        dff = dff[dff.TIMESTAMP.str.startswith(x_val[:10])]
                        index = np.arange(0, len(dff['VARIABLE_NAME']))
                        dff.reset_index(drop=True, inplace=True)
                        dff.set_index(index, inplace=True)
                        dff = dff[(dff['TIMESTAMP'] == x_val)]
                    a = []
                    a.append(dff.index)

                    for i in range(len(a)):
                        for j in a:
                            leftchild.append(j[i])
                    if len(leftchild) > 2:
                        leftchild.pop(0)
                    return (leftchild, leftchild)
                    # else: return (no_update, no_update)
                else:
                    return (no_update, no_update)
        else:
            return (no_update, no_update)


@app.callback([Output('leftIntegralFirst', 'value'), Output('leftIntegralSecond', 'value')],
              [Input('pointLeftFirst', 'children'), Input('pointLeftSecond', 'children'),Input('firstChoosenValue', 'value')],
               )
def display_hover_data(leftchild, rightchild, firstchoosen):
    if leftchild == None or rightchild == None or leftchild == [] or rightchild == []:
        raise PreventUpdate

    minchild = 0
    maxchild = 0
    if firstchoosen != None and len(leftchild) == 2:
        for i in range(len(leftchild)):
            if leftchild[0] < leftchild[1]:
                minchild = leftchild[0]
                maxchild = leftchild[1]
            else:
                minchild = leftchild[1]
                maxchild = leftchild[0]
        return ('T ' + str(minchild), 'T ' + str(maxchild))
    elif firstchoosen == None :
        return '',''
    else:
        return (no_update, no_update)

@app.callback([Output('leftIntegralFirstdb', 'value'), Output('leftIntegralSeconddb', 'value')],
              [Input('pointLeftFirstdb', 'children'), Input('pointLeftSeconddb', 'children'), Input('firstChoosenValuedb', 'value')],
              )
def display_hover_data_db1(leftchild, rightchild,firstchoosen):
    if leftchild == None or rightchild == None or leftchild == [] or rightchild == []:
        raise PreventUpdate
    minchild = 0
    maxchild = 0
    if firstchoosen != None and len(leftchild) == 2:
        for i in range(len(leftchild)):
            if leftchild[0] < leftchild[1]:
                minchild = leftchild[0]
                maxchild = leftchild[1]
            else:
                minchild = leftchild[1]
                maxchild = leftchild[0]
        return ('T ' + str(minchild), 'T ' + str(maxchild))
    elif firstchoosen == None :
        return '',''
    else:
        return (no_update, no_update)

@app.callback(
    [Output('pointLeftFirstTab4', 'children'),
     Output('pointLeftSecondTab4', 'children')],
    [Input('graph4', 'clickData'),
     Input('radiographtab4hidden', 'children'),
     Input('firstChoosenValueTab4', 'value'),
     ],  # describe variable of shift
    [State('tab4hiddenValuey_axissecond', 'children'),
     State('tab4hiddenValuex_axissecond', 'children'),
     State('tab2hiddenValuey_axis', 'children'),
     State('tab2hiddenValuex_axis', 'children'),
     State('pointLeftFirstTab4', 'children'),
     State('pointLeftSecondTab4', 'children'),
     State('retrieve', 'children'),
     State('shift_x_axistab4', 'value'),  # shifting value of x_axis
     State('output_s', 'children')  # it takes values of tabdropdowntop and topdropdowntoptab4
     ]
)
def valintTab4(clickData4, radioval, firstchoosen, valysecond, valxsecond, valy, valx, leftchild, rightchild, retrieve,
               shift_x, container):
    if clickData4 == None or clickData4 == [] or firstchoosen == None or firstchoosen == [] or retrieve == None or retrieve == []:
        raise PreventUpdate
    spaceList1 = []
    zero = 0
    spaceList2 = []
    if len(retrieve) > 0:
        df = pd.read_excel('appending.xlsx')
        df = df.loc[:, ~df.columns.str.match('Unnamed')]
        if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
            a = df['ID'].unique()
            dff2 = pd.DataFrame([])
            for i in a:
                dff = df[df['ID'] == i]
                index = np.arange(0, len(dff))
                dff.reset_index(drop=True, inplace=True)
                dff.set_index(index, inplace=True)
                # dff.drop(['Unnamed: 0', 'Unnamed: 0.1', 'Unnamed: 0.1.1'], axis=1, inplace=True)
                dff = dff.pivot(values='Value', columns='ID')
                dff2 = pd.concat([dff2, dff], axis=1)
            df = dff2.copy()

        else :
            df['index'] = df.index
            df.dropna(axis=0, inplace=True)
        for i in range(len(container)):
            spaceList1.append(zero)
            zero += 1
            spaceList2.append(container[i])
        zippedval = [i for i in list(zip(spaceList1, spaceList2))]
        curvenumber = clickData4['points'][0]['curveNumber']
        for k in zippedval:
            if k[1] == firstchoosen:
                if k[0] == curvenumber:
                    if radioval == 'choosevalue':
                        if firstchoosen[-1].isdigit() == 1 and firstchoosen[:2].startswith('Tb') != 1:
                            if valxsecond != []:
                                t = valxsecond.index(firstchoosen)
                                m = valysecond[t]
                                x_val = clickData4['points'][0]['x']
                                dff = df[df[m] == x_val]
                                a = []
                                a.append(dff[firstchoosen].index)
                                for i in range(len(a)):
                                    for j in a:
                                        leftchild.append(j[i])

                                if len(leftchild) > 2:
                                    leftchild.pop(0)
                                return (leftchild, leftchild)
                            else:
                                m = firstchoosen[-1:]
                                m = 'T' + m
                                x_val = clickData4['points'][0]['x']
                                dff = df[df[m] == x_val]
                                a = []
                                a.append(dff[firstchoosen].index)
                                for i in range(len(a)):
                                    for j in a:
                                        leftchild.append(j[i])

                                if len(leftchild) > 2:
                                    leftchild.pop(0)
                                return (leftchild, leftchild)
                        elif firstchoosen[-2].isdigit() == 1:
                            if valxsecond != []:
                                t = valxsecond.index(firstchoosen)
                                m = valysecond[t]
                                x_val = clickData4['points'][0]['x']
                                dff = df[df[m] == x_val]
                                a = []
                                a.append(dff[firstchoosen].index)
                                for i in range(len(a)):
                                    for j in a:
                                        leftchild.append(j[i])

                                if len(leftchild) > 2:
                                    leftchild.pop(0)
                                return (leftchild, leftchild)
                            else:
                                m = firstchoosen[-2:]
                                m = 'T' + m
                                x_val = clickData4['points'][0]['x']
                                dff = df[df[m] == x_val]
                                a = []
                                a.append(dff[firstchoosen].index)
                                for i in range(len(a)):
                                    for j in a:
                                        leftchild.append(j[i])
                                if len(leftchild) > 2:
                                    leftchild.pop(0)
                                return (leftchild, leftchild)
                        else:
                            if valxsecond != []:
                                if firstchoosen in valxsecond :
                                    t = valxsecond.index(firstchoosen)
                                    m = valysecond[t]
                                    x_val = clickData4['points'][0]['x']
                                    dff = df[df[m] == x_val]
                                    a = []
                                    a.append(dff[firstchoosen].index)
                                    for i in range(len(a)):
                                        for j in a:
                                            leftchild.append(j[i])

                                    if len(leftchild) > 2:
                                        leftchild.pop(0)
                                    return (leftchild, leftchild)
                                else:
                                    return (no_update, no_update)
                            else:
                                return (no_update, no_update)

                    elif radioval == 'optionlibre':
                        if valx != []:
                            t = valx.index(firstchoosen)
                            m = valy[t]
                            x_val = clickData4['points'][0]['x']
                            dff = df[df[m] == x_val]
                            a = []
                            a.append(dff[valx].index)
                            for i in range(len(a)):
                                for j in a:
                                    leftchild.append(j[i])

                            if len(leftchild) > 2:
                                leftchild.pop(0)

                            return (leftchild, leftchild)
                        else:
                            return (no_update, no_update)

                    else:
                        return (no_update, no_update)
                else:
                    return (no_update, no_update)
            # else:
            #     return (no_update, no_update)
    else:
        return (no_update, no_update)

@app.callback([Output('leftIntegralFirstTab4', 'value'),
               Output('leftIntegralSecondTab4', 'value')],
              [Input('pointLeftFirstTab4', 'children'),
               Input('pointLeftSecondTab4', 'children'),
               Input('firstChoosenValueTab4', 'value'),
               Input('radiographtab4', 'value')], )
def display_hover_dataTab4(leftchild, rightchild, firstchoosen, radioval):
    if leftchild == None or rightchild == None or firstchoosen == None:
        raise PreventUpdate
    if radioval == 'optionlibre' :
        if firstchoosen != None and len(leftchild) == 2:
            for i in range(len(leftchild)):
                if leftchild[0] < leftchild[1]:
                    minchild = leftchild[0]
                    maxchild = leftchild[1]
                else:
                    minchild = leftchild[1]
                    maxchild = leftchild[0]
            return ('T ' + str(minchild), 'T ' + str(maxchild))
        elif firstchoosen == None :
            return '',''
        else:
            return no_update, no_update
    if radioval == 'choosevalue' :
        if firstchoosen != None and len(leftchild) == 2:
            for i in range(len(leftchild)):
                if leftchild[0] < leftchild[1]:
                    minchild = leftchild[0]
                    maxchild = leftchild[1]
                else:
                    minchild = leftchild[1]
                    maxchild = leftchild[0]
            return ('T ' + str(minchild), 'T ' + str(maxchild))
        elif firstchoosen == None :
            return '',''
        else:
            return no_update, no_update
    else:
        return no_update, no_update


@app.callback(
    [Output('pointRightFirst', 'children'),
     Output('pointRightSecond', 'children')],
    [Input('graph', 'clickData'),
     Input('secondChoosenValue', 'value')],
    [State('leftSideChecklistValueHidden', 'children'),
     State('pointRightFirst', 'children'),
     State('pointRightSecond', 'children'),
     State('shift_x_axis', 'value'),
     State('checklistvaleurhidden2', "children"),
     State('retrieve', 'children')]
)
def valint2(clickData, secondchoosen, value, leftchild, rightchild, shift_x,rightside, retrieve):
    if value is [] or value is None or clickData == None or secondchoosen == None or rightside == None or retrieve == None or retrieve == []:
        raise PreventUpdate

    spaceList1 = []
    zero = 0
    spaceList2 = []
    if len(retrieve) > 0:
        df = pd.read_excel('appending.xlsx')
        df = df.loc[:, ~df.columns.str.match('Unnamed')]
        df['index'] = df.index
        for i in range(len(rightside)):
            spaceList1.append(zero)
            zero += 1
            spaceList2.append(rightside[i])
        for i in range(len(value)):
            spaceList1.append(zero)
            zero += 1
            spaceList2.append(value[i])
        zippedval = [i for i in list(zip(spaceList1, spaceList2))]
        curvenumber = clickData['points'][0]['curveNumber']
        for k in zippedval:
            if k[1] == secondchoosen:
                if k[0] == curvenumber:
                    x_val = clickData['points'][0]['x']
                    if 'date' in df.columns:
                        dff = df[df['date'] == x_val]
                    elif 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                        dff = df.loc[df['ID'] == secondchoosen]
                        dff = dff.copy()
                        index = np.arange(0, len(dff['ID']))
                        dff.reset_index(drop=True, inplace=True)
                        dff.set_index(index, inplace=True)
                        if x_val[-3] == '.':
                            x_val = x_val + '0000+00:00'
                        elif x_val[-1] == '.':
                            x_val = x_val + '00000+00:00'
                        else:
                            x_val += '000+00:00'
                        dff = dff[dff['Date'] == x_val]

                    else:
                        a = ''
                        for v in df.columns:
                            if 'Temps' in v:
                                a += v
                                dff = df[df[v] == x_val]
                                if shift_x != 0:
                                    x_val -= shift_x
                                    dff = df[df[v] == x_val]
                    a = []
                    if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                        a.append([dff[dff['ID'] == secondchoosen].index][0])
                    else : a.append(dff[secondchoosen].index)
                    for i in range(len(a)):
                        for j in a:
                            leftchild.append(j[i])
                    if len(leftchild) > 2:
                        leftchild.pop(0)
                    return (leftchild, leftchild)
                else:
                    return (no_update, no_update)
            # else : return (no_update, no_update)
    else:
        return (no_update, no_update)

@app.callback(
    [Output('pointRightFirstdb', 'children'),
     Output('pointRightSeconddb', 'children')],
    [Input('getdbgraph', 'clickData'),
     Input('secondChoosenValuedb', 'value'), ],
    [State('dbvalname', 'value'), State('pointRightFirstdb', 'children'),
     State('pointRightSeconddb', 'children'),
     State('memory-output', 'data'),
     State('dbvalchoosen', 'value'), State('db_name', 'value')
     ]
)
def valintdb2(clickData, secondchoosen, value, leftchild, rightchild, retrieve, dbch, dbname):
    if value == [] or value == None or clickData == None or clickData == [] or secondchoosen == None or \
            retrieve == None or retrieve == []:
        raise PreventUpdate
    spaceList1 = []
    zero = 0
    spaceList2 = []
    if retrieve != []:
        df = pd.DataFrame(retrieve)
        if dbname == 'rcckn':
            if dbch == 'send_variablevalues':
                df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'TIMESTAMP',
                              'PROCESSED', 'TIMED_OUT', 'UNREFERENCED']
            if dbch == 'received_variablevalues':
                df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'LOCAL_TIMESTAMP',
                              'REMOTE_ID', 'REMOTE_TIMESTAMP', 'REMOTE_MESSAGE_ID', 'PROCESSED', 'TIMED_OUT',
                              'CONVERTED_NUM_VALUE']
        if dbname == 'enerbat':
            df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']

        df['index'] = df.index

        for i in range(len(value)):
            spaceList1.append(zero)
            zero += 1
            spaceList2.append(value[i])
        zippedval = [i for i in list(zip(spaceList1, spaceList2))]
        curvenumber = clickData['points'][0]['curveNumber']
        for k in zippedval:
            if k[1] == secondchoosen:
                if k[0] == curvenumber:
                    x_val = clickData['points'][0]['x']
                    x_val = x_val[:10] + 'T' + x_val[11:]
                    if len(x_val) <=16 :
                        x_val += ':00'
                    dff = df[df['VARIABLE_NAME'] == secondchoosen]
                    if dbch == 'send_variablevalues':
                        dff = dff[dff.TIMESTAMP.str.startswith(x_val[:10])]
                        index = np.arange(0, len(dff['VARIABLE_NAME']))
                        dff.reset_index(drop=True, inplace=True)
                        dff.set_index(index, inplace=True)
                        dff = dff[(dff['TIMESTAMP'] == x_val)]
                    if dbch == 'received_variablevalues':
                        dff = dff[dff.REMOTE_TIMESTAMP.str.startswith(x_val[:10])]
                        index = np.arange(0, len(dff['VARIABLE_NAME']))
                        dff.reset_index(drop=True, inplace=True)
                        dff.set_index(index, inplace=True)
                        dff = dff[(dff['REMOTE_TIMESTAMP'] == x_val)]
                    elif dbch != 'received_variablevalues' and dbch != 'send_variablevalues':
                        dff = dff[dff.TIMESTAMP.str.startswith(x_val[:10])]
                        index = np.arange(0, len(dff['VARIABLE_NAME']))
                        dff.reset_index(drop=True, inplace=True)
                        dff.set_index(index, inplace=True)
                        dff = dff[(dff['TIMESTAMP'] == x_val)]
                    a = []
                    a.append(dff.index)
                    for i in range(len(a)):
                        for j in a:
                            leftchild.append(j[i])
                    if len(leftchild) > 2:
                        leftchild.pop(0)
                    return (leftchild, leftchild)
                    # else: return (no_update, no_update)
                else:
                    return (no_update, no_update)
        else:
            return (no_update, no_update)


@app.callback([Output('rightIntegralFirstdb', 'value'), Output('rightIntegralSeconddb', 'value')],
              [Input('pointRightFirstdb', 'children'), Input('pointRightSeconddb', 'children'),Input('secondChoosenValuedb', 'value')],
              )
def display_hover_data_db2(leftchild, rightchild,secondchoosen):
    if leftchild == None or rightchild == None or leftchild == [] or rightchild == []:
        raise PreventUpdate
    minchild = 0
    maxchild = 0
    if secondchoosen != None and len(leftchild) == 2:
        for i in range(len(leftchild)):
            if leftchild[0] < leftchild[1]:
                minchild = leftchild[0]
                maxchild = leftchild[1]
            else:
                minchild = leftchild[1]
                maxchild = leftchild[0]
        return ('T ' + str(minchild), 'T ' + str(maxchild))
    elif secondchoosen == None :
        return '',''
    else:
        return (no_update, no_update)

@app.callback(
    [Output('rightIntegralFirst', 'value'), Output('rightIntegralSecond', 'value')],
    [Input('pointRightFirst', 'children'), Input('pointRightSecond', 'children'),
    Input('secondChoosenValue', 'value')], )
def display_hover_data2(leftchild, rightchild1, secondchoosen):
    minchild = 0
    maxchild = 0
    if secondchoosen != None and len(leftchild) == 2:

        for i in range(len(leftchild)):
            if leftchild[0] < leftchild[1]:
                minchild = leftchild[0]
                maxchild = leftchild[1]
            else:
                minchild = leftchild[1]
                maxchild = leftchild[0]
        return ('T ' + str(minchild), 'T ' + str(maxchild))
    elif secondchoosen == None :
        return '',''
    else:
        return (no_update, no_update)


@app.callback(
    [Output('pointRightFirstTab4', 'children'),
     Output('pointRightSecondTab4', 'children')],
    [Input('graph4', 'clickData'),
     Input('radiographtab4hidden', 'children'),
     Input('secondChoosenValueTab4', 'value'),
     ],
    [State('tab4hiddenValuey_axissecond', 'children'),
     State('tab4hiddenValuex_axissecond', 'children'),
     State('tab2hiddenValuey_axis', 'children'),
     State('tab2hiddenValuex_axis', 'children'),
     State('pointRightFirstTab4', 'children'),
     State('pointRightSecondTab4', 'children'),
     State('retrieve', 'children'),
     State('output_s', 'children'),
     State('shift_x_axistab4', 'value'), ]
)
def valintTab4_2(clickData, radioval, secondchoosen, valysecond, valxsecond, valy, valx, leftchild, rightchild,
                 retrieve, container, shift_x):
    if clickData == None or container is [] or container is None or secondchoosen == None or secondchoosen == [] or retrieve == None or retrieve == []:
        raise PreventUpdate

    spaceList1 = []
    zero = 0
    spaceList2 = []
    if len(retrieve) > 0:
        df = pd.read_excel('appending.xlsx')
        df = df.loc[:, ~df.columns.str.match('Unnamed')]
        if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
            a = df['ID'].unique()
            dff2 = pd.DataFrame([])
            for i in a:
                dff = df[df['ID'] == i]
                index = np.arange(0, len(dff))
                dff.reset_index(drop=True, inplace=True)
                dff.set_index(index, inplace=True)
                dff = dff.pivot(values='Value', columns='ID')
                dff2 = pd.concat([dff2, dff], axis=1)
            df = dff2.copy()
        else :
            df['index'] = df.index
        for i in range(len(container)):
            spaceList1.append(zero)
            zero += 1
            spaceList2.append(container[i])
        zippedval = [i for i in list(zip(spaceList1, spaceList2))]
        curvenumber = clickData['points'][0]['curveNumber']
        for k in zippedval:
            if k[1] == secondchoosen:
                if k[0] == curvenumber:
                    if radioval == "choosevalue":
                        if secondchoosen[-1].isdigit() == 1 and secondchoosen[:2].startswith('Tb') !=1:
                            if valxsecond != []:
                                t = valxsecond.index(secondchoosen)
                                m = valysecond[t]
                                x_val = clickData['points'][0]['x']
                                dff = df[df[m] == x_val]
                                a = []
                                a.append(dff[secondchoosen].index)
                                for i in range(len(a)):
                                    for j in a:
                                        leftchild.append(j[i])
                                    if len(leftchild) > 2:
                                        leftchild.pop(0)
                                return (leftchild, leftchild)
                            else:
                                m = secondchoosen[-1:]
                                m = 'T' + m
                                x_val = clickData['points'][0]['x']
                                dff = df[df[m] == x_val]
                                a = []
                                a.append(dff[secondchoosen].index)
                                for i in range(len(a)):
                                    for j in a:
                                        leftchild.append(j[i])
                                if len(leftchild) > 2:
                                    leftchild.pop(0)
                                return (leftchild, leftchild)
                        elif secondchoosen[-2].isdigit() == 1:
                            if valxsecond != []:
                                t = valxsecond.index(secondchoosen)
                                m = valysecond[t]
                                x_val = clickData['points'][0]['x']
                                dff = df[df[m] == x_val]
                                a = []
                                a.append(dff[secondchoosen].index)
                                for i in range(len(a)):
                                    for j in a:
                                        leftchild.append(j[i])
                                if len(leftchild) > 2:
                                    leftchild.pop(0)
                                return (leftchild, leftchild)
                            else:
                                m = secondchoosen[-2:]
                                m = 'T' + m
                                x_val = clickData['points'][0]['x']
                                dff = df[df[m] == x_val]
                                a = []
                                a.append(dff[secondchoosen].index)
                                for i in range(len(a)):
                                    for j in a:
                                        leftchild.append(j[i])
                                if len(leftchild) > 2:
                                    leftchild.pop(0)
                                return (leftchild, leftchild)
                        else:
                            if valxsecond != []:
                                t = valxsecond.index(secondchoosen)
                                m = valysecond[t]
                                x_val = clickData['points'][0]['x']
                                dff = df[df[m] == x_val]
                                a = []
                                a.append(dff[secondchoosen].index)
                                for i in range(len(a)):
                                    for j in a:
                                        leftchild.append(j[i])

                                if len(leftchild) > 2:
                                    leftchild.pop(0)
                                return (leftchild, leftchild)
                            else:
                                return no_update, no_update
                    elif radioval == 'optionlibre':
                        if valx != []:
                            t = valx.index(secondchoosen)
                            m = valy[0]
                            x_val = clickData['points'][0]['x']
                            dff = df[df[m] == x_val]
                            a = []
                            a.append(dff[secondchoosen].index)
                            for i in range(len(a)):
                                for j in a:
                                    leftchild.append(j[i])

                            if len(leftchild) > 2:
                                leftchild.pop(0)
                            return (leftchild, leftchild)
                        else:
                            return (no_update, no_update)

                    else:
                        return (no_update, no_update)
                else:
                    return (no_update, no_update)

    else:
        return (no_update, no_update)

@app.callback(
    [Output('rightIntegralFirstTab4', 'value'),
     Output('rightIntegralSecondTab4', 'value')],
    [Input('pointRightFirstTab4', 'children'),
     Input('pointRightSecondTab4', 'children'),
     Input('secondChoosenValueTab4', 'value'),
     Input('radiographtab4', 'value')], )
def display_hover_data4(leftchild, rightchild, secondchoosen, radioval):
    if leftchild == None or rightchild == None or leftchild == [] or rightchild == [] or secondchoosen == None:
        raise PreventUpdate

    if radioval == 'optionlibre':
        if secondchoosen != None and len(leftchild) == 2:
            for i in range(len(leftchild)):
                if leftchild[0] < leftchild[1]:
                    minchild = leftchild[0]
                    maxchild = leftchild[1]
                else:
                    minchild = leftchild[1]
                    maxchild = leftchild[0]
            return ('T ' + str(minchild), 'T ' + str(maxchild))
        elif secondchoosen == None:
            return '', ''
        else:
            return no_update, no_update
    if radioval == 'choosevalue':
        if secondchoosen != None and len(leftchild) == 2:
            for i in range(len(leftchild)):
                if leftchild[0] < leftchild[1]:
                    minchild = leftchild[0]
                    maxchild = leftchild[1]
                else:
                    minchild = leftchild[1]
                    maxchild = leftchild[0]
            return ('T ' + str(minchild), 'T ' + str(maxchild))
        elif secondchoosen == None:
            return '', ''
        else:
            return no_update, no_update
    else:
        return ('', '')


@app.callback(Output('leftIntegral', 'value'),
              [Input('leftIntegralFirst', 'value'),
               Input('leftIntegralSecond', 'value'),
               Input('firstChoosenValue', 'value'),
               ], [State('retrieve', 'children')]
              )
def integralCalculation(st1left, st1right, valuechoosenleft, retrieve):
    if st1left == None or st1right == None or valuechoosenleft == [] or retrieve == None or retrieve == []:
        raise PreventUpdate
    if st1left.startswith('T') == 1 and st1right.startswith('T') == 1:
        st1left = st1left[2:]
        st1right = st1right[2:]
    elif st1left.startswith('T') == 1 and st1right.isnumeric() == 1:
        st1left = st1left[2:]
        st1right = st1right
    elif st1left.isnumeric() == 1 and st1right.isnumeric() == 1:
        st1left = st1left
        st1right = st1right
    elif st1left.isnumeric() == 1 and st1right.startswith('T') == 1:
        st1left = st1left
        st1right = st1right[2:]
    if len(retrieve) > 0:
        if st1left != '' and st1right != '':
            df = pd.read_excel('appending.xlsx')
            df = df.loc[:, ~df.columns.str.match('Unnamed')]
            df['index'] = df.index
            df = df.reindex(columns=sorted(df.columns, reverse=True))
            if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                df = df.loc[df['ID'] == valuechoosenleft]
                df = df.copy()
                index = np.arange(0, len(df['ID']))
                df.reset_index(drop=True, inplace=True)
                df.set_index(index, inplace=True)
                dff1 =  df[(df[df['ID'] == valuechoosenleft].index >= float(st1left)) & (df[df['ID'] == valuechoosenleft].index <= float(st1right))]
                c = dff1['Value']
                area1 = abs(trapz((abs(c)), dx=1))
                return area1

            else :
                dff1 = df[(df[valuechoosenleft].index >= float(st1left)) & (df[valuechoosenleft].index <= float(st1right)) |
                      (df[valuechoosenleft].index >= float(st1right)) & (df[valuechoosenleft].index <= float(st1left))]
            for i in df.columns:
                if i.startswith('Temps'):
                    dff1 = dff1.groupby(i).mean()

                c = dff1[valuechoosenleft]
                area1 = abs(trapz((abs(c)), dx=1))

                return area1
        elif (st1left == '' and st1right != '') or (st1left != '' and st1right == ''):
            return 'total integration'
        elif (st1left == '' and st1right == '') and valuechoosenleft != None:
            return 'total integration'
        elif st1left != '' and st1right != '' and valuechoosenleft == None:
            return 'total integration'
        elif st1left == '' and st1right != '' and valuechoosenleft == None:
            return 'total integration'

@app.callback(Output('leftIntegraldb', 'value'),
              [Input('leftIntegralFirstdb', 'value'),
               Input('leftIntegralSeconddb', 'value'),
               Input('firstChoosenValuedb', 'value'), ],
              [State('memory-output', 'data'),
               State('dbvalchoosen', 'value'), State('db_name', 'value'), State('dbvaldate', 'value')]
              )
def integralCalculation(st1left, st1right, valuechoosenleft, retrieve, dbch, dbname, valdate):
    if st1left == None or st1right == None or valuechoosenleft == None or valuechoosenleft == [] or retrieve == None or retrieve == []:
        raise PreventUpdate
    if st1left.startswith('T') == 1 and st1right.startswith('T') == 1:
        st1left = st1left[2:]
        st1right = st1right[2:]
    elif st1left.startswith('T') == 1 and st1right.isnumeric() == 1:
        st1left = st1left[2:]
        st1right = st1right
    elif st1left.isnumeric() == 1 and st1right.isnumeric() == 1:
        st1left = st1left
        st1right = st1right
    elif st1left.isnumeric() == 1 and st1right.startswith('T') == 1:
        st1left = st1left
        st1right = st1right[2:]
    if retrieve != []:
        if st1left != '' and st1right != '':
            df = pd.DataFrame(retrieve)
            if dbname == 'rcckn':
                if dbch == 'send_variablevalues':
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'TIMESTAMP',
                                  'PROCESSED', 'TIMED_OUT', 'UNREFERENCED']
                if dbch == 'received_variablevalues':
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'LOCAL_TIMESTAMP',
                                  'REMOTE_ID', 'REMOTE_TIMESTAMP', 'REMOTE_MESSAGE_ID', 'PROCESSED', 'TIMED_OUT',
                                  'CONVERTED_NUM_VALUE']
            if dbname == 'enerbat':
                df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
            df1 = df[df['VARIABLE_NAME'] == valuechoosenleft]
            if dbch == 'send_variablevalues':
                df1 = df1[df1.TIMESTAMP.str.startswith(valdate[0])]
            if dbch == 'received_variablevalues':
                df1 = df1[df1.REMOTE_TIMESTAMP.str.startswith(valdate[0])]
            else:
                df1 = df1[df1.TIMESTAMP.str.startswith(valdate[0])]
            index = np.arange(0, len(df1['VARIABLE_NAME']))
            df1.reset_index(drop=True, inplace=True)
            df1.set_index(index, inplace=True)
            dff2 = df1[(df1.index >= float(st1left)) & (df1.index <= float(st1right)) |
                       (df1.index >= float(st1right)) & (df1.index <= float(st1left))]
            c = dff2['VARIABLE_NUM_VALUE']
            area1 = abs(trapz((abs(c)), dx=1))

            return area1
        elif (st1left == '' and st1right != '') or (st1left != '' and st1right == ''):
            return 'total integration'
        elif (st1left == '' and st1right == '') and valuechoosenleft != None:
            return 'total integration'
        elif st1left != '' and st1right != '' and valuechoosenleft == None:
            return 'total integration'
        elif st1left == '' and st1right != '' and valuechoosenleft == None:
            return 'total integration'


@app.callback(Output('leftIntegralTab4', 'value'),
              [Input('leftIntegralFirstTab4', 'value'),
               Input('leftIntegralSecondTab4', 'value'),
               Input('firstChoosenValueTab4', 'value'),
               ], [State('retrieve', 'children')]
              )
def integralCalculationtab4(st1left, st1right, valuechoosenleft, retrieve):
    if st1left == None or st1right == None or valuechoosenleft == None or valuechoosenleft == [] or retrieve == None or retrieve == []:
        raise PreventUpdate
    if st1left.startswith('T') == 1 and st1right.startswith('T') == 1:
        st1left = st1left[2:]
        st1right = st1right[2:]
    elif st1left.startswith('T') == 1 and st1right.isnumeric() == 1:
        st1left = st1left[2:]
        st1right = st1right
    elif st1left.isnumeric() == 1 and st1right.isnumeric() == 1:
        st1left = st1left
        st1right = st1right
    elif st1left.isnumeric() == 1 and st1right.startswith('T') == 1:
        st1left = st1left
        st1right = st1right[2:]
    if len(retrieve) > 0:
        if st1left != '' and st1right != '':
            df = pd.read_excel('appending.xlsx')
            df = df.loc[:, ~df.columns.str.match('Unnamed')]
            if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                a = df['ID'].unique()
                dff2 = pd.DataFrame([])
                for i in a:
                    dff = df[df['ID'] == i]

                    index = np.arange(0, len(dff))
                    dff.reset_index(drop=True, inplace=True)
                    dff.set_index(index, inplace=True)
                    # dff.drop(['Unnamed: 0', 'Unnamed: 0.1', 'Unnamed: 0.1.1'], axis=1, inplace=True)
                    dff = dff.pivot(values='Value', columns='ID')
                    dff2 = pd.concat([dff2, dff], axis=1)
                df = dff2.copy()

            else :
                df['index'] = df.index
                df = df.reindex(columns=sorted(df.columns, reverse=True))
            dff1 = df[(df[valuechoosenleft].index >= float(st1left)) & (df[valuechoosenleft].index <= float(st1right)) |
                      (df[valuechoosenleft].index >= float(st1right)) & (df[valuechoosenleft].index <= float(st1left))]
            for i in df.columns:
                if i.startswith('Temps'):
                    dff1 = dff1.groupby(i).mean()
            c = dff1[valuechoosenleft]
            area1 = abs(trapz(abs(c), dx=1))

            return area1
        elif (st1left == '' and st1right != '') or (st1left != '' and st1right == ''):
            return 'total integration'
        elif (st1left == '' and st1right == '') and valuechoosenleft != None:
            return 'total integration'
        elif st1left != '' and st1right != '' and valuechoosenleft == None:
            return 'total integration'
        elif st1left == '' and st1right != '' and valuechoosenleft == None:
            return 'total integration'


@app.callback(Output('rightIntegral', 'value'),
              [Input('rightIntegralFirst', 'value'),
               Input('rightIntegralSecond', 'value'),
               Input('secondChoosenValue', 'value'),
               ], [State('retrieve', 'children')]
              )
def integralCalculation2(st2left, st2right, valuechoosenright, retrieve):
    if st2left == None or st2right == None or valuechoosenright == [] or retrieve == None or retrieve == []:
        raise PreventUpdate
    if st2left.startswith('T') == 1 and st2right.startswith('T') == 1:
        st2left = st2left[2:]
        st2right = st2right[2:]
    elif st2left.startswith('T') == 1 and st2right.isnumeric() == 1:
        st2left = st2left[2:]
        st2right = st2right
    elif st2left.isnumeric() == 1 and st2right.isnumeric() == 1:
        st2left = st2left
        st2right = st2right
    elif st2left.isnumeric() == 1 and st2right.startswith('T') == 1:
        st2left = st2left
        st2right = st2right[2:]
    if len(retrieve) > 0:
        if st2left != '' and st2right != '':
            df = pd.read_excel('appending.xlsx')
            df = df.loc[:, ~df.columns.str.match('Unnamed')]
            df['index'] = df.index
            df = df.reindex(columns=sorted(df.columns, reverse=True))
            if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                df = df.loc[df['ID'] == valuechoosenright]
                df = df.copy()
                index = np.arange(0, len(df['ID']))
                df.reset_index(drop=True, inplace=True)
                df.set_index(index, inplace=True)
                dff2 = df[(df[df['ID'] == valuechoosenright].index >= float(st2left)) & (
                            df[df['ID'] == valuechoosenright].index <= float(st2right))]
                f = dff2['Value']
                area1 = abs(trapz((abs(f)), dx=1))

                return area1

            else :
                dff2 = df[
                    (df[valuechoosenright].index >= float(st2left)) & (df[valuechoosenright].index <= float(st2right)) |
                    (df[valuechoosenright].index >= float(st2right)) & (df[valuechoosenright].index <= float(st2left))]
                for i in df.columns:
                    if i.startswith('Temps'):
                        dff2 = dff2.groupby(i).mean()

                f = dff2[valuechoosenright]
                area2 = abs(trapz(abs(f), dx=1))
                return area2
        elif (st2left == '' and st2right != '') or (st2left != '' and st2right == ''):
            return 'total integration'
        elif (st2left == '' and st2right == '') and valuechoosenright != None:
            return 'total integration'
        elif st2left != '' and st2right != '' and valuechoosenright == None:
            return 'total integration'
        elif st2left == '' and st2right != '' and valuechoosenright == None:
            return 'total integration'


@app.callback(Output('rightIntegraldb', 'value'),
              [Input('rightIntegralFirstdb', 'value'),
               Input('rightIntegralSeconddb', 'value'),
               Input('secondChoosenValuedb', 'value'), ],
              [State('memory-output', 'data'),
               State('dbvalchoosen', 'value'), State('db_name', 'value'), State('dbvaldate', 'value')]
              )
def integralCalculationdb(st2left, st2right, valuechoosenright, retrieve, dbch, dbname, valdate):
    if st2left == None or st2right == None or valuechoosenright == [] or retrieve == None or retrieve == []:
        raise PreventUpdate
    if st2left.startswith('T') == 1 and st2right.startswith('T') == 1:
        st2left = st2left[2:]
        st2right = st2right[2:]
    elif st2left.startswith('T') == 1 and st2right.isnumeric() == 1:
        st2left = st2left[2:]
        st2right = st2right
    elif st2left.isnumeric() == 1 and st2right.isnumeric() == 1:
        st2left = st2left
        st2right = st2right
    elif st2left.isnumeric() == 1 and st2right.startswith('T') == 1:
        st2left = st2left
        st2right = st2right[2:]
    if retrieve != []:
        if st2left != '' and st2right != '':
            df = pd.DataFrame(retrieve)
            if dbname == 'rcckn':
                if dbch == 'send_variablevalues':
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'TIMESTAMP',
                                  'PROCESSED', 'TIMED_OUT', 'UNREFERENCED']
                if dbch == 'received_variablevalues':
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'LOCAL_TIMESTAMP',
                                  'REMOTE_ID', 'REMOTE_TIMESTAMP', 'REMOTE_MESSAGE_ID', 'PROCESSED', 'TIMED_OUT',
                                  'CONVERTED_NUM_VALUE']
            if dbname == 'enerbat':
                df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
            df1 = df[df['VARIABLE_NAME'] == valuechoosenright]
            if dbch == 'send_variablevalues':
                df1 = df1[df1.TIMESTAMP.str.startswith(valdate[0])]
            if dbch == 'received_variablevalues':
                df1 = df1[df1.REMOTE_TIMESTAMP.str.startswith(valdate[0])]
            index = np.arange(0, len(df1['VARIABLE_NAME']))
            df1.reset_index(drop=True, inplace=True)
            df1.set_index(index, inplace=True)
            dff2 = df1[(df1.index >= float(st2left)) & (df1.index <= float(st2right)) |
                       (df1.index >= float(st2right)) & (df1.index <= float(st2left))]
            c = dff2['VARIABLE_NUM_VALUE']
            area1 = abs(trapz((abs(c)), dx=1))

            return area1
        elif (st2left == '' and st2right != '') or (st2left != '' and st2right == ''):
            return 'total integration'
        elif (st2left == '' and st2right == '') and valuechoosenright != None:
            return 'total integration'
        elif st2left != '' and st2right != '' and valuechoosenright == None:
            return 'total integration'
        elif st2left == '' and st2right != '' and valuechoosenright == None:
            return 'total integration'


@app.callback(Output('rightIntegralTab4', 'value'),
              [Input('rightIntegralFirstTab4', 'value'),
               Input('rightIntegralSecondTab4', 'value'),
               Input('secondChoosenValueTab4', 'value'),
               ], [State('retrieve', 'children')]
              )
def integralCalculation4(st2left, st2right, valuechoosenright, retrieve):
    if st2left == None or st2right == None or  valuechoosenright == [] or retrieve == None or retrieve == []:
        raise PreventUpdate
    if st2left.startswith('T') == 1 and st2right.startswith('T') == 1:
        st2left = st2left[2:]
        st2right = st2right[2:]
    elif st2left.startswith('T') == 1 and st2right.isnumeric() == 1:
        st2left = st2left[2:]
        st2right = st2right
    elif st2left.isnumeric() == 1 and st2right.isnumeric() == 1:
        st2left = st2left
        st2right = st2right
    elif st2left.isnumeric() == 1 and st2right.startswith('T') == 1:
        st2left = st2left
        st2right = st2right[2:]
    if len(retrieve) > 0:
        if st2left != '' and st2right != '':
            df = pd.read_excel('appending.xlsx')
            df = df.loc[:, ~df.columns.str.match('Unnamed')]
            if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                a = df['ID'].unique()
                dff2 = pd.DataFrame([])
                for i in a:
                    dff = df[df['ID'] == i]

                    index = np.arange(0, len(dff))
                    dff.reset_index(drop=True, inplace=True)
                    dff.set_index(index, inplace=True)
                    # dff.drop(['Unnamed: 0', 'Unnamed: 0.1', 'Unnamed: 0.1.1'], axis=1, inplace=True)
                    dff = dff.pivot(values='Value', columns='ID')
                    dff2 = pd.concat([dff2, dff], axis=1)
                df = dff2.copy()
            else:
                df['index'] = df.index
                df = df.reindex(columns=sorted(df.columns, reverse=True))
            dff2 = df[
                (df[valuechoosenright].index >= float(st2left)) & (df[valuechoosenright].index <= float(st2right)) |
                (df[valuechoosenright].index >= float(st2right)) & (df[valuechoosenright].index <= float(st2left))]
            for i in df.columns:
                if i.startswith('Temps'):
                    dff2 = dff2.groupby(i).mean()
            f = dff2[valuechoosenright]
            area2 = abs(trapz(abs(f), dx=1))
            return area2
        elif (st2left == '' and st2right != '') or (st2left != '' and st2right == ''):
            return 'total integration'
        elif (st2left == '' and st2right == '') and valuechoosenright != None:
            return 'total integration'
        elif st2left != '' and st2right != '' and valuechoosenright == None:
            return 'total integration'
        elif st2left == '' and st2right != '' and valuechoosenright == None:
            return 'total integration'


@app.callback(Output('operation', 'value'),
              [Input('leftIntegral', 'value'),
               Input('rightIntegral', 'value'),
               Input('operateur', 'value')],
              )
def differanceintegration(value1, value2, ops):
    if value1 == None or value2 == None:
        raise PreventUpdate
    if ops == ['Plus']:
        return float(value1 + value2)
    elif ops == ['Moins']:
        return float(value1 - value2)
    elif ops == ['Multiplie']:
        return float(value1 * value2)
    elif ops == ['Division']:
        return float(value1 / value2)
    elif ops == []:
        return []


@app.callback(Output('operationTab4', 'value'),
              [Input('leftIntegralTab4', 'value'),
               Input('rightIntegralTab4', 'value'),
               Input('operateurTab4', 'value')],
              )
def differanceintegrationTab4(value1, value2, ops):
    if value1 == None or value2 == None:
        raise PreventUpdate
    if ops == ['Plus']:
        return float(value1 + value2)
    elif ops == ['Moins']:
        return float(value1 - value2)
    elif ops == ['Multiplie']:
        return float(value1 * value2)
    elif ops == ['Division']:
        return float(value1 / value2)
    elif ops == []:
        return []


@app.callback(Output('operationdb', 'value'),
              [Input('leftIntegraldb', 'value'),
               Input('rightIntegraldb', 'value'),
               Input('operateurdb', 'value')],
              )
def differanceintegrationdb(value1, value2, ops):
    if value1 == None or value2 == None:
        raise PreventUpdate
    if ops == ['Plus']:
        return float(value1 + value2)
    elif ops == ['Moins']:
        return float(value1 - value2)
    elif ops == ['Multiplie']:
        return float(value1 * value2)
    elif ops == ['Division']:
        return float(value1 / value2)
    elif ops == []:
        return []

@app.callback(Output('intersection', 'value'),
              [Input('pointLeftFirst', 'children'),
               Input('pointRightFirst', 'children'),
               Input('firstChoosenValue', 'value'),
               Input('secondChoosenValue', 'value'),
               Input('leftIntegralFirst', 'value'),
               Input('rightIntegralFirst', 'value'), ],
              [State('intersection', 'value'), State('retrieve', 'children'),
               ]
              )
def differanceCalculation(firstshape, secondshape, valuechoosenleft, valuechoosenright, leftfirst, rightfirst, diff, retrieve):
    if retrieve == None or retrieve == []:
        raise PreventUpdate
    differance = []
    if len(firstshape) == 2 and len(secondshape) == 2:
        a = int(firstshape[0])
        c = int(secondshape[0])
        b = int(firstshape[1])
        d = int(secondshape[1])
        if set(range(a, b)).issuperset(set(range(c, d))) == 1:
            differance.append(c)
            differance.append(d)
        elif set(range(c, d)).issuperset(set(range(a, b))) == 1:
            differance.append(a)
            differance.append(b)
        elif len(set(range(a, b)).intersection(set(range(c, d)))) >= 1 or len(
                set(range(c, d)).intersection(set(range(a, b)))) >= 1:
            if a <= c:
                if len(differance) == 2:
                    differance.pop(0)
                    differance.append(b)
                differance.append(b)
            if a >= c:
                if len(differance) == 2:
                    differance.pop(0)
                    differance.append(a)
                differance.append(a)
            if b <= d:
                 if len(differance) == 2:
                    differance.pop(0)
                    differance.append(c)
                 differance.append(c)
            if b >= d:
                if len(differance) == 2:
                    differance.pop(0)
                    differance.append(d)
                differance.append(d)
        else:
            return ['intersection']
        differancelast = []
        if len(retrieve) > 0 and valuechoosenright != None and valuechoosenleft != None and leftfirst != None and rightfirst != None:
            first = differance[0]
            second = differance[1]
            df = pd.read_excel('appending.xlsx')
            df = df.loc[:, ~df.columns.str.match('Unnamed')]
            if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                a = df['ID'].unique()
                dff2 = pd.DataFrame([])
                for i in a:
                    dff = df[df['ID'] == i]

                    index = np.arange(0, len(dff))
                    dff.reset_index(drop=True, inplace=True)
                    dff.set_index(index, inplace=True)
                    # dff.drop(['Unnamed: 0', 'Unnamed: 0.1', 'Unnamed: 0.1.1'], axis=1, inplace=True)
                    dff = dff.pivot(values='Value', columns='ID')
                    dff2 = pd.concat([dff2, dff], axis=1)
                df = dff2.copy()

            else:
                df['index'] = df.index
                df = df.reindex(columns=sorted(df.columns, reverse=True))
            dff = df[(df[valuechoosenright].index >= float(first)) & (df[valuechoosenright].index <= float(second)) |
                     (df[valuechoosenright].index >= float(second)) & (df[valuechoosenright].index <= float(first))]
            l = dff[valuechoosenright]

            dff2 = df[(df[valuechoosenleft].index >= float(first)) & (df[valuechoosenleft].index <= float(second)) |
                      (df[valuechoosenleft].index >= float(second)) & (df[valuechoosenleft].index <= float(first))]
            r = dff2[valuechoosenleft]
            tt = []
            yy = []
            for i in l:
                tt.append(i)
            for i in r:
                yy.append(i)
            for i in range(len(tt)):
                if tt[i] <= yy[i]:
                    differancelast.append(tt[i])
                if yy[i] < tt[i]:
                    differancelast.append(yy[i])
            diff = (abs(trapz(differancelast, dx=1)))
            return diff


@app.callback(Output('intersectiondb', 'value'),
              [Input('firstChoosenValuedb', 'value'),
               Input('secondChoosenValuedb', 'value'),
               Input('leftIntegralFirstdb', 'value'),
               Input('rightIntegralFirstdb', 'value'),
               Input('leftIntegralSeconddb', 'value'),
               Input('rightIntegralSeconddb', 'value'),
               ],
              [State('intersectiondb', 'value'), State('memory-output', 'data'),
               State('dbvalchoosen', 'value'), State('db_name', 'value'), State('dbvaldate', 'value')
               ]
              )
def differanceCalculationdb(valuechoosenleft, valuechoosenright, leftfirst, rightfirst, leftsecond, rightsecond,
                            diff, retrieve, dbch, dbname, dbdate):
    if retrieve == None or retrieve == [] or leftfirst == None or rightfirst == None or leftsecond == None or rightsecond == None:
        raise PreventUpdate

    if valuechoosenright != None and valuechoosenleft != None:
        differance = []
        st1left = leftfirst[2:]
        a,b,c,d = 0,0,0,0
        if st1left != '':
            a = int(st1left)
        else : a == 0
        st1right = leftsecond[2:]
        if st1right != '':
            b = int(st1right)
        else : b == 0
        st2left = rightfirst[2:]
        if st2left != '':
            c = int(st2left)
        else : c == 0
        st2right = rightsecond[2:]
        if st2right != '':
            d = int(st2right)
        else : d == 0
        if set(range(a, b)).issuperset(set(range(c, d))) == 1:
            differance.append(c)
            differance.append(d)
        elif set(range(c, d)).issuperset(set(range(a, b))) == 1:
            differance.append(a)
            differance.append(b)
        elif len(set(range(a, b)).intersection(set(range(c, d)))) >= 1 or len(
                set(range(c, d)).intersection(set(range(a, b)))) >= 1:
            if a <= c:
                if len(differance) == 2:
                    differance.pop(0)
                    differance.append(b)
                differance.append(b)
            if a >= c:
                if len(differance) == 2:
                    differance.pop(0)
                    differance.append(a)
                differance.append(a)
            if b <= d:
                if len(differance) == 2:
                    differance.pop(0)
                    differance.append(c)
                differance.append(c)
            if b >= d:
                if len(differance) == 2:
                    differance.pop(0)
                    differance.append(d)
                differance.append(d)
        else:
            return ['intersection']
        df1 = pd.DataFrame(retrieve)
        dates = []
        if dbname == 'rcckn':
            if dbch == 'send_variablevalues':
                df1.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'TIMESTAMP',
                               'PROCESSED', 'TIMED_OUT', 'UNREFERENCED']
                for col in df1['TIMESTAMP']:
                    dates.append(col[:10])
            if dbch == 'received_variablevalues':
                df1.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'LOCAL_TIMESTAMP',
                               'REMOTE_ID', 'REMOTE_TIMESTAMP', 'REMOTE_MESSAGE_ID', 'PROCESSED', 'TIMED_OUT',
                               'CONVERTED_NUM_VALUE']
                for col in df1['REMOTE_TIMESTAMP']:
                    dates.append(col[:10])
        if dbname == 'enerbat':
            df1.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
            for col in df1['TIMESTAMP']:
                dates.append(col[:10])
        #

        df1['dates'] = dates
        first_df = df1[df1['VARIABLE_NAME'] == valuechoosenleft]
        second_df = df1[df1['VARIABLE_NAME'] == valuechoosenright]
        first_df = first_df[first_df['dates'].isin(dbdate)]
        index = np.arange(0, len(first_df['VARIABLE_NAME']))
        first_df.reset_index(drop=True, inplace=True)
        first_df.set_index(index, inplace=True)
        first_df = first_df[(first_df.index >= float(differance[0])) & (first_df.index <= float(differance[1]))]
        first_df = first_df['VARIABLE_NUM_VALUE']

        second_df = second_df[second_df['dates'].isin(dbdate)]
        index = np.arange(0, len(second_df['VARIABLE_NAME']))
        second_df.reset_index(drop=True, inplace=True)
        second_df.set_index(index, inplace=True)
        second_df = second_df[(second_df.index >= float(differance[0])) & (second_df.index <= float(differance[1]))]
        second_df = second_df['VARIABLE_NUM_VALUE']
        min_val = []
        for i, j in zip(first_df, second_df):
            if i <= j:
                min_val.append(i)
            if j < i:
                min_val.append(j)

        diff = (abs(trapz(min_val, dx=1)))
        return diff
    else:
        return ['intersection']


@app.callback(Output('intersectionTab4', 'value'),
              [Input('pointLeftFirstTab4', 'children'),
               Input('pointRightFirstTab4', 'children'),
               Input('firstChoosenValueTab4', 'value'),
               Input('secondChoosenValueTab4', 'value'),
               Input('leftIntegralFirstTab4', 'value'),
               Input('rightIntegralFirstTab4', 'value'), ],
              [State('intersectionTab4', 'value'), State('retrieve', 'children'),
               ]
              )
def differanceCalculation4(firstshape, secondshape, valuechoosenleft, valuechoosenright, leftfirst, rightfirst, diff,
                           retrieve):

    if retrieve == None or retrieve == []:
        raise PreventUpdate
    differance = []
    if len(firstshape) == 2 and len(secondshape) == 2:
        a = int(firstshape[0])
        c = int(secondshape[0])
        b = int(firstshape[1])
        d = int(secondshape[1])
        if set(range(a, b)).issuperset(set(range(c, d))) == 1:
            differance.append(c)
            differance.append(d)
        elif set(range(c, d)).issuperset(set(range(a, b))) == 1:
            differance.append(a)
            differance.append(b)
        elif len(set(range(a, b)).intersection(set(range(c, d)))) >= 1 or len(
                set(range(c, d)).intersection(set(range(a, b)))) >= 1:
            if a <= c:
                if len(differance) == 2:
                    differance.pop(0)
                    differance.append(b)
                differance.append(b)
            if a >= c:
                if len(differance) == 2:
                    differance.pop(0)
                    differance.append(a)
                differance.append(a)
            if b <= d:
                if len(differance) == 2:
                    differance.pop(0)
                    differance.append(c)
                differance.append(c)
            if b >= d:
                if len(differance) == 2:
                    differance.pop(0)
                    differance.append(d)
                differance.append(d)
        else:
            return ['intersection']
        differancelast = []
        if len(retrieve) > 0 and valuechoosenright != None and valuechoosenleft != None and leftfirst != None and rightfirst != None:
            first = differance[0]
            second = differance[1]
            df = pd.read_excel('appending.xlsx')
            df = df.loc[:, ~df.columns.str.match('Unnamed')]
            if 'ID' and 'Value' and 'Quality' and 'Date' in df.columns:
                a = df['ID'].unique()
                dff2 = pd.DataFrame([])
                for i in a:
                    dff = df[df['ID'] == i]

                    index = np.arange(0, len(dff))
                    dff.reset_index(drop=True, inplace=True)
                    dff.set_index(index, inplace=True)
                    # dff.drop(['Unnamed: 0', 'Unnamed: 0.1', 'Unnamed: 0.1.1'], axis=1, inplace=True)
                    dff = dff.pivot(values='Value', columns='ID')
                    dff2 = pd.concat([dff2, dff], axis=1)
                df = dff2.copy()

            else :
                df['index'] = df.index
                df = df.reindex(columns=sorted(df.columns, reverse=True))
            dff = df[(df[valuechoosenright].index >= float(first)) & (df[valuechoosenright].index <= float(second)) |
                     (df[valuechoosenright].index >= float(second)) & (df[valuechoosenright].index <= float(first))]
            l = dff[valuechoosenright]

            dff2 = df[(df[valuechoosenleft].index >= float(first)) & (df[valuechoosenleft].index <= float(second)) |
                      (df[valuechoosenleft].index >= float(second)) & (df[valuechoosenleft].index <= float(first))]
            r = dff2[valuechoosenleft]
            tt = []
            yy = []
            for i in l:
                tt.append(i)
            for i in r:
                yy.append(i)
            for i in range(len(tt)):
                if tt[i] <= yy[i]:
                    differancelast.append(tt[i])
                if yy[i] < tt[i]:
                    differancelast.append(yy[i])
            diff = (abs(trapz(differancelast, dx=1)))
            return diff


@app.callback(Output('writeexcelhidden', 'children'),
              [Input('write_excel', 'n_clicks')],
              [State('firstChoosenValue', 'value'),
               State('leftIntegralFirst', 'value'),
               State('leftIntegralSecond', 'value'),
               State('leftIntegral', 'value'),
               State('secondChoosenValue', 'value'),
               State('rightIntegralFirst', 'value'),
               State('rightIntegralSecond', 'value'),
               State('rightIntegral', 'value'),
               State('operation', 'value'),
               State('intersection', 'value'),
               State('inputRightX_axis', 'value'),
               State('inputRightY_axis', 'value')
               ],
              )
def write_excel(nc, a, b, c, d, e, f, g, h, i, j,ref_X, ref_Y):
    if nc > 0:
        now = datetime.datetime.now()
        if i == []:
            i = None
        if j == ['intersection']:
            j = None
        x = ( a, b, c, d, e, f, g, h, i, j,ref_X,ref_Y,now)
        if x != None: return x


@app.callback(Output('writeexcelhiddenTab4', 'children'),
              [Input('write_excelTab4', 'n_clicks')],
              [State('firstChoosenValueTab4', 'value'),
               State('leftIntegralFirstTab4', 'value'),
               State('leftIntegralSecondTab4', 'value'),
               State('leftIntegralTab4', 'value'),
               State('secondChoosenValueTab4', 'value'),
               State('rightIntegralFirstTab4', 'value'),
               State('rightIntegralSecondTab4', 'value'),
               State('rightIntegralTab4', 'value'),
               State('operationTab4', 'value'),
               State('intersectionTab4', 'value'),

               ],
              )
def write_excelTab4(nc, a, b, c, d, e, f, g, h, i, j):
    if nc > 0:
        now = datetime.datetime.now()
        if i == []:
            i = None
        if j == ['intersection']:
            j = None
        x = (now, a, b, c, d, e, f, g, h, i, j)

        if x != None: return x


@app.callback(Output('hiddenrecord3', 'children'),
              [Input('writeexcelhidden', 'children'), Input('writeexcelhiddenTab4', 'children')],
              )
def pasfunc(hiddenvalchild, hiddenvalchild4):
    if hiddenvalchild == None and hiddenvalchild4 == None:
        raise PreventUpdate
    if hiddenvalchild != None:
        return hiddenvalchild
    if hiddenvalchild4 != None:
        return hiddenvalchild4


@app.callback(Output('hiddenrecord4', 'children'),
              [Input('hiddenrecord3', 'children')],
              [State('hiddenrecord4', 'children')], )
def lastfunc(hiddenvalchild, lastvalchild):
    lastvalchild = hiddenvalchild + lastvalchild
    return lastvalchild


@app.callback(Output('hiddenrecord1', 'children'),
              [Input('hiddenrecord4', 'children')],
              )
def exportdata(valueparse):
    a_parse = []
    t_parse = []
    for i in valueparse:
        if i == None:
            a_parse.append('None')
        else:
            a_parse.append(i)
        if len(a_parse) % 13 == 0:
            t_parse.append(a_parse)
            a_parse = []
    t_parse.insert(0, ['firstChoosenValue', 'leftIntegralFirst', 'leftIntegralSecond', 'leftIntegral',
                       'secondChoosenValue','rightIntegralFirst', 'rightIntegralSecond', 'rightIntegral', 'operation',
                       'intersection','referance X', 'referance Y','time'])

    df = pd.DataFrame(t_parse)
    df.to_excel('new_fichier.xlsx')

@app.callback(Output('writeexcelhiddendb', 'children'),
              [Input('write_exceldb', 'n_clicks')],
              [State('firstChoosenValuedb', 'value'),
               State('leftIntegralFirstdb', 'value'),
               State('leftIntegralSeconddb', 'value'),
               State('leftIntegraldb', 'value'),
               State('secondChoosenValuedb', 'value'),
               State('rightIntegralFirstdb', 'value'),
               State('rightIntegralSeconddb', 'value'),
               State('rightIntegraldb', 'value'),
               State('operationdb', 'value'),
               State('intersectiondb', 'value'),
               ],
              )
def write_exceldb(nc, a, b, c, d, e, f, g, h, i, j):
    if nc > 0:
        now = datetime.datetime.now()
        if i == []:
            i = None
        if j == ['intersection']:
            j = None
        x = (now, a, b, c, d, e, f, g, h, i, j)

        if x != None: return x


@app.callback(Output('hiddenrecord3db', 'children'),
              [Input('writeexcelhiddendb', 'children'), Input('writeexcelhiddendb', 'children')],
              )
def pasfuncdb(hiddenvalchild, hiddenvalchild4):
    if hiddenvalchild == None and hiddenvalchild4 == None:
        raise PreventUpdate
    if hiddenvalchild != None:
        return hiddenvalchild
    if hiddenvalchild4 != None:
        return hiddenvalchild4


@app.callback(Output('hiddenrecord4db', 'children'),
              [Input('hiddenrecord3db', 'children')],
              [State('hiddenrecord4db', 'children')], )
def lastfuncdb(hiddenvalchild, lastvalchild):
    lastvalchild = hiddenvalchild + lastvalchild
    return lastvalchild


@app.callback(Output('hiddenrecord1db', 'children'),
              [Input('hiddenrecord4db', 'children')],
              )
def exportdatadb(valueparse):
    a_parse = []
    t_parse = []
    for i in valueparse:
        if i == None:
            a_parse.append('None')
        else:
            a_parse.append(i)
        if len(a_parse) % 11 == 0:
            t_parse.append(a_parse)
            a_parse = []
    t_parse.insert(0, ['time', 'firstChoosenValue', 'leftIntegralFirst', 'leftIntegralSecond', 'leftIntegral',
                       'secondChoosenValue','rightIntegralFirst', 'rightIntegralSecond', 'rightIntegral', 'operation',
                       'intersection'])

    df = pd.DataFrame(t_parse)
    df.columns = df.iloc[0]
    df.drop(df.index[0], inplace=True, axis=0)
    df.to_excel('new_fichier.xlsx')


@app.server.route("/download_excel/")
def download_excel():
    # Create DF
    dff = pd.read_excel("new_fichier.xlsx")
    dff.columns = dff.iloc[0]
    dff.drop(dff.index[0], inplace=True, axis=0)
    dff = dff.filter(regex='^\D')
    # Convert DF
    buf = io.BytesIO()
    excel_writer = pd.ExcelWriter(buf, engine="xlsxwriter")
    dff.to_excel(excel_writer, sheet_name="sheet1")
    excel_writer.save()
    excel_data = buf.getvalue()
    buf.seek(0)
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        attachment_filename="LERMAB.xlsx",
        as_attachment=True,
        cache_timeout=0
    )


@app.server.route("/download_exceldb/")
def download_exceldb():
    # Create DF
    dff = pd.read_excel("new_fichier.xlsx")
    # Convert DF
    buf = io.BytesIO()
    excel_writer = pd.ExcelWriter(buf, engine="xlsxwriter")
    dff.to_excel(excel_writer, sheet_name="sheet1")
    excel_writer.save()
    excel_data = buf.getvalue()
    buf.seek(0)
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        attachment_filename="LERMAB.xlsx",
        as_attachment=True,
        cache_timeout=0
    )

@app.callback(Output('dbvalchoosen', 'options'),
              [Input('db_name', 'value')], [State('db_Ip', 'value'),State('db_password', 'value')])
def relationdb(dbname, ipval, password):
    if dbname == None:
        raise PreventUpdate
    ipadress = "193.54.2.211"
    server = SSHTunnelForwarder(
        (ipadress, 22),
        ssh_username='soudani',
        ssh_password="univ484067152",
        remote_bind_address=(ipadress, 3306))

    server.start()

    try:
        conn = mariadb.connect(
            user = ipval,
            password = password,
            host = ipadress,
            port = 3306,
            database = dbname)

    except mariadb.Error as e:
        print(f"Error connecting to MariaDB Platform: {e}")
        sys.exit(1)
    if dbname == 'rcckn':
        cur = conn.cursor()
        cur.execute(f"select table_name from information_schema.tables where TABLE_SCHEMA= 'rcckn'")
        val = cur.fetchall()

        return [{'label': i[0], 'value': i[0]} for i in val if
                i[0] != 'app_variablerequest' and i[0] != 'send_controlvalues' and
                i[0] != 'received_ack' and i[0] != 'send_vw_variablerequestdestination'
                and i[0] != 'flyway_schema_history' and i[0] != 'app_vw_messaging_followup' and
                i[0] != 'received_variablerequest' and i[0] != 'received_controlvalues'
                and i[0] != 'app_system_properties' and i[0] != 'tbl_sites' and i[0] != 'tbl_inventory'
                and i[0] != 'send_messages' and i[0] != 'send_variablevaluesmessage']
    elif dbname == 'enerbat':
        cur = conn.cursor()
        cur.execute(f"select table_name from information_schema.tables where TABLE_SCHEMA= 'enerbat'")
        val = cur.fetchall()

        return [{'label': i[0], 'value': i[0]} for i in val]
    else:
        no_update


@app.callback(Output('prvalchoosen', 'options'),
              [Input('prname', 'value')],[State('username_pr', 'value'), State('password_pr', 'value')])
def relationpr(prname, user, passw):
    if prname == None:
        raise PreventUpdate
    ipadress = "193.54.2.211"
    server = SSHTunnelForwarder(
        (ipadress, 22),
        ssh_username='soudani',
        ssh_password="univ484067152",
        remote_bind_address=(ipadress, 3306))

    server.start()
    try:
        conn = mariadb.connect(
            user='dashapp' if user == None else user,
            password='dashapp' if passw == None else passw,
            host=ipadress,
            port=3306,
            database=prname)
        cur = conn.cursor()

    except mariadb.Error as e:
        print(f"Error connecting to MariaDB Platform: {e}")
        sys.exit(1)
    if prname == 'rcckn':
        cur = conn.cursor()
        cur.execute(f"select table_name from information_schema.tables where TABLE_SCHEMA= 'rcckn'")
        val = cur.fetchall()

        return [{'label': i[0], 'value': i[0]} for i in val if
                i[0] != 'app_variablerequest' and i[0] != 'send_controlvalues' and
                i[0] != 'received_ack' and i[0] != 'send_vw_variablerequestdestination' and i[
                    0] != 'flyway_schema_history'
                and i[0] != 'app_vw_messaging_followup' and i[0] != 'received_variablerequest' and i[
                    0] != 'received_controlvalues'
                and i[0] != 'app_system_properties' and i[0] != 'tbl_sites' and i[0] != 'tbl_inventory' and i[
                    0] != 'send_messages'
                and i[0] != 'send_variablevaluesmessage']
    elif prname == 'enerbat':
        cur = conn.cursor()
        cur.execute(f"select table_name from information_schema.tables where TABLE_SCHEMA= 'enerbat'")
        val = cur.fetchall()

        return [{'label': i[0], 'value': i[0]} for i in val]
    else:
        no_update


@app.callback([Output('dbvalname', 'options'), Output('dbvaldate', 'options')],
              [Input('activatedb', 'n_clicks'), Input('deactivatedb', 'n_clicks')],
              [State('dbvalchoosen', 'value'), State('db_name', 'value'), State('db_Ip', 'value'),State('db_password', 'value')])
def dbname(nc, nc2, dbch, dbname, ipval, password):
    if dbname == None:
        raise PreventUpdate
    ipadress = "193.54.2.211"
    q1 = dash.callback_context.triggered[0]["prop_id"].split(".")[0]
    if q1 == 'activatedb':

        server = SSHTunnelForwarder(
            (ipadress, 22),
            ssh_username='soudani',
            ssh_password="univ484067152",
            remote_bind_address=(ipadress, 3306))

        server.start()
        try:
            conn = mariadb.connect(
                user=ipval,
                password=password,
                host=ipadress,
                port=3306,
                database=dbname)

        except mariadb.Error as e:
            print(f"Error connecting to MariaDB Platform: {e}")
            sys.exit(1)
            # Get Cursor

        if dbname == 'rcckn':
            if dbch == 'received_variablevalues':
                cur1 = conn.cursor()
                cur1.execute("SELECT DISTINCT VARIABLE_NAME FROM received_variablevalues ")
                t1 = cur1.fetchall()
                name = [i[0] for i in t1]
                cur2 = conn.cursor()
                cur2.execute("SELECT DISTINCT REMOTE_TIMESTAMP FROM received_variablevalues ")
                t2 = cur2.fetchall()
                str_list = [i[0] for i in t2]
                df = pd.DataFrame(str_list)
                df.columns = ['REMOTE_TIMESTAMP']
                df['REMOTE_TIMESTAMP'] = df.REMOTE_TIMESTAMP.apply(pd.to_datetime)
                df["day"] = df.REMOTE_TIMESTAMP.dt.day
                df["month"] = df.REMOTE_TIMESTAMP.dt.month
                df["year"] = df.REMOTE_TIMESTAMP.dt.year
                a = [str(i) + '-' + str(j) + '-' + str(k) for i, j, k in zip(df["year"], df["month"], df["day"])]
                a = list(set(a))
                b = pd.to_datetime(a)
                b = sorted(b)
                str_list = [t.strftime("%Y-%m-%d") for t in b]
                return [{'label': i, 'value': i} for i in name], [{'label': i, 'value': i} for i in str_list]
            elif dbch == "send_variablevalues":
                cur1 = conn.cursor()
                cur1.execute("SELECT DISTINCT VARIABLE_NAME FROM send_variablevalues ")
                t1 = cur1.fetchall()
                name = [i[0] for i in t1]
                cur2 = conn.cursor()
                cur2.execute("SELECT DISTINCT TIMESTAMP FROM send_variablevalues ")
                t2 = cur2.fetchall()
                str_list = [i[0] for i in t2]
                df = pd.DataFrame(str_list)
                df.columns = ['TIMESTAMP']
                df['TIMESTAMP'] = df.TIMESTAMP.apply(pd.to_datetime)
                df["day"] = df.TIMESTAMP.dt.day
                df["month"] = df.TIMESTAMP.dt.month
                df["year"] = df.TIMESTAMP.dt.year
                a = [str(i) + '-' + str(j) + '-' + str(k) for i, j, k in zip(df["year"], df["month"], df["day"])]
                a = list(set(a))
                b = pd.to_datetime(a)
                b = sorted(b)
                str_list = [t.strftime("%Y-%m-%d") for t in b]
                return [{'label': i, 'value': i} for i in name], [{'label': i, 'value': i} for i in str_list]
            elif (dbch != 'received_variablevalues' or dbch != 'send_variablevalues') and dbch != None:
                cur = conn.cursor()
                cur.execute(
                    f"SELECT column_name FROM INFORMATION_SCHEMA.COLUMNS WHERE TABLE_NAME = '{dbch}' ORDER BY ORDINAL_POSITION")
                t = cur.fetchall()
                x = [ i[0] for i in t]
                if 'REMOTE_TIMESTAMP' in x:
                    cur1 = conn.cursor()
                    cur1.execute(f"SELECT DISTINCT VARIABLE_NAME FROM {dbch} ")
                    t1 = cur1.fetchall()
                    name = [i[0] for i in t1]
                    cur2 = conn.cursor()
                    cur2.execute(f"SELECT DISTINCT REMOTE_TIMESTAMP FROM {dbch} ")
                    t2 = cur2.fetchall()
                    str_list = [i[0] for i in t2]
                    df = pd.DataFrame(str_list)
                    df.columns = ['REMOTE_TIMESTAMP']
                    df['REMOTE_TIMESTAMP'] = df.REMOTE_TIMESTAMP.apply(pd.to_datetime)
                    df["day"] = df.REMOTE_TIMESTAMP.dt.day
                    df["month"] = df.REMOTE_TIMESTAMP.dt.month
                    df["year"] = df.REMOTE_TIMESTAMP.dt.year
                    a = [str(i) + '-' + str(j) + '-' + str(k) for i, j, k in zip(df["year"], df["month"], df["day"])]
                    a = list(set(a))
                    b = pd.to_datetime(a)
                    b = sorted(b)
                    str_list = [t.strftime("%Y-%m-%d") for t in b]
                    return [{'label': i, 'value': i} for i in name], [{'label': i, 'value': i} for i in str_list]

                else:
                    cur1 = conn.cursor()
                    cur1.execute(f"SELECT DISTINCT VARIABLE_NAME FROM {dbch} ")
                    t1 = cur1.fetchall()
                    name = [i[0] for i in t1]
                    cur2 = conn.cursor()
                    cur2.execute(f"SELECT DISTINCT TIMESTAMP FROM {dbch}  ")
                    t2 = cur2.fetchall()
                    str_list = [i[0] for i in t2]
                    df = pd.DataFrame(str_list)
                    df.columns = ['TIMESTAMP']
                    df['TIMESTAMP'] = df.TIMESTAMP.apply(pd.to_datetime)
                    df["day"] = df.TIMESTAMP.dt.day
                    df["month"] = df.TIMESTAMP.dt.month
                    df["year"] = df.TIMESTAMP.dt.year
                    a = [str(i) + '-' + str(j) + '-' + str(k) for i, j, k in zip(df["year"], df["month"], df["day"])]
                    a = list(set(a))
                    b = pd.to_datetime(a)
                    b = sorted(b)
                    str_list = [t.strftime("%Y-%m-%d") for t in b]
                    return [{'label': i, 'value': i} for i in name], [{'label': i, 'value': i} for i in str_list]
            elif dbch == None:
                return no_update, no_update
        if dbname == 'enerbat':
            if dbch != None:
                cur1 = conn.cursor()
                cur1.execute(f"SELECT DISTINCT VARIABLE_NAME FROM {dbch} ")
                t1 = cur1.fetchall()
                name = [i[0] for i in t1]
                cur2 = conn.cursor()
                cur2.execute(f"SELECT DISTINCT TIMESTAMP FROM {dbch}  ")
                t2 = cur2.fetchall()
                str_list = [i[0] for i in t2]
                df = pd.DataFrame(str_list)
                df.columns = ['TIMESTAMP']
                df['TIMESTAMP'] = df.TIMESTAMP.apply(pd.to_datetime)
                df["day"] = df.TIMESTAMP.dt.day
                df["month"] = df.TIMESTAMP.dt.month
                df["year"] = df.TIMESTAMP.dt.year
                a = [str(i) + '-' + str(j) + '-' + str(k) for i, j, k in zip(df["year"], df["month"], df["day"])]
                a = list(set(a))
                b = pd.to_datetime(a)
                b = sorted(b)
                str_list = [t.strftime("%Y-%m-%d") for t in b]
                return [{'label': i, 'value': i} for i in name], [{'label': i, 'value': i} for i in str_list]
            else:
                return no_update, no_update
        else : return no_update, no_update
    if q1 == 'deactivatedb':
        kk = [{'label': i, 'value': i} for i in '']
        return [{'label': i, 'value': i} for i in ''], [{'label': i, 'value': i} for i in '']
    else:
        return no_update, no_update


@app.callback([Output('prvalname', 'options'), Output('prvaldate', 'options')],
              [Input('my-toggle-switch-pr-db', 'on'),Input('interval_component_pr_db', 'n_intervals'),Input('ok_reel_pr', 'n_clicks')],
              [State('prvalchoosen', 'value'), State('prname', 'value'),State('username_pr', 'value'), State('password_pr', 'value')])
def prname(on,interval,ok,  prch, prname, user, passw):
    if prname == None :
        raise PreventUpdate
    ipadress = "193.54.2.211"
    q1 = dash.callback_context.triggered[0]["prop_id"].split(".")[0]
    if q1 == 'ok_reel_pr' :
        if user != None or passw != None :
            user = user
            password = passw
    if on == 1:
        try:
            conn = mariadb.connect(
                user='dashapp' if user == None else user,
                password='dashapp' if passw == None else passw,
                host=ipadress,
                port=3306,
                database=prname)

        except mariadb.Error as e:
            print(f"Error connecting to MariaDB Platform: {e}")
            sys.exit(1)

        if prname == 'rcckn':
            if prch == 'received_variablevalues':
                cur1 = conn.cursor()
                cur1.execute("SELECT DISTINCT VARIABLE_NAME FROM received_variablevalues ")
                t1 = cur1.fetchall()
                name = [i[0] for i in t1]
                cur2 = conn.cursor()
                cur2.execute("SELECT DISTINCT REMOTE_TIMESTAMP FROM received_variablevalues ")
                t2 = cur2.fetchall()

                str_list = [i[0] for i in t2]
                df = pd.DataFrame(str_list)
                df.columns = ['REMOTE_TIMESTAMP']
                df['REMOTE_TIMESTAMP'] = df.REMOTE_TIMESTAMP.apply(pd.to_datetime)
                df["day"] = df.REMOTE_TIMESTAMP.dt.day
                df["month"] = df.REMOTE_TIMESTAMP.dt.month
                df["year"] = df.REMOTE_TIMESTAMP.dt.year
                a = [str(i) + '-' + str(j) + '-' + str(k) for i, j, k in zip(df["year"], df["month"], df["day"])]
                a = list(set(a))
                b = pd.to_datetime(a)
                b = sorted(b)
                str_list = [t.strftime("%Y-%m-%d") for t in b]
                return [{'label': i, 'value': i} for i in name], [{'label': i, 'value': i} for i in str_list]
            elif prch == "send_variablevalues" :
                cur1 = conn.cursor()
                cur1.execute("SELECT DISTINCT VARIABLE_NAME FROM send_variablevalues ")
                t1 = cur1.fetchall()
                name = [i[0] for i in t1]
                cur2 = conn.cursor()
                cur2.execute("SELECT DISTINCT TIMESTAMP FROM send_variablevalues ")
                t2 = cur2.fetchall()

                str_list = [i[0] for i in t2]
                df = pd.DataFrame(str_list)
                df.columns = ['TIMESTAMP']
                df['TIMESTAMP'] = df.TIMESTAMP.apply(pd.to_datetime)
                df["day"] = df.TIMESTAMP.dt.day
                df["month"] = df.TIMESTAMP.dt.month
                df["year"] = df.TIMESTAMP.dt.year
                a = [str(i) + '-' + str(j) + '-' + str(k) for i, j, k in zip(df["year"], df["month"], df["day"])]
                a = list(set(a))
                b = pd.to_datetime(a)
                b = sorted(b)
                str_list = [t.strftime("%Y-%m-%d") for t in b]
                return [{'label': i, 'value': i} for i in name], [{'label': i, 'value': i} for i in str_list]
            elif prch == None: raise PreventUpdate
            else:
                cur1 = conn.cursor()
                cur1.execute(f"SELECT DISTINCT VARIABLE_NAME FROM {prch} ")
                t1 = cur1.fetchall()
                name = [i[0] for i in t1]
                cur2 = conn.cursor()
                cur2.execute(f"SELECT DISTINCT TIMESTAMP FROM {prch}  ")
                t2 = cur2.fetchall()

                str_list = [i[0] for i in t2]
                df = pd.DataFrame(str_list)
                if df.empty !=1 :
                    df.columns = ['TIMESTAMP']
                    df['TIMESTAMP'] = df.TIMESTAMP.apply(pd.to_datetime)
                    df["day"] = df.TIMESTAMP.dt.day
                    df["month"] = df.TIMESTAMP.dt.month
                    df["year"] = df.TIMESTAMP.dt.year
                    a = [str(i) + '-' + str(j) + '-' + str(k) for i, j, k in zip(df["year"], df["month"], df["day"])]
                    a = list(set(a))
                    b = pd.to_datetime(a)
                    b = sorted(b)
                    str_list = [t.strftime("%Y-%m-%d") for t in b]
                    return [{'label': i, 'value': i} for i in name], [{'label': i, 'value': i} for i in str_list]
                else : raise PreventUpdate

        if prname == 'enerbat':
            if prch != None:
                cur1 = conn.cursor()
                cur1.execute(f"SELECT DISTINCT VARIABLE_NAME FROM {prch} ")
                t1 = cur1.fetchall()
                name = [i[0] for i in t1]
                cur2 = conn.cursor()
                cur2.execute(f"SELECT DISTINCT TIMESTAMP FROM {prch}  ")
                t2 = cur2.fetchall()

                str_list = [i[0] for i in t2]
                df = pd.DataFrame(str_list)
                df.columns = ['TIMESTAMP']
                df['TIMESTAMP'] = df.TIMESTAMP.apply(pd.to_datetime)
                df["day"] = df.TIMESTAMP.dt.day
                df["month"] = df.TIMESTAMP.dt.month
                df["year"] = df.TIMESTAMP.dt.year
                a = [str(i) + '-' + str(j) + '-' + str(k) for i, j, k in zip(df["year"], df["month"], df["day"])]
                a = list(set(a))
                b = pd.to_datetime(a)
                b = sorted(b)
                str_list = [t.strftime("%Y-%m-%d") for t in b]
                return [{'label': i, 'value': i} for i in name], [{'label': i, 'value': i} for i in str_list]
            elif prch == None:
                raise PreventUpdate
        else:
            no_update, no_update
    if on == 0:
        kk = [{'label': i, 'value': i} for i in '']
        return [{'label': i, 'value': i} for i in ''], [{'label': i, 'value': i} for i in '']
    else:
        no_update, no_update

@app.callback(ServersideOutput('memory-output', 'data'),
              [Input('dbvalname', 'value'), Input('dbvaldate', 'value')],
              [State('dbvalchoosen', 'value'), State('db_name', 'value'), State('db_Ip', 'value'),State('db_password', 'value')])
def dbname(valname, valdate, dbch, dbname, ipval, password):
    if dbname == None or valname == None or valdate == None:
        raise PreventUpdate
    ipadress = "193.54.2.211"
    server = SSHTunnelForwarder(
        (ipadress, 22),
        ssh_username='soudani',
        ssh_password="univ484067152",
        remote_bind_address=(ipadress, 3306))

    server.start()


    try:
        conn = mariadb.connect(
            user=ipval,
            password=password,
            host=ipadress,
            port=3306,
            database=dbname)

    except mariadb.Error as e:
        print(f"Error connecting to MariaDB Platform: {e}")
        sys.exit(1)
        # Get Cursor

    if dbname == 'rcckn':
        if dbch == 'received_variablevalues':
            cur1 = conn.cursor()
            if len(valname) == 1:
                cur1.execute(f"SELECT * FROM received_variablevalues WHERE VARIABLE_NAME = '{valname[0]}'")
            elif len(valname) > 1:
                valname = tuple(valname)
                cur1.execute(f"SELECT * FROM received_variablevalues WHERE VARIABLE_NAME IN {valname}")
            else: return no_update
            t1 = cur1.fetchall()
            return t1
        elif dbch == "send_variablevalues":
            cur1 = conn.cursor()
            if len(valname) == 1:
                cur1.execute(f"SELECT * FROM send_variablevalues WHERE VARIABLE_NAME = '{valname[0]}'")
            elif len(valname) > 1:
                valname = tuple(valname)
                cur1.execute(f"SELECT * FROM send_variablevalues WHERE VARIABLE_NAME IN {valname}")
            else: return no_update
            t1 = cur1.fetchall()
            return t1
        elif dbch != "send_variablevalues" or dbch != "received_variablevalues":
            cur1 = conn.cursor()
            if len(valname) == 1:
                cur1.execute(f"SELECT * FROM {dbch} WHERE VARIABLE_NAME = '{valname[0]}'")
            elif len(valname) > 1:
                valname = tuple(valname)
                cur1.execute(f"SELECT * FROM {dbch} WHERE VARIABLE_NAME IN {valname}")
            else  : return no_update
            t1 = cur1.fetchall()
            return t1
    if dbname == 'enerbat':
        if dbch != None:
            cur1 = conn.cursor()
            if len(valname) == 1:
                cur1.execute(f"SELECT * FROM {dbch} WHERE VARIABLE_NAME = '{valname[0]}'")
            elif len(valname) > 1:
                valname = tuple(valname)
                cur1.execute(f"SELECT * FROM {dbch} WHERE VARIABLE_NAME IN {valname}")
            else: return no_update
            t1 = cur1.fetchall()
            return t1
@app.callback(ServersideOutput('memory-outputpr', 'data'),
              [Input('prvalname', 'value'), Input('prvaldate', 'value')],
              [State('prvalchoosen', 'value'), State('prname', 'value'),
               State('username_pr', 'value'), State('password_pr', 'value')])
def dbname(valname, valdate, dbch, dbname, user, passw):
    if dbname == None or valname == None or valdate == None:
        raise PreventUpdate
    ipadress = "193.54.2.211"
    server = SSHTunnelForwarder(
        (ipadress, 22),
        ssh_username='soudani',
        ssh_password="univ484067152",
        remote_bind_address=(ipadress, 3306))

    server.start()

    try:
        conn = mariadb.connect(
            user='dashapp' if user == None else user,
            password='dashapp' if passw == None else passw,
            host=ipadress,
            port=3306,
            database=dbname)

    except mariadb.Error as e:
        print(f"Error connecting to MariaDB Platform: {e}")
        sys.exit(1)
        # Get Cursor

    if dbname == 'rcckn':
        if dbch == 'received_variablevalues':
            cur1 = conn.cursor()
            if len(valname) == 1:
                cur1.execute(f"SELECT * FROM received_variablevalues WHERE VARIABLE_NAME = '{valname[0]}'")
            elif len(valname) > 1:
                valname = tuple(valname)
                cur1.execute(f"SELECT * FROM received_variablevalues WHERE VARIABLE_NAME IN {valname}")
            else:
                return no_update
            t1 = cur1.fetchall()
            return t1
        elif dbch == "send_variablevalues":
            cur1 = conn.cursor()
            if len(valname) == 1:
                cur1.execute(f"SELECT * FROM send_variablevalues WHERE VARIABLE_NAME = '{valname[0]}'")
            elif len(valname) > 1:
                valname = tuple(valname)
                cur1.execute(f"SELECT * FROM send_variablevalues WHERE VARIABLE_NAME IN {valname}")
            else:
                return no_update
            t1 = cur1.fetchall()
            return t1
        else:

            cur1 = conn.cursor()
            if len(valname) == 1:
                cur1.execute(f"SELECT * FROM {dbch} WHERE VARIABLE_NAME = '{valname[0]}'")
            elif len(valname) > 1:
                valname = tuple(valname)
                cur1.execute(f"SELECT * FROM {dbch} WHERE VARIABLE_NAME IN {valname}")
            else:
                return no_update
            t1 = cur1.fetchall()
            return t1
    if dbname == 'enerbat':
        if dbch != None:
            cur1 = conn.cursor()
            if len(valname) == 1:
                cur1.execute(f"SELECT * FROM {dbch} WHERE VARIABLE_NAME = '{valname[0]}'")
            elif len(valname) > 1:
                valname = tuple(valname)
                cur1.execute(f"SELECT * FROM {dbch} WHERE VARIABLE_NAME IN {valname}")
            else:
                return no_update
            t1 = cur1.fetchall()
            return t1
@app.callback([Output('getdbtable', 'data'),
               Output('getdbtable', 'columns'), ],
              [Input('memory-output', 'data'), Input('dbvalname', 'value'),
               Input('dbvaldate', 'value'), Input('deactivatedb', 'n_clicks')],
              [State('dbvalchoosen', 'value'), State('db_name', 'value'), ]
              )
def on_data_set_table(data, valname, valdate, nc2, dbch, dbname):
    if data is None or valname == None or valdate == None or dbch == None or dbname == None:
        raise PreventUpdate
    df = pd.DataFrame(data)
    if dbname == 'rcckn':
        if dbch != None and len(df.columns) == 11:
            if valdate != '' or valname != []:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'LOCAL_TIMESTAMP',
                                  'REMOTE_ID', 'REMOTE_TIMESTAMP', 'REMOTE_MESSAGE_ID', 'PROCESSED', 'TIMED_OUT',
                                  'CONVERTED_NUM_VALUE']
                    df['REMOTE_TIMESTAMP'] = df['REMOTE_TIMESTAMP'].astype('string')
                    a = []
                    for col in df['REMOTE_TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdate)):
                        valdate_new.append(valdate[i][:10])
                    df1 = df[df['dates'].isin(valdate_new)]
                    a = df1.loc[df1['VARIABLE_NAME'].isin(valname)]
                    x = a.to_dict('records')
                    return x, [{'name': i, 'id': i} for i in df.columns if i.startswith('Unn') != 1 or i != 'dates']
                else:
                    raise PreventUpdate
            else:
                raise PreventUpdate
        if dbch != None and len(df.columns) == 8:
            if valdate != '' or valname != []:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'TIMESTAMP',
                                  'PROCESSED', 'TIMED_OUT', 'UNREFERENCED']
                    df['TIMESTAMP'] = df['TIMESTAMP'].astype('string')
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdate)):
                        valdate_new.append(valdate[i][:10])
                    df1 = df[df['dates'].isin(valdate_new)]
                    a = df1.loc[df1['VARIABLE_NAME'].isin(valname)]
                    x = a.to_dict('records')
                    return x, [{'name': i, 'id': i} for i in df.columns if i.startswith('Unn') != 1 or i != 'dates']
                else:
                    raise PreventUpdate
            else:
                raise PreventUpdate
        if dbch != None and len(df.columns) == 4:
            if valdate != '' or valname != []:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
                    df['TIMESTAMP'] = df['TIMESTAMP'].astype('string')
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdate)):
                        valdate_new.append(valdate[i][:10])
                    df1 = df[df['dates'].isin(valdate_new)]
                    a = df1.loc[df1['VARIABLE_NAME'].isin(valname)]
                    x = a.to_dict('records')
                    return x, [{'name': i, 'id': i} for i in df.columns if i.startswith('Unn') != 1 or i != 'dates']
                else:
                    raise PreventUpdate
            else:
                raise PreventUpdate
        if dbch == None :
            if valdate != '' or valname != []:
                raise PreventUpdate
            if valdate == '' or valname != []:
                raise PreventUpdate
            if valdate != '' or valname == []:
                raise PreventUpdate
    if dbname == 'enerbat':
        if dbch != None:
            if valdate != '' or valname != []:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
                    df['TIMESTAMP'] = df['TIMESTAMP'].astype('string')
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdate)):
                        valdate_new.append(valdate[i][:10])
                    df1 = df[df['dates'].isin(valdate_new)]
                    a = df1.loc[df1['VARIABLE_NAME'].isin(valname)]
                    x = a.to_dict('records')
                    return x, [{'name': i, 'id': i} for i in df.columns if i.startswith('Unn') != 1 or i != 'dates']
                else:
                    raise PreventUpdate
@app.callback([Output('getrealtable', 'data'),
               Output('getrealtable', 'columns'), ],
              [Input('get_data_from_modbus', 'data')]
              )
def on_data_set_table(data):
    if data == None :
        raise PreventUpdate
    if data != None :
        df = pd.DataFrame(data, columns = ['VARIABLE_NAME', 'VARIABLE_NUM_VALUE','QUALITY', 'TIMESTAMP'])
        x = df.to_dict('records')
        return x, [{'name': i, 'id': i} for i in df.columns]


@app.callback([Output('getprtable', 'data'),
               Output('getprtable', 'columns'), ],
              [ Input('memory-outputpr', 'data'), Input('prvalname', 'value'),
               Input('prvaldate', 'value'),Input('interval_component_pr_db', 'n_intervals')],
              [State('prvalchoosen', 'value'), State('prname', 'value'), ]
              )
def on_data_set_tablepr(data, valname, valdate,interval, prch, prname):
    if valname == None or valdate == None or prch == None or prname == None :
        raise PreventUpdate
    df = pd.DataFrame(data)
    if prname == 'rcckn':
        if prch == 'received_variablevalues':
            if valdate != '' or valname != []:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'LOCAL_TIMESTAMP',
                                  'REMOTE_ID', 'REMOTE_TIMESTAMP', 'REMOTE_MESSAGE_ID', 'PROCESSED', 'TIMED_OUT',
                                  'CONVERTED_NUM_VALUE']
                    df['REMOTE_TIMESTAMP'] = df['REMOTE_TIMESTAMP'].astype('string')
                    a = []
                    for col in df['REMOTE_TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdate)):
                        valdate_new.append(valdate[i][:10])
                    df1 = df[df['dates'].isin(valdate_new)]
                    a = df1.loc[df1['VARIABLE_NAME'].isin(valname)]
                    x = a.to_dict('records')
                    return x, [{'name': i, 'id': i} for i in df.columns if i.startswith('Unn') != 1 or i != 'dates']
                else:
                    raise PreventUpdate
            else:
                raise PreventUpdate
        if prch == 'send_variablevalues':
            if valdate != '' or valname != []:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'TIMESTAMP',
                                  'PROCESSED', 'TIMED_OUT', 'UNREFERENCED']
                    df['TIMESTAMP'] = df['TIMESTAMP'].astype('string')
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdate)):
                        valdate_new.append(valdate[i][:10])
                    df1 = df[df['dates'].isin(valdate_new)]
                    a = df1.loc[df1['VARIABLE_NAME'].isin(valname)]
                    x = a.to_dict('records')
                    return x, [{'name': i, 'id': i} for i in df.columns if i.startswith('Unn') != 1 or i != 'dates']
                else:
                    raise PreventUpdate
        if prch != 'send_variablevalues' or prch != 'received_variablevalues':
            if valdate != '' or valname != []:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
                    df['TIMESTAMP'] = df['TIMESTAMP'].astype('string')
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdate)):
                        valdate_new.append(valdate[i][:10])
                    df1 = df[df['dates'].isin(valdate_new)]
                    a = df1.loc[df1['VARIABLE_NAME'].isin(valname)]
                    x = a.to_dict('records')
                    return x, [{'name': i, 'id': i} for i in df.columns if i.startswith('Unn') != 1 or i != 'dates']

                else:
                    raise PreventUpdate
            else:
                raise PreventUpdate
    if prname == 'enerbat':
        if prch != None:
            if valdate != '' or valname != []:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
                    df['TIMESTAMP'] = df['TIMESTAMP'].astype('string')
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdate)):
                        valdate_new.append(valdate[i][:10])
                    df1 = df[df['dates'].isin(valdate_new)]
                    a = df1.loc[df1['VARIABLE_NAME'].isin(valname)]
                    x = a.to_dict('records')
                    return x, [{'name': i, 'id': i} for i in df.columns if i.startswith('Unn') != 1 or i != 'dates']
                else:
                    raise PreventUpdate


@app.callback(
    [Output('firstChoosenValuedb', 'options'),
     Output('secondChoosenValuedb', 'options')],
    [Input('dbvalname', 'value')], )
def containerdb(val1):
    if val1 == None or val1 == []:
        raise PreventUpdate

    return [{'label': i, 'value': i} for i in val1], [{'label': i, 'value': i} for i in val1]


@app.callback(Output('getdbgraph', 'figure'),
              [Input("selectyaxisdb", "on"),
               Input('memory-output', 'data'),
               Input('dbvalname', 'value'),
               Input('dbvaldate', 'value'),
               Input('sliderWidthdb', 'value'),
               Input('sliderHeightdb', 'value'),
               Input('radiographdb', 'value'),
               Input('hiddenTextxaxisdb', 'children'),
               Input('hiddenTextyaxisdb', 'children'),
               Input('hiddenTexty1axisdb', 'children'),
               Input('hiddenTexty2axisdb', 'children'),
               Input('hiddenTextHeaderdb', 'children'),
               Input('hiddenTextNotedb', 'children'),
               Input('characteredb', 'value'),
               ],
              [State('dbvalchoosen', 'value'), State('db_name', 'value'), ])
def on_data_set_graph(on, data, valy, valdat, sliderw, sliderh,radio,g1, g2, y1,y2,head, note,car, dbch, dbname):
    if data == None or valy == [] or valdat == [] or valdat == None:
        raise PreventUpdate
    df = pd.DataFrame(data)
    fig = go.Figure()
    if dbname == 'rcckn':
        if dbch == 'received_variablevalues' or len(df.columns) == 11:
            if df.empty != 1:
                df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'LOCAL_TIMESTAMP',
                              'REMOTE_ID', 'REMOTE_TIMESTAMP', 'REMOTE_MESSAGE_ID', 'PROCESSED', 'TIMED_OUT', 'CONVERTED_NUM_VALUE']
                a = []
                for col in df['REMOTE_TIMESTAMP']:
                    a.append(col[:10])
                df['dates'] = a
                valdate_new = []
                for i in range(len(valdat)):
                    valdate_new.append(valdat[i][:10])
                for j in range(len(valy)):
                    for k in range(len(valdate_new)):
                        a = df[df['VARIABLE_NAME'] == valy[j]]
                        a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                        b = df[df['VARIABLE_NAME'] == valy[j]]['REMOTE_TIMESTAMP']
                        b = [i for i in b if i.startswith(valdate_new[k])]
                        time.sleep(1)
                        if on == 1:
                            if len(valy) == 1:
                                fig.add_trace(go.Scattergl(x=b, y=a, mode=radio, marker=dict(line=dict(width=0.2, color='blue')),
                                         name="{}".format(valy[j]))),
                            elif len(valy) > 1:
                                fig.add_trace(
                                    go.Scattergl(x=b, y=a, mode=radio, marker=dict(line=dict(width=0.2, color='blue')),
                                                 name="{}".format(valy[j]), yaxis = f'y{j+1}')),
                        else :
                            fig.add_trace(
                                go.Scattergl(x=b, y=a, mode=radio, marker=dict(line=dict(width=0.2, color='blue')),
                                             name="{}".format(valy[j]))),

                    fig.update_xaxes(
                        tickfont_size=car,
                        tickangle=90,
                        title_text='' if g1 == [] else g1[-1],
                        title_standoff=25, title_font={"size": car}, ),

                    fig.update_yaxes(
                        tickfont_size=car,
                        title_text='' if g2 == [] else g2[-1],
                        title_standoff=25,
                        titlefont=dict(
                            color="#636EFA",
                            size=car
                        ),
                        tickfont=dict(
                            color="#636EFA"
                        )),

                    fig.update_layout(
                        title_text=head[-1] if len(head) > 0 else "{}".format(valy[0]),
                        title_font={"size": car},
                        width=sliderw,
                        height=sliderh,
                        legend=dict(

                            traceorder="normal",
                            font=dict(
                                family="sans-serif",
                                size=car,
                                color=colors['figure_text']
                            ),
                            yanchor="bottom",
                            y=-0.7,
                            xanchor="right",
                            x=1,
                            bgcolor=colors['background'],
                            borderwidth=1,
                            orientation='h',
                        ),
                        yaxis2=dict(
                            title_text='y2_axis' if y1 == [] else y1[-1],
                            titlefont=dict(
                                color="#EF553B",
                                size=car
                            ),
                            tickfont=dict(
                                color="#EF553B"
                            ),
                            title_standoff=10,
                            anchor="free",
                            overlaying="y",
                            side="right",
                            position=1

                        ),

                        yaxis3=dict(
                            title_text='y3_axis' if y2 == [] else y2[-1],
                            titlefont=dict(
                                color="#00CC96",
                                size=car
                            ),
                            tickfont=dict(
                                color="#00CC96",
                            ),
                            title_standoff=10,
                            anchor="free",
                            overlaying="y",
                            side="right",
                            position=0.87
                        ),

                        margin=dict(
                            l=50,
                            r=50,
                            b=100,
                            t=50,
                            pad=4

                        ),
                        paper_bgcolor="LightSteelBlue",
                        plot_bgcolor=colors['background'], ),
                    note_new = []
                    if note != None and len(note) > 0:
                        note = note[-1].split(' ')
                    x = ''
                    for i in note:
                        note_new.append(i)
                        if len(note_new) % 4 == 0:
                            note_new[-1] += '<br>'
                    x += ' '.join(note_new)
                    note = [x]
                    fig.add_annotation(text=note[-1] if len(note) > 0 else '', align='left',
                                       bordercolor='black',
                                       borderwidth=1, xref="paper", yref="paper", bgcolor="white", font=dict(
                            family="Arial, sans-serif", size=car, ),
                                       x=1, y=1, showarrow=False)
                    fig.update_layout()
                return fig
            else:
                raise PreventUpdate
        elif dbch == 'send_variablevalues' or len(df.columns) == 8:
            if df.empty != 1:
                df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'TIMESTAMP',
                              'PROCESSED', 'TIMED_OUT', 'UNREFERENCED']
                a = []
                for col in df['TIMESTAMP']:
                    a.append(col[:10])
                df['dates'] = a
                valdate_new = []
                for i in range(len(valdat)):
                    valdate_new.append(valdat[i])
                for j in range(len(valy)):
                    for k in range(len(valdate_new)):
                        a = df[df['VARIABLE_NAME'] == valy[j]]
                        a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                        b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                        b = [i for i in b if i.startswith(valdate_new[k])]
                        if on == 1:
                            if len(valy) == 1:
                                fig.add_trace(
                                    go.Scattergl(x=b, y=a, mode=radio, marker=dict(line=dict(width=0.2, color='blue')),
                                                 name="{}".format(valy[j]))),
                            elif len(valy) > 1:
                                fig.add_trace(
                                    go.Scattergl(x=b, y=a, mode=radio, marker=dict(line=dict(width=0.2, color='blue')),
                                                 name="{}".format(valy[j]), yaxis=f'y{j + 1}')),
                        else:
                            fig.add_trace(
                                go.Scattergl(x=b, y=a, mode=radio, marker=dict(line=dict(width=0.2, color='blue')),
                                             name="{}".format(valy[j]))),
                        fig.update_xaxes(
                            tickfont_size=car,
                            tickangle=90,
                            title_text='' if g1 == [] else g1[-1],
                            title_standoff=25, title_font={"size": car}, ),

                        fig.update_yaxes(
                            tickfont_size=car,
                            title_text='' if g2 == [] else g2[-1],
                            title_standoff=25,
                            titlefont=dict(
                                color="#636EFA",
                                size=car
                            ),
                            tickfont=dict(
                                color="#636EFA"
                            )),

                        fig.update_layout(
                            title_text=head[-1] if len(head) > 0 else "{}".format(valy[0]),
                            title_font={"size": car},
                            width=sliderw,
                            height=sliderh,
                            legend=dict(

                                traceorder="normal",
                                font=dict(
                                    family="sans-serif",
                                    size=car,
                                    color=colors['figure_text']
                                ),
                                yanchor="bottom",
                                y=-0.7,
                                xanchor="right",
                                x=1,
                                bgcolor=colors['background'],
                                borderwidth=1,
                                orientation='h',
                            ),
                            yaxis2=dict(
                                title_text='y2_axis' if y1 == [] else y1[-1],
                                titlefont=dict(
                                    color="#EF553B",
                                    size=car
                                ),
                                tickfont=dict(
                                    color="#EF553B"
                                ),
                                title_standoff=10,
                                anchor="free",
                                overlaying="y",
                                side="right",
                                position=1

                            ),

                            yaxis3=dict(
                                title_text='y3_axis' if y2 == [] else y2[-1],
                                titlefont=dict(
                                    color="#00CC96",
                                    size=car
                                ),
                                tickfont=dict(
                                    color="#00CC96",
                                ),
                                title_standoff=10,
                                anchor="free",
                                overlaying="y",
                                side="right",
                                position=0.87
                            ),

                            margin=dict(
                                l=50,
                                r=50,
                                b=100,
                                t=50,
                                pad=4

                            ),
                            paper_bgcolor="LightSteelBlue",
                            plot_bgcolor=colors['background'], ),
                        note_new = []
                        if note != None and len(note) > 0:
                            note = note[-1].split(' ')
                        x = ''
                        for i in note:
                            note_new.append(i)
                            if len(note_new) % 4 == 0:
                                note_new[-1] += '<br>'
                        x += ' '.join(note_new)
                        note = [x]
                        fig.add_annotation(text=note[-1] if len(note) > 0 else '', align='left',
                                           bordercolor='black',
                                           borderwidth=1, xref="paper", yref="paper", bgcolor="white", font=dict(
                                family="Arial, sans-serif", size=car, ),
                                           x=1, y=1, showarrow=False)
                        fig.update_layout()
                return fig
            else:
                raise PreventUpdate
        elif len(df.columns) == 4 :
            if df.empty != 1:
                df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
                a = []
                for col in df['TIMESTAMP']:
                    a.append(col[:10])
                df['dates'] = a
                valdate_new = []
                for i in range(len(valdat)):
                    valdate_new.append(valdat[i])
                for j in range(len(valy)):
                    for k in range(len(valdate_new)):
                        a = df[df['VARIABLE_NAME'] == valy[j]]
                        a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                        b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                        b = [i for i in b if i.startswith(valdate_new[k])]
                        if on == 1:
                            if len(valy) == 1:
                                fig.add_trace(
                                    go.Scattergl(x=b, y=a, mode=radio, marker=dict(line=dict(width=0.2, color='blue')),
                                                 name="{}".format(valy[j]))),
                            elif len(valy) > 1:
                                fig.add_trace(
                                    go.Scattergl(x=b, y=a, mode=radio, marker=dict(line=dict(width=0.2, color='blue')),
                                                 name="{}".format(valy[j]), yaxis=f'y{j + 1}')),
                        else:
                            fig.add_trace(
                                go.Scattergl(x=b, y=a, mode=radio, marker=dict(line=dict(width=0.2, color='blue')),
                                             name="{}".format(valy[j]))),
                        fig.update_xaxes(
                            tickfont_size=car,
                            tickangle=90,
                            title_text='' if g1 == [] else g1[-1],
                            title_standoff=25, title_font={"size": car}, ),

                        fig.update_yaxes(
                            tickfont_size=car,
                            title_text='' if g2 == [] else g2[-1],
                            title_standoff=25,
                            titlefont=dict(
                                color="#636EFA",
                                size=car
                            ),
                            tickfont=dict(
                                color="#636EFA"
                            )),

                        fig.update_layout(
                            title_text=head[-1] if len(head) > 0 else "{}".format(valy[0]),
                            title_font={"size": car},
                            width=sliderw,
                            height=sliderh,
                            legend=dict(

                                traceorder="normal",
                                font=dict(
                                    family="sans-serif",
                                    size=car,
                                    color=colors['figure_text']
                                ),
                                yanchor="bottom",
                                y=-0.7,
                                xanchor="right",
                                x=1,
                                bgcolor=colors['background'],
                                borderwidth=1,
                                orientation='h',
                            ),
                            yaxis2=dict(
                                title_text='y2_axis' if y1 == [] else y1[-1],
                                titlefont=dict(
                                    color="#EF553B",
                                    size=car
                                ),
                                tickfont=dict(
                                    color="#EF553B"
                                ),
                                title_standoff=10,
                                anchor="free",
                                overlaying="y",
                                side="right",
                                position=1

                            ),

                            yaxis3=dict(
                                title_text='y3_axis' if y2 == [] else y2[-1],
                                titlefont=dict(
                                    color="#00CC96",
                                    size=car
                                ),
                                tickfont=dict(
                                    color="#00CC96",
                                ),
                                title_standoff=10,
                                anchor="free",
                                overlaying="y",
                                side="right",
                                position=0.87
                            ),

                            margin=dict(
                                l=50,
                                r=50,
                                b=100,
                                t=50,
                                pad=4

                            ),
                            paper_bgcolor="LightSteelBlue",
                            plot_bgcolor=colors['background'], ),
                        note_new = []
                        if note != None and len(note) > 0:
                            note = note[-1].split(' ')
                        x = ''
                        for i in note:
                            note_new.append(i)
                            if len(note_new) % 4 == 0:
                                note_new[-1] += '<br>'
                        x += ' '.join(note_new)
                        note = [x]
                        fig.add_annotation(text=note[-1] if len(note) > 0 else '', align='left',
                                           bordercolor='black',
                                           borderwidth=1, xref="paper", yref="paper", bgcolor="white", font=dict(
                                family="Arial, sans-serif", size=car, ),
                                           x=1, y=1, showarrow=False)
                        fig.update_layout()
                return fig
            else:
                raise PreventUpdate
        else:
            raise PreventUpdate
    if dbname == 'enerbat':
        if df.empty != 1:
            df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
            a = []
            for col in df['TIMESTAMP']:
                a.append(col[:10])
            df['dates'] = a
            valdate_new = []
            for i in range(len(valdat)):
                valdate_new.append(valdat[i][:10])
            for j in range(len(valy)):
                for k in range(len(valdate_new)):
                    a = df[df['VARIABLE_NAME'] == valy[j]]
                    a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                    b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                    b = [i for i in b if i.startswith(valdate_new[k])]
                    if on == 1:
                        if len(valy) == 1:
                            fig.add_trace(
                                go.Scattergl(x=b, y=a, mode=radio, marker=dict(line=dict(width=0.2, color='blue')),
                                             name="{}".format(valy[j]))),
                        elif len(valy) > 1:
                            fig.add_trace(
                                go.Scattergl(x=b, y=a, mode=radio, marker=dict(line=dict(width=0.2, color='blue')),
                                             name="{}".format(valy[j]), yaxis=f'y{j + 1}')),
                    else:
                        fig.add_trace(
                            go.Scattergl(x=b, y=a, mode=radio, marker=dict(line=dict(width=0.2, color='blue')),
                                         name="{}".format(valy[j]))),
                    fig.update_xaxes(
                        tickfont_size=car,
                        tickangle=90,
                        title_text='' if g1 == [] else g1[-1],
                        title_standoff=25, title_font={"size": car}, ),

                    fig.update_yaxes(
                        tickfont_size=car,
                        title_text='' if g2 == [] else g2[-1],
                        title_standoff=25,
                        titlefont=dict(
                            color="#636EFA",
                            size=car
                        ),
                        tickfont=dict(
                            color="#636EFA"
                        )),

                    fig.update_layout(
                        title_text=head[-1] if len(head) > 0 else "{}".format(valy[0]),
                        title_font={"size": car},
                        width=sliderw,
                        height=sliderh,
                        legend=dict(

                            traceorder="normal",
                            font=dict(
                                family="sans-serif",
                                size=car,
                                color=colors['figure_text']
                            ),
                            yanchor="bottom",
                            y=-0.6,
                            xanchor="right",
                            x=1,
                            bgcolor=colors['background'],
                            borderwidth=1,
                            orientation='h',
                        ),
                        yaxis2=dict(
                            title_text='y2_axis' if y1 == [] else y1[-1],
                            titlefont=dict(
                                color="#EF553B",
                                size=car
                            ),
                            tickfont=dict(
                                color="#EF553B"
                            ),
                            title_standoff=10,
                            anchor="free",
                            overlaying="y",
                            side="right",
                            position=1

                        ),

                        yaxis3=dict(
                            title_text='y3_axis' if y2 == [] else y2[-1],
                            titlefont=dict(
                                color="#00CC96",
                                size=car
                            ),
                            tickfont=dict(
                                color="#00CC96",
                            ),
                            title_standoff=10,
                            anchor="free",
                            overlaying="y",
                            side="right",
                            position=0.87
                        ),

                        margin=dict(
                            l=50,
                            r=50,
                            b=100,
                            t=50,
                            pad=4

                        ),
                        paper_bgcolor="LightSteelBlue",
                        plot_bgcolor=colors['background'], ),
                    note_new = []
                    if note != None and len(note) > 0:
                        note = note[-1].split(' ')
                    x = ''
                    for i in note:
                        note_new.append(i)
                        if len(note_new) % 4 == 0:
                            note_new[-1] += '<br>'
                    x += ' '.join(note_new)
                    note = [x]
                    fig.add_annotation(text=note[-1] if len(note) > 0 else '', align='left',
                                       bordercolor='black',
                                       borderwidth=1, xref="paper", yref="paper", bgcolor="white", font=dict(
                            family="Arial, sans-serif", size=car, ),
                                       x=1, y=1, showarrow=False)
                    fig.update_layout()
            return fig
        else:
            raise PreventUpdate
@app.callback(Output('graphreal', 'figure'),
              [Input('get_data_from_modbus', 'data'),
               Input('realvalue', 'value'),
               Input('sliderWidthreel', 'value'),
               Input('sliderHeightreel', 'value')])
def graphreelTime(data, val, sliderwidth, sliderheight):
    if val == None:
        raise PreventUpdate
    df = pd.DataFrame(data, columns=['ID', 'Value', 'Quality', 'TIMESTAMP'])
    df1 = df[df['ID'].isin(val)]
    a = df1.loc[df1['ID'].isin(val)]
    fig = go.Figure()
    for i in val:
        fig.add_trace(go.Scattergl(x=a[a['ID'] == i]['TIMESTAMP'], y=a[a['ID'] == i]['Value'], mode='lines',
                                   marker=dict(line=dict(width=0.2, color='white')), name="{}".format(i),
                                   ))
    fig.update_layout(

        autosize=False,
        width=sliderwidth,
        height=sliderheight,

        margin=dict(
            l=50,
            r=50,
            b=100,
            t=50,
            pad=4
        ),
        paper_bgcolor="LightSteelBlue",
        template="plotly_white"
    ),

    return fig

@app.callback([Output('firstgraph_pr_real', 'options'), Output('secondgraph_pr_real', 'options'),
               Output('thirdgraph_pr_real', 'options'),Output('fourgraph_pr_real', 'options')],
              [Input('realvalue_pr', 'value')])

def delivre_dropdown(values) :
    if values == None :
        raise PreventUpdate
    return  [{'label' : i[16:], 'value' : i} for i in values], [{'label' :i[16:], 'value' : i} for i in values],[{'label' : i[16:], 'value' : i} for i in values],[{'label' : i[16:], 'value' : i} for i in values]

@app.callback([Output('firstgraph_pr_db', 'options'), Output('secondgraph_pr_db', 'options'),
               Output('thirdgraph_pr_db', 'options'),Output('fourgraph_pr_db', 'options')],
              [Input('prvalname', 'value')])

def delivre_dropdown_db(values) :
    if values == None :
        raise PreventUpdate
    return  [{'label' : i, 'value' : i} for i in values],[{'label' : i, 'value' : i} for i in values],[{'label': i, 'value': i} for i in values],[{'label' : i, 'value' : i} for i in values]

@app.callback(Output('getprgraph', 'figure'),
              [Input('memory-outputpr', 'data'),
                Input('get_data_from_modbus_pr', 'data'),
               Input('firstgraph_pr_real', 'value'),
               Input('firstgraph_pr_db', 'value'),
               Input('prvaldate', 'value'),
               Input('sliderWidthpr', 'value'),
               Input('sliderHeightpr', 'value'),Input('interval_component_pr_db', 'n_intervals') ],
              [State('prvalchoosen', 'value'), State('prname', 'value'),State('prvalname', 'value') ])
def on_data_set_graph(data2,data, realval,valy, valdat, sliderw, sliderh,interval, prch, prname, prvalname):
    if realval == None and valy == None:
        raise PreventUpdate
    df = pd.DataFrame(data2)
    df2 = pd.DataFrame(data, columns=['ID', 'Value', 'Quality', 'TIMESTAMP'])
    fig = go.Figure()
    if realval != None and valy != None:
        for i in realval:
            fig.add_trace(
                go.Scattergl(x=df2[df2['ID'] == i[16:]]['TIMESTAMP'], y=df2[df2['ID'] == i[16:]]['Value'], mode='lines',
                             marker=dict(line=dict(width=0.2, color='white')), name="{}".format(i),
                             ))
        if prname == 'rcckn':
            if prch == 'received_variablevalues' or len(df.columns) == 11:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'LOCAL_TIMESTAMP',
                                  'REMOTE_ID',
                                  'REMOTE_TIMESTAMP', 'REMOTE_MESSAGE_ID', 'PROCESSED', 'TIMED_OUT',
                                  'CONVERTED_NUM_VALUE']
                    a = []
                    for col in df['REMOTE_TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i][:10])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['REMOTE_TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(
                                go.Scattergl(x=b, y=a, mode='markers', marker=dict(line=dict(width=0.2, color='white')),
                                             name="{}/{}".format(valy[j], valdate_new[k]))),
                        fig.update_layout(
                            autosize=True,
                            width=sliderw,
                            height=sliderh,
                            margin=dict(
                                l=100,
                                r=60,
                                b=40,
                                t=40,
                                pad=2
                            ),

                            uirevision=valy[j]),
                    return fig
                else:
                    raise PreventUpdate
            elif prch == 'send_variablevalues' or len(df.columns) == 8:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'TIMESTAMP',
                                  'PROCESSED', 'TIMED_OUT', 'UNREFERENCED']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                       marker=dict(
                                                           line=dict(
                                                               width=0.2,
                                                               color='white')),
                                                       name="{}/{}".format(valy[j], valdate_new[k]))),
                            fig.update_layout(
                                autosize=True,
                                width=sliderw,
                                height=sliderh,
                                margin=dict(
                                    l=100,
                                    r=60,
                                    b=40,
                                    t=40,
                                    pad=2
                                ),

                                uirevision=valy[j], ),
                    return fig
                else:
                    raise PreventUpdate
            else:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                       marker=dict(
                                                           line=dict(
                                                               width=0.2,
                                                               color='white')),
                                                       name="{}/{}".format(valy[j], valdate_new[k]))),
                            fig.update_layout(
                                autosize=True,
                                width=sliderw,
                                height=sliderh,
                                margin=dict(
                                    l=100,
                                    r=60,
                                    b=40,
                                    t=40,
                                    pad=2
                                ),

                                uirevision=valy[j], ),
                    return fig
                else:
                    raise PreventUpdate

        if prname == 'enerbat':
            if df.empty != 1:
                df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
                df = df[df['VARIABLE_NAME'].isin(prvalname)]
                a = []
                for col in df['TIMESTAMP']:
                    a.append(col[:10])
                df['dates'] = a
                valdate_new = []
                for i in range(len(valdat)):
                    valdate_new.append(valdat[i][:10])
                for j in range(len(valy)):
                    for k in range(len(valdate_new)):
                        a = df[df['VARIABLE_NAME'] == valy[j]]
                        a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                        b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                        b = [i for i in b if i.startswith(valdate_new[k])]
                        fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                   marker=dict(
                                                       line=dict(
                                                           width=0.2,
                                                           color='white')),
                                                   name="{}/{}".format(valy[j], valdate_new[k]))),
                        fig.update_layout(
                            autosize=True,
                            width=sliderw,
                            height=sliderh,
                            margin=dict(
                                l=100,
                                r=60,
                                b=40,
                                t=40,
                                pad=2
                            ),

                            uirevision=valy[j], ),
                return fig
    if realval != None:
        for i in realval:
            fig.add_trace(go.Scattergl(x=df2[df2['ID'] == i[16:]]['TIMESTAMP'], y=df2[df2['ID'] == i[16:]]['Value'], mode='lines',
                                       marker=dict(line=dict(width=0.2, color='white')), name="{}".format(i),
                                       ))
            fig.update_layout(
                autosize=True,
                width=sliderw,
                height=sliderh,
                margin=dict(
                    l=100,
                    r=60,
                    b=40,
                    t=40,
                    pad=2
                ),

                uirevision=i[16:]),
            return fig
    if valy != None:
        if prname == 'rcckn':
            if prch == 'received_variablevalues' or len(df.columns) == 11:
                if df.empty != 1:

                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'LOCAL_TIMESTAMP',
                                  'REMOTE_ID',
                                  'REMOTE_TIMESTAMP', 'REMOTE_MESSAGE_ID', 'PROCESSED', 'TIMED_OUT', 'CONVERTED_NUM_VALUE']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['REMOTE_TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i][:10])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['REMOTE_TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(
                                go.Scattergl(x=b, y=a, mode='markers', marker=dict(line=dict(width=0.2, color='white')),
                                             name="{}/{}".format(valy[j], valdate_new[k]))),
                        fig.update_layout(
                            autosize=True,
                            width=sliderw,
                            height=sliderh,
                            margin=dict(
                                l=100,
                                r=60,
                                b=40,
                                t=40,
                                pad=2
                            ),

                            uirevision=valy[j]),
                    return fig
                else:
                    raise PreventUpdate
            elif prch == 'send_variablevalues' or len(df.columns) == 8:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'TIMESTAMP',
                                  'PROCESSED', 'TIMED_OUT', 'UNREFERENCED']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                       marker=dict(
                                                           line=dict(
                                                               width=0.2,
                                                               color='white')),
                                                       name="{}/{}".format(valy[j], valdate_new[k]))),
                            fig.update_layout(
                                autosize=True,
                                width=sliderw,
                                height=sliderh,
                                margin=dict(
                                    l=100,
                                    r=60,
                                    b=40,
                                    t=40,
                                    pad=2
                                ),

                                uirevision=valy[j], ),
                    return fig
                else:
                    raise PreventUpdate
            elif len(df.columns) == 4 :
                if df.empty != 1:

                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE',  'TIMESTAMP']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                       marker=dict(
                                                           line=dict(
                                                               width=0.2,
                                                               color='white')),
                                                       name="{}/{}".format(valy[j], valdate_new[k]))),
                            fig.update_layout(
                                autosize=True,
                                width=sliderw,
                                height=sliderh,
                                margin=dict(
                                    l=100,
                                    r=60,
                                    b=40,
                                    t=40,
                                    pad=2
                                ),

                                uirevision=valy[j], ),
                    return fig
                else:
                    raise PreventUpdate

        if prname == 'enerbat':
            if df.empty != 1:
                df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
                df = df[df['VARIABLE_NAME'].isin(prvalname)]
                #
                a = []
                for col in df['TIMESTAMP']:
                    a.append(col[:10])
                df['dates'] = a
                valdate_new = []
                for i in range(len(valdat)):
                    valdate_new.append(valdat[i][:10])
                for j in range(len(valy)):
                    for k in range(len(valdate_new)):
                        a = df[df['VARIABLE_NAME'] == valy[j]]
                        a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                        b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                        b = [i for i in b if i.startswith(valdate_new[k])]
                        fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                   marker=dict(
                                                       line=dict(
                                                           width=0.2,
                                                           color='white')),
                                                   name="{}/{}".format(valy[j], valdate_new[k]))),
                        fig.update_layout(

                            width=sliderw,
                            height=sliderh,
                            margin=dict(
                                l=100,
                                r=60,
                                b=40,
                                t=40,
                                pad=2
                            ),

                            uirevision=valy[j], ),
                return fig
            else :
                raise PreventUpdate


@app.callback(Output('getprgraph2', 'figure'),
              [Input('memory-outputpr', 'data'),Input('get_data_from_modbus_pr', 'data'),
               Input('secondgraph_pr_real', 'value'),
               Input('secondgraph_pr_db', 'value'),
               Input('prvaldate', 'value'),
               Input('sliderWidthpr2', 'value'),
               Input('sliderHeightpr2', 'value'),Input('interval_component_pr_db', 'n_intervals') ],
              [State('prvalchoosen', 'value'), State('prname', 'value'),State('prvalname', 'value')  ])
def on_data_set_graph2(data2,data, realval,valy, valdat, sliderw, sliderh,interval, prch, prname,prvalname):
    if realval == None and valy == None:
        raise PreventUpdate
    df = pd.DataFrame(data2)
    df2 = pd.DataFrame(data, columns=['ID', 'Value', 'Quality', 'TIMESTAMP'])
    fig = go.Figure()
    if realval != None and valy != None:
        for i in realval:
            fig.add_trace(
                go.Scattergl(x=df2[df2['ID'] == i[16:]]['TIMESTAMP'], y=df2[df2['ID'] == i[16:]]['Value'], mode='lines',
                             marker=dict(line=dict(width=0.2, color='white')), name="{}".format(i),
                             ))
        if prname == 'rcckn':
            if prch == 'received_variablevalues' or len(df.columns) == 11:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'LOCAL_TIMESTAMP',
                                  'REMOTE_ID',
                                  'REMOTE_TIMESTAMP', 'REMOTE_MESSAGE_ID', 'PROCESSED', 'TIMED_OUT',
                                  'CONVERTED_NUM_VALUE']
                    a = []
                    for col in df['REMOTE_TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i][:10])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['REMOTE_TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(
                                go.Scattergl(x=b, y=a, mode='markers', marker=dict(line=dict(width=0.2, color='white')),
                                             name="{}/{}".format(valy[j], valdate_new[k]))),
                        fig.update_layout(
                            autosize=True,
                            width=sliderw,
                            height=sliderh,
                            margin=dict(
                                l=100,
                                r=60,
                                b=40,
                                t=40,
                                pad=2
                            ),

                            uirevision=valy[j]),
                    return fig
                else:
                    raise PreventUpdate
            elif prch == 'send_variablevalues' or len(df.columns) == 8:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'TIMESTAMP',
                                  'PROCESSED', 'TIMED_OUT', 'UNREFERENCED']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                       marker=dict(
                                                           line=dict(
                                                               width=0.2,
                                                               color='white')),
                                                       name="{}/{}".format(valy[j], valdate_new[k]))),
                            fig.update_layout(
                                autosize=True,
                                width=sliderw,
                                height=sliderh,
                                margin=dict(
                                    l=100,
                                    r=60,
                                    b=40,
                                    t=40,
                                    pad=2
                                ),

                                uirevision=valy[j], ),
                    return fig
                else:
                    raise PreventUpdate
            else:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                       marker=dict(
                                                           line=dict(
                                                               width=0.2,
                                                               color='white')),
                                                       name="{}/{}".format(valy[j], valdate_new[k]))),
                            fig.update_layout(
                                autosize=True,
                                width=sliderw,
                                height=sliderh,
                                margin=dict(
                                    l=100,
                                    r=60,
                                    b=40,
                                    t=40,
                                    pad=2
                                ),

                                uirevision=valy[j], ),
                    return fig
                else:
                    raise PreventUpdate

        if prname == 'enerbat':
            if df.empty != 1:
                df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
                df = df[df['VARIABLE_NAME'].isin(prvalname)]
                a = []
                for col in df['TIMESTAMP']:
                    a.append(col[:10])
                df['dates'] = a
                valdate_new = []
                for i in range(len(valdat)):
                    valdate_new.append(valdat[i][:10])
                for j in range(len(valy)):
                    for k in range(len(valdate_new)):
                        a = df[df['VARIABLE_NAME'] == valy[j]]
                        a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                        b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                        b = [i for i in b if i.startswith(valdate_new[k])]
                        fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                   marker=dict(
                                                       line=dict(
                                                           width=0.2,
                                                           color='white')),
                                                   name="{}/{}".format(valy[j], valdate_new[k]))),
                        fig.update_layout(
                            autosize=True,
                            width=sliderw,
                            height=sliderh,
                            margin=dict(
                                l=100,
                                r=60,
                                b=40,
                                t=40,
                                pad=2
                            ),

                            uirevision=valy[j], ),
                return fig
    if realval != None:
        for i in realval:
            fig.add_trace(
                go.Scattergl(x=df2[df2['ID'] == i[16:]]['TIMESTAMP'], y=df2[df2['ID'] == i[16:]]['Value'], mode='lines',
                             marker=dict(line=dict(width=0.2, color='white')), name="{}".format(i),
                             ))
            fig.update_layout(
                autosize=True,
                width=sliderw,
                height=sliderh,
                margin=dict(
                    l=100,
                    r=60,
                    b=40,
                    t=40,
                    pad=2
                ),

                uirevision=i[16:]),
            return fig
    if valy != None:
        if prname == 'rcckn':
            if prch == 'received_variablevalues' or len(df.columns) == 11:
                if df.empty != 1:

                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'LOCAL_TIMESTAMP',
                                  'REMOTE_ID',
                                  'REMOTE_TIMESTAMP', 'REMOTE_MESSAGE_ID', 'PROCESSED', 'TIMED_OUT',
                                  'CONVERTED_NUM_VALUE']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['REMOTE_TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i][:10])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['REMOTE_TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(
                                go.Scattergl(x=b, y=a, mode='markers', marker=dict(line=dict(width=0.2, color='white')),
                                             name="{}/{}".format(valy[j], valdate_new[k]))),
                        fig.update_layout(
                            autosize=True,
                            width=sliderw,
                            height=sliderh,
                            margin=dict(
                                l=100,
                                r=60,
                                b=40,
                                t=40,
                                pad=2
                            ),

                            uirevision=valy[j]),
                    return fig
                else:
                    raise PreventUpdate
            elif prch == 'send_variablevalues' or len(df.columns) == 8:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'TIMESTAMP',
                                  'PROCESSED', 'TIMED_OUT', 'UNREFERENCED']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                       marker=dict(
                                                           line=dict(
                                                               width=0.2,
                                                               color='white')),
                                                       name="{}/{}".format(valy[j], valdate_new[k]))),
                            fig.update_layout(
                                autosize=True,
                                width=sliderw,
                                height=sliderh,
                                margin=dict(
                                    l=100,
                                    r=60,
                                    b=40,
                                    t=40,
                                    pad=2
                                ),

                                uirevision=valy[j], ),
                    return fig
                else:
                    raise PreventUpdate
            elif len(df.columns) == 4:
                if df.empty != 1:

                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                       marker=dict(
                                                           line=dict(
                                                               width=0.2,
                                                               color='white')),
                                                       name="{}/{}".format(valy[j], valdate_new[k]))),
                            fig.update_layout(
                                autosize=True,
                                width=sliderw,
                                height=sliderh,
                                margin=dict(
                                    l=100,
                                    r=60,
                                    b=40,
                                    t=40,
                                    pad=2
                                ),

                                uirevision=valy[j], ),
                    return fig
                else:
                    raise PreventUpdate

        if prname == 'enerbat':
            if df.empty != 1:
                df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
                df = df[df['VARIABLE_NAME'].isin(prvalname)]
                a = []
                for col in df['TIMESTAMP']:
                    a.append(col[:10])
                df['dates'] = a
                valdate_new = []
                for i in range(len(valdat)):
                    valdate_new.append(valdat[i][:10])
                for j in range(len(valy)):
                    for k in range(len(valdate_new)):
                        a = df[df['VARIABLE_NAME'] == valy[j]]
                        a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                        b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                        b = [i for i in b if i.startswith(valdate_new[k])]
                        fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                   marker=dict(
                                                       line=dict(
                                                           width=0.2,
                                                           color='white')),
                                                   name="{}/{}".format(valy[j], valdate_new[k]))),
                        fig.update_layout(

                            width=sliderw,
                            height=sliderh,
                            margin=dict(
                                l=100,
                                r=60,
                                b=40,
                                t=40,
                                pad=2
                            ),

                            uirevision=valy[j], ),
                return fig
            else:
                raise PreventUpdate


@app.callback(Output('getprgraph3', 'figure'),
              [Input('memory-outputpr', 'data'),Input('get_data_from_modbus_pr', 'data'),
               Input('thirdgraph_pr_real', 'value'),
               Input('thirdgraph_pr_db', 'value'),
               Input('prvaldate', 'value'),
               Input('sliderWidthpr3', 'value'),
               Input('sliderHeightpr3', 'value'),Input('interval_component_pr_db', 'n_intervals') ],
              [State('prvalchoosen', 'value'), State('prname', 'value'),State('prvalname', 'value')  ])
def on_data_set_graph3(data2,data, realval,valy, valdat, sliderw, sliderh,interval, prch, prname,prvalname):
    if realval == None and valy == None:
        raise PreventUpdate
    df = pd.DataFrame(data2)
    df2 = pd.DataFrame(data, columns=['ID', 'Value', 'Quality', 'TIMESTAMP'])
    fig = go.Figure()
    if realval != None and valy != None:
        for i in realval:
            fig.add_trace(
                go.Scattergl(x=df2[df2['ID'] == i[16:]]['TIMESTAMP'], y=df2[df2['ID'] == i[16:]]['Value'], mode='lines',
                             marker=dict(line=dict(width=0.2, color='white')), name="{}".format(i),
                             ))
        if prname == 'rcckn':
            if prch == 'received_variablevalues' or len(df.columns) == 11:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'LOCAL_TIMESTAMP',
                                  'REMOTE_ID',
                                  'REMOTE_TIMESTAMP', 'REMOTE_MESSAGE_ID', 'PROCESSED', 'TIMED_OUT',
                                  'CONVERTED_NUM_VALUE']
                    a = []
                    for col in df['REMOTE_TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i][:10])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['REMOTE_TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(
                                go.Scattergl(x=b, y=a, mode='markers', marker=dict(line=dict(width=0.2, color='white')),
                                             name="{}/{}".format(valy[j], valdate_new[k]))),
                        fig.update_layout(
                            autosize=True,
                            width=sliderw,
                            height=sliderh,
                            margin=dict(
                                l=100,
                                r=60,
                                b=40,
                                t=40,
                                pad=2
                            ),

                            uirevision=valy[j]),
                    return fig
                else:
                    raise PreventUpdate
            elif prch == 'send_variablevalues' or len(df.columns) == 8:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'TIMESTAMP',
                                  'PROCESSED', 'TIMED_OUT', 'UNREFERENCED']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                       marker=dict(
                                                           line=dict(
                                                               width=0.2,
                                                               color='white')),
                                                       name="{}/{}".format(valy[j], valdate_new[k]))),
                            fig.update_layout(
                                autosize=True,
                                width=sliderw,
                                height=sliderh,
                                margin=dict(
                                    l=100,
                                    r=60,
                                    b=40,
                                    t=40,
                                    pad=2
                                ),

                                uirevision=valy[j], ),
                    return fig
                else:
                    raise PreventUpdate
            else:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                       marker=dict(
                                                           line=dict(
                                                               width=0.2,
                                                               color='white')),
                                                       name="{}/{}".format(valy[j], valdate_new[k]))),
                            fig.update_layout(
                                autosize=True,
                                width=sliderw,
                                height=sliderh,
                                margin=dict(
                                    l=100,
                                    r=60,
                                    b=40,
                                    t=40,
                                    pad=2
                                ),

                                uirevision=valy[j], ),
                    return fig
                else:
                    raise PreventUpdate

        if prname == 'enerbat':
            if df.empty != 1:
                df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
                df = df[df['VARIABLE_NAME'].isin(prvalname)]
                a = []
                for col in df['TIMESTAMP']:
                    a.append(col[:10])
                df['dates'] = a
                valdate_new = []
                for i in range(len(valdat)):
                    valdate_new.append(valdat[i][:10])
                for j in range(len(valy)):
                    for k in range(len(valdate_new)):
                        a = df[df['VARIABLE_NAME'] == valy[j]]
                        a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                        b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                        b = [i for i in b if i.startswith(valdate_new[k])]
                        fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                   marker=dict(
                                                       line=dict(
                                                           width=0.2,
                                                           color='white')),
                                                   name="{}/{}".format(valy[j], valdate_new[k]))),
                        fig.update_layout(
                            autosize=True,
                            width=sliderw,
                            height=sliderh,
                            margin=dict(
                                l=100,
                                r=60,
                                b=40,
                                t=40,
                                pad=2
                            ),

                            uirevision=valy[j], ),
                return fig
    if realval != None:
        for i in realval:
            fig.add_trace(
                go.Scattergl(x=df2[df2['ID'] == i[16:]]['TIMESTAMP'], y=df2[df2['ID'] == i[16:]]['Value'], mode='lines',
                             marker=dict(line=dict(width=0.2, color='white')), name="{}".format(i),
                             ))
            fig.update_layout(
                autosize=True,
                width=sliderw,
                height=sliderh,
                margin=dict(
                    l=100,
                    r=60,
                    b=40,
                    t=40,
                    pad=2
                ),

                uirevision=i[16:]),
            return fig
    if valy != None:
        if prname == 'rcckn':
            if prch == 'received_variablevalues' or len(df.columns) == 11:
                if df.empty != 1:

                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'LOCAL_TIMESTAMP',
                                  'REMOTE_ID',
                                  'REMOTE_TIMESTAMP', 'REMOTE_MESSAGE_ID', 'PROCESSED', 'TIMED_OUT',
                                  'CONVERTED_NUM_VALUE']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['REMOTE_TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i][:10])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['REMOTE_TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(
                                go.Scattergl(x=b, y=a, mode='markers', marker=dict(line=dict(width=0.2, color='white')),
                                             name="{}/{}".format(valy[j], valdate_new[k]))),
                        fig.update_layout(
                            autosize=True,
                            width=sliderw,
                            height=sliderh,
                            margin=dict(
                                l=100,
                                r=60,
                                b=40,
                                t=40,
                                pad=2
                            ),

                            uirevision=valy[j]),
                    return fig
                else:
                    raise PreventUpdate
            elif prch == 'send_variablevalues' or len(df.columns) == 8:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'TIMESTAMP',
                                  'PROCESSED', 'TIMED_OUT', 'UNREFERENCED']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                       marker=dict(
                                                           line=dict(
                                                               width=0.2,
                                                               color='white')),
                                                       name="{}/{}".format(valy[j], valdate_new[k]))),
                            fig.update_layout(
                                autosize=True,
                                width=sliderw,
                                height=sliderh,
                                margin=dict(
                                    l=100,
                                    r=60,
                                    b=40,
                                    t=40,
                                    pad=2
                                ),

                                uirevision=valy[j], ),
                    return fig
                else:
                    raise PreventUpdate
            elif len(df.columns) == 4:
                if df.empty != 1:

                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                       marker=dict(
                                                           line=dict(
                                                               width=0.2,
                                                               color='white')),
                                                       name="{}/{}".format(valy[j], valdate_new[k]))),
                            fig.update_layout(
                                autosize=True,
                                width=sliderw,
                                height=sliderh,
                                margin=dict(
                                    l=100,
                                    r=60,
                                    b=40,
                                    t=40,
                                    pad=2
                                ),

                                uirevision=valy[j], ),
                    return fig
                else:
                    raise PreventUpdate

        if prname == 'enerbat':
            if df.empty != 1:
                df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
                df = df[df['VARIABLE_NAME'].isin(prvalname)]
                #
                a = []
                for col in df['TIMESTAMP']:
                    a.append(col[:10])
                df['dates'] = a
                valdate_new = []
                for i in range(len(valdat)):
                    valdate_new.append(valdat[i][:10])
                for j in range(len(valy)):
                    for k in range(len(valdate_new)):
                        a = df[df['VARIABLE_NAME'] == valy[j]]
                        a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                        b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                        b = [i for i in b if i.startswith(valdate_new[k])]
                        fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                   marker=dict(
                                                       line=dict(
                                                           width=0.2,
                                                           color='white')),
                                                   name="{}/{}".format(valy[j], valdate_new[k]))),
                        fig.update_layout(

                            width=sliderw,
                            height=sliderh,
                            margin=dict(
                                l=100,
                                r=60,
                                b=40,
                                t=40,
                                pad=2
                            ),

                            uirevision=valy[j], ),
                return fig
            else:
                raise PreventUpdate


@app.callback(Output('getprgraph4', 'figure'),
              [Input('memory-outputpr', 'data'),Input('get_data_from_modbus_pr', 'data'),
               Input('fourgraph_pr_real', 'value'),
               Input('fourgraph_pr_db', 'value'),
               Input('prvaldate', 'value'),
               Input('sliderWidthpr4', 'value'),
               Input('sliderHeightpr4', 'value'),Input('interval_component_pr_db', 'n_intervals') ],
              [State('prvalchoosen', 'value'), State('prname', 'value'),State('prvalname', 'value')  ])
def on_data_set_graph4(data2,data, realval,valy, valdat, sliderw, sliderh,interval, prch, prname, prvalname):
    if realval == None and valy == None:
        raise PreventUpdate
    df = pd.DataFrame(data2)
    df2 = pd.DataFrame(data, columns=['ID', 'Value', 'Quality', 'TIMESTAMP'])
    fig = go.Figure()
    if realval != None and valy != None:
        for i in realval:
            fig.add_trace(
                go.Scattergl(x=df2[df2['ID'] == i[16:]]['TIMESTAMP'], y=df2[df2['ID'] == i[16:]]['Value'], mode='lines',
                             marker=dict(line=dict(width=0.2, color='white')), name="{}".format(i),
                             ))
        if prname == 'rcckn':
            if prch == 'received_variablevalues' or len(df.columns) == 11:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'LOCAL_TIMESTAMP',
                                  'REMOTE_ID',
                                  'REMOTE_TIMESTAMP', 'REMOTE_MESSAGE_ID', 'PROCESSED', 'TIMED_OUT',
                                  'CONVERTED_NUM_VALUE']
                    a = []
                    for col in df['REMOTE_TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i][:10])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['REMOTE_TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(
                                go.Scattergl(x=b, y=a, mode='markers', marker=dict(line=dict(width=0.2, color='white')),
                                             name="{}/{}".format(valy[j], valdate_new[k]))),
                        fig.update_layout(
                            autosize=True,
                            width=sliderw,
                            height=sliderh,
                            margin=dict(
                                l=100,
                                r=60,
                                b=40,
                                t=40,
                                pad=2
                            ),

                            uirevision=valy[j]),
                    return fig
                else:
                    raise PreventUpdate
            elif prch == 'send_variablevalues' or len(df.columns) == 8:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'TIMESTAMP',
                                  'PROCESSED', 'TIMED_OUT', 'UNREFERENCED']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                       marker=dict(
                                                           line=dict(
                                                               width=0.2,
                                                               color='white')),
                                                       name="{}/{}".format(valy[j], valdate_new[k]))),
                            fig.update_layout(
                                autosize=True,
                                width=sliderw,
                                height=sliderh,
                                margin=dict(
                                    l=100,
                                    r=60,
                                    b=40,
                                    t=40,
                                    pad=2
                                ),

                                uirevision=valy[j], ),
                    return fig
                else:
                    raise PreventUpdate
            else:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                       marker=dict(
                                                           line=dict(
                                                               width=0.2,
                                                               color='white')),
                                                       name="{}/{}".format(valy[j], valdate_new[k]))),
                            fig.update_layout(
                                autosize=True,
                                width=sliderw,
                                height=sliderh,
                                margin=dict(
                                    l=100,
                                    r=60,
                                    b=40,
                                    t=40,
                                    pad=2
                                ),

                                uirevision=valy[j], ),
                    return fig
                else:
                    raise PreventUpdate

        if prname == 'enerbat':
            if df.empty != 1:
                df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
                df = df[df['VARIABLE_NAME'].isin(prvalname)]
                a = []
                for col in df['TIMESTAMP']:
                    a.append(col[:10])
                df['dates'] = a
                valdate_new = []
                for i in range(len(valdat)):
                    valdate_new.append(valdat[i][:10])
                for j in range(len(valy)):
                    for k in range(len(valdate_new)):
                        a = df[df['VARIABLE_NAME'] == valy[j]]
                        a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                        b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                        b = [i for i in b if i.startswith(valdate_new[k])]
                        fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                   marker=dict(
                                                       line=dict(
                                                           width=0.2,
                                                           color='white')),
                                                   name="{}/{}".format(valy[j], valdate_new[k]))),
                        fig.update_layout(
                            autosize=True,
                            width=sliderw,
                            height=sliderh,
                            margin=dict(
                                l=100,
                                r=60,
                                b=40,
                                t=40,
                                pad=2
                            ),

                            uirevision=valy[j], ),
                return fig
    if realval != None:
        for i in realval:
            fig.add_trace(
                go.Scattergl(x=df2[df2['ID'] == i[16:]]['TIMESTAMP'], y=df2[df2['ID'] == i[16:]]['Value'], mode='lines',
                             marker=dict(line=dict(width=0.2, color='white')), name="{}".format(i),
                             ))
            fig.update_layout(
                autosize=True,
                width=sliderw,
                height=sliderh,
                margin=dict(
                    l=100,
                    r=60,
                    b=40,
                    t=40,
                    pad=2
                ),

                uirevision=i[16:]),
            return fig
    if valy != None:
        if prname == 'rcckn':
            if prch == 'received_variablevalues' or len(df.columns) == 11:
                if df.empty != 1:

                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'LOCAL_TIMESTAMP',
                                  'REMOTE_ID',
                                  'REMOTE_TIMESTAMP', 'REMOTE_MESSAGE_ID', 'PROCESSED', 'TIMED_OUT',
                                  'CONVERTED_NUM_VALUE']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['REMOTE_TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i][:10])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['REMOTE_TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(
                                go.Scattergl(x=b, y=a, mode='markers', marker=dict(line=dict(width=0.2, color='white')),
                                             name="{}/{}".format(valy[j], valdate_new[k]))),
                        fig.update_layout(
                            autosize=True,
                            width=sliderw,
                            height=sliderh,
                            margin=dict(
                                l=100,
                                r=60,
                                b=40,
                                t=40,
                                pad=2
                            ),

                            uirevision=valy[j]),
                    return fig
                else:
                    raise PreventUpdate
            elif prch == 'send_variablevalues' or len(df.columns) == 8:
                if df.empty != 1:
                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'TIMESTAMP',
                                  'PROCESSED', 'TIMED_OUT', 'UNREFERENCED']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                       marker=dict(
                                                           line=dict(
                                                               width=0.2,
                                                               color='white')),
                                                       name="{}/{}".format(valy[j], valdate_new[k]))),
                            fig.update_layout(
                                autosize=True,
                                width=sliderw,
                                height=sliderh,
                                margin=dict(
                                    l=100,
                                    r=60,
                                    b=40,
                                    t=40,
                                    pad=2
                                ),

                                uirevision=valy[j], ),
                    return fig
                else:
                    raise PreventUpdate
            elif len(df.columns) == 4:
                if df.empty != 1:

                    df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE','TIMESTAMP']
                    df = df[df['VARIABLE_NAME'].isin(prvalname)]
                    a = []
                    for col in df['TIMESTAMP']:
                        a.append(col[:10])
                    df['dates'] = a
                    valdate_new = []
                    for i in range(len(valdat)):
                        valdate_new.append(valdat[i])
                    for j in range(len(valy)):
                        for k in range(len(valdate_new)):
                            a = df[df['VARIABLE_NAME'] == valy[j]]
                            a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                            b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                            b = [i for i in b if i.startswith(valdate_new[k])]
                            fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                       marker=dict(
                                                           line=dict(
                                                               width=0.2,
                                                               color='white')),
                                                       name="{}/{}".format(valy[j], valdate_new[k]))),
                            fig.update_layout(
                                autosize=True,
                                width=sliderw,
                                height=sliderh,
                                margin=dict(
                                    l=100,
                                    r=60,
                                    b=40,
                                    t=40,
                                    pad=2
                                ),

                                uirevision=valy[j], ),
                    return fig
                else:
                    raise PreventUpdate

        if prname == 'enerbat':
            if df.empty != 1:
                df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'TIMESTAMP']
                df = df[df['VARIABLE_NAME'].isin(prvalname)]
                #
                a = []
                for col in df['TIMESTAMP']:
                    a.append(col[:10])
                df['dates'] = a
                valdate_new = []
                for i in range(len(valdat)):
                    valdate_new.append(valdat[i][:10])
                for j in range(len(valy)):
                    for k in range(len(valdate_new)):
                        a = df[df['VARIABLE_NAME'] == valy[j]]
                        a = a[a['dates'].isin(valdate_new)]['VARIABLE_NUM_VALUE']
                        b = df[df['VARIABLE_NAME'] == valy[j]]['TIMESTAMP']
                        b = [i for i in b if i.startswith(valdate_new[k])]
                        fig.add_trace(go.Scattergl(x=b, y=a, mode='markers',
                                                   marker=dict(
                                                       line=dict(
                                                           width=0.2,
                                                           color='white')),
                                                   name="{}/{}".format(valy[j], valdate_new[k]))),
                        fig.update_layout(

                            width=sliderw,
                            height=sliderh,
                            margin=dict(
                                l=100,
                                r=60,
                                b=40,
                                t=40,
                                pad=2
                            ),

                            uirevision=valy[j], ),
                return fig
            else:
                raise PreventUpdate

if __name__ == '__main__':
    app.run_server(debug=True, host='0.0.0.0', port=8049)